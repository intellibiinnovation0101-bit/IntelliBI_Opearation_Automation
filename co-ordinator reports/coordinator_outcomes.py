"""
================================================================================
  IntelliBI Operations Automation
  COORDINATOR ACTUAL PERFORMANCE (OUTCOMES)
  (co-ordinator reports / coordinator_outcomes.py)
  ------------------------------------------------------------------------------
  The "what result did the work produce?" half of the Coordinator Task
  Performance report (pyCoordinatorTaskPerformanceReport.py):

        Effort (tasks generated / completed)  ->  Outcome (actual performance)

  One measure per task group. Every measure REUSES the calculation of the report
  that owns it — nothing is re-defined here:

    Attendance   "Overall Att %" of the Attendance report's Session Summary
                 (pyAttendaceFeedbackReport.overall_att_kpi — the function the
                 Session Summary itself calls).
                   Daily  : the Daily report window, yesterday 12:00 PM -> the
                            time of that day's Coordinator task list (daily formula)
                   Period : sessions dated in the period (period formula)
    Assignment   "Total Submission %" READ from the corresponding generated
                 Assignment Submission Performance report (its Summary tab):
                   Daily D : that morning's Daily report (deadlines of D-1 — the
                             report the Assignment script generates on day D)
                   Weekly / Monthly / Manual : the report of the same period
                 Not generated yet / unreadable -> calculated now with that
                 report's own functions (build_dataset + summarise) for the same
                 period and clearly labelled "provisional".
    Admission    Resolution % = learners on the day's Admission Formalities list
                 whose admission form is SIGNED at the evening re-check
                 (pyAdmissionFormalitiesReport.build_rows on refreshed signer data)
                 ÷ learners on the list that are still applicable.
                 A follow-up marked Done does NOT count — only the signature does.
    Wise & IV    Resolution % = validation issues on the day's Wise & Interview
                 Feedback Validation list that are FIXED at the evening re-check
                 (pyWiseDataValidationReport rules re-run on refreshed data; the
                 interview feedback re-checked in the Interview Consolidate Sheet)
                 ÷ issues still applicable. Issue level: a student with a missing
                 email AND an invalid phone is 2 issues.
    Instructor   Sessions without escalation % = held sessions of the day's
                 window NOT on the Instructor Follow-Ups list ÷ held sessions.
                 One session counts once however many reasons it was flagged
                 for; the reasons are shown separately.
    Interview    "Overall Interview Attendance %" of the IntelliBI Interview
                 Consolidated Report (interviewed ÷ scheduled, same attended rule).
                   Daily  : interviews dated yesterday .. today
                   Weekly : Monday .. today (or the week's Sunday)
                   Monthly: the 1st .. today (or the month's last day)

  "No longer applicable" items (learner no longer active, record deleted) are
  shown but excluded from the denominator. An item whose source could not be
  read is "Not checked" — excluded and reported, never guessed.

  Admission / Wise need the EVENING state of the morning list, which cannot be
  re-created later. Each delivered Daily performance report therefore stores
  every checked item on the reporting PC (cache/coordinator_outcome_checks, one
  JSON per day); Weekly / Monthly reports read those stored checks back (the day
  being reported today is checked live). Reports made before 07-Oct-2026 kept
  them in an "Outcome Detail" tab, which is still read for those days.

  This module holds the pure calculators (unit-tested with synthetic data) and
  OutcomeEngine, which applies them to periods. All Google access lives in
  GoogleOutcomeSources, so tests use fake sources.
================================================================================
"""
from __future__ import annotations

import logging
import os
import re
import unicodedata
from collections import OrderedDict, defaultdict
from datetime import date, datetime, time, timedelta

log = logging.getLogger("CoordinatorOutcomes")

# =============================================================================
#  DEFINITIONS
# =============================================================================
GROUP_ORDER = ["attendance", "assignment", "admission", "wise", "instructor", "interview"]

# key -> (measure name, short definition, source report)
MEASURES = OrderedDict([
    ("attendance", ("Overall Att %",
                    "Present ÷ attendance-applicable learner-sessions",
                    "Attendance & Feedback report · Session Summary")),
    ("assignment", ("Total Submission %",
                    "Submitted ÷ expected submissions of the assignments due",
                    "Assignment Submission Performance report")),
    ("admission", ("Formalities Resolved %",
                   "Morning-list learners whose admission form is signed by the evening check",
                   "Admission Formalities (Zoho Sign) — evening re-check")),
    ("wise", ("Issues Resolved %",
              "Morning-list validation issues fixed by the evening check",
              "Wise Data Validation + Interview Consolidate Sheet — evening re-check")),
    ("instructor", ("Sessions Without Escalation %",
                    "Held sessions NOT on the Instructor Follow-Ups list ÷ held sessions",
                    "Attendance data + Instructor Follow-Ups list")),
    ("interview", ("Overall Interview Attendance %",
                   "Interviewed ÷ scheduled interviews",
                   "IntelliBI Interview Consolidated Report")),
])

# basis wording: (numerator noun, denominator noun)
BASIS = {
    "attendance": ("present", "applicable learner-sessions"),
    "assignment": ("submitted", "expected submissions"),
    "admission": ("resolved", "pending learners"),
    "wise": ("fixed", "issues"),
    "instructor": ("without escalation", "held sessions"),
    "interview": ("interviewed", "scheduled"),
}

RES_RESOLVED = "Resolved"
RES_PENDING = "Still pending"
RES_NA = "No longer applicable"
RES_NOT_CHECKED = "Not checked"
RES_ESCALATED = "Escalated"
RES_CLEAN = "No escalation"

ST_MEASURED = "measured"
ST_NO_DATA = "no_data"            # nothing to measure in the period (e.g. no interviews)
ST_NOT_CHECKED = "not_checked"    # source unreadable / evening check not stored

DETAIL_TAB = "Outcome Detail"        # legacy tab name (reports before 07-Oct-2026)
DETAIL_COLS = ["Report Day", "Task Group", "Section", "Item", "Issue", "Morning Status",
               "Evening Status", "Result", "Checked At", "Issue Key"]

# Instructor Follow-Ups "Why Flagged" reasons (pyCoordinatorTaskListReport
# INSTR_REASON_LABELS) — matched on the reason's own leading words.
INSTR_REASON_TEXT = [
    ("session cancelled", "cancelled"),
    ("started late", "started_late"),
    ("not started 5 min early", "not_early"),
    ("session underrun", "underrun"),
    ("feedback missing", "feedback_missing"),
    ("low feedback rate", "low_feedback_rate"),
]
INSTR_REASON_LABELS = {
    "cancelled": "Session cancelled",
    "started_late": "Started late",
    "not_early": "Not started 5 min early",
    "underrun": "Session underrun (>= 30 min short)",
    "feedback_missing": "Instructor feedback missing",
    "low_feedback_rate": "Low student feedback rate",
}

# Wise & Interview Feedback Validation tab sections (banner titles) and the
# status columns each one carries (the tab's own data headers).
WISE_STATUS_VALUES = {"invalid", "missing", "warning"}       # = the tab's _WISE_SEV keys
WISE_SECTIONS = {
    "student": ("STUDENT VALIDATION", ["Name", "Email", "Phone", "Tag", "Note", "Picture"]),
    "course": ("COURSE VALIDATION", ["Title", "Subtitle", "Tag"]),
    "instructor": ("INSTRUCTOR VALIDATION", []),
    "interview": ("INTERVIEW FEEDBACK NOT COMPLETED", []),
}
WISE_SECTION_LABEL = {"student": "Student", "course": "Course", "instructor": "Instructor",
                      "interview": "Interview feedback"}


# =============================================================================
#  SMALL HELPERS
# =============================================================================
def _n(v) -> str:
    """Comparable text: NFKC, case-folded, single spaces; '—' / 'nan' = blank."""
    s = unicodedata.normalize("NFKC", str(v if v is not None else "")).strip()
    if s in ("—", "-", "nan", "None", "NaT"):
        return ""
    return " ".join(s.split()).casefold()


def _blank(v) -> str:
    s = str(v if v is not None else "").strip()
    return "" if s in ("—", "nan", "None") else s


def pct(num, den, nd=1):
    return round(num / den * 100.0, nd) if den else None


def outcome(key, num, den, pct_value=None, *, state=None, note="", breakdown=None, extra=None):
    """One group's actual-performance figure for one period."""
    if state is None:
        state = ST_MEASURED if den else ST_NO_DATA
    if state != ST_MEASURED:
        p = None
    else:
        p = pct_value if pct_value is not None else pct(num, den)
    o = {"key": key, "measure": MEASURES[key][0], "num": int(num or 0), "den": int(den or 0),
         "pct": p, "state": state, "note": note, "breakdown": list(breakdown or [])}
    if extra:
        o.update(extra)
    return o


def not_checked(key, note):
    return outcome(key, 0, 0, state=ST_NOT_CHECKED, note=note)


def basis_text(o) -> str:
    if o is None:
        return ""
    if o["state"] == ST_NOT_CHECKED:
        return "not checked"
    if o["state"] == ST_NO_DATA:
        return "nothing to measure"
    a, b = BASIS[o["key"]]
    return f"{o['num']} {a} of {o['den']} {b}"


def detail(day, group, section, item, issue, morning, evening, result, checked_at, key):
    return {"day": day, "group": group, "section": section, "item": item, "issue": issue,
            "morning": morning, "evening": evening, "result": result,
            "checked_at": checked_at, "key": key}


def detail_row_values(d) -> list:
    """Detail dict -> a stored row (DETAIL_COLS order)."""
    ca = d.get("checked_at")
    return [d["day"].strftime("%Y-%m-%d") if isinstance(d["day"], date) else str(d["day"]),
            d["group"], d["section"], d["item"], d["issue"], d["morning"], d["evening"],
            d["result"], ca.strftime("%Y-%m-%d %H:%M") if isinstance(ca, datetime) else (ca or ""),
            d["key"]]


def detail_from_values(vals) -> dict | None:
    """Stored row -> detail dict (None for a malformed row)."""
    vals = list(vals) + [""] * (len(DETAIL_COLS) - len(vals))
    try:
        day = datetime.strptime(str(vals[0]).strip()[:10], "%Y-%m-%d").date()
    except ValueError:
        return None
    group = str(vals[1] or "").strip()
    if group not in GROUP_ORDER:
        return None
    return detail(day, group, *(str(v or "") for v in vals[2:8]), str(vals[8] or ""),
                  str(vals[9] or ""))


def resolution_outcome(key, details, note=""):
    """Admission / Wise: Resolved ÷ (Resolved + Still pending). No longer
    applicable and Not checked items are excluded (and shown in the breakdown)."""
    c = defaultdict(int)
    for d in details:
        c[d["result"]] += 1
    den = c[RES_RESOLVED] + c[RES_PENDING]
    br = [(RES_RESOLVED, c[RES_RESOLVED]), (RES_PENDING, c[RES_PENDING]),
          (RES_NA, c[RES_NA]), (RES_NOT_CHECKED, c[RES_NOT_CHECKED])]
    if not den and c[RES_NOT_CHECKED]:
        return outcome(key, 0, 0, state=ST_NOT_CHECKED, breakdown=br,
                       note=note or "source could not be read at the evening check")
    return outcome(key, c[RES_RESOLVED], den, breakdown=br, note=note)


# =============================================================================
#  1. ATTENDANCE  (AR.overall_att_kpi — the Session Summary's own function)
# =============================================================================
def attendance_daily_window(day: date, cutoff: datetime):
    """The Attendance Daily report window: yesterday 12:00 PM -> report time."""
    return datetime.combine(day - timedelta(days=1), time(12, 0)), cutoff


def filter_window(df, start_dt, end_dt):
    if df is None or df.empty or "_dt" not in df.columns:
        return df.iloc[0:0].copy() if df is not None else df
    import pandas as pd
    m = df["_dt"].apply(lambda x: x is not None and not pd.isna(x) and start_dt <= x <= end_dt)
    return df[m].copy()


def filter_dates(df, start: date, end: date):
    if df is None or df.empty or "_date" not in df.columns:
        return df.iloc[0:0].copy() if df is not None else df
    import pandas as pd
    m = df["_date"].apply(lambda d: d is not None and not pd.isna(d) and start <= d <= end)
    return df[m].copy()


def attendance_outcome(AR, att_f, period: bool):
    k = AR.overall_att_kpi(att_f, period=period)
    return outcome("attendance", k["present"], k["applicable"],
                   k["pct"] if k["applicable"] else None,
                   extra={"absent": k["absent"]})


# =============================================================================
#  2. ASSIGNMENT  (pyAssignmentSubmissionPerformanceReport.build_dataset/summarise)
# =============================================================================
def assignment_report_job(CP, kind: str, start: date, end: date) -> dict:
    """The Assignment Submission Performance report that corresponds to a
    Coordinator report period. The Assignment script's AUTO plan makes, on run
    day D, the Daily report of D-1 (deadlines that have fully passed), the
    Weekly report of the previous Mon-Sun on Monday and the Monthly report of the
    previous month on the 1st. So:
        Coordinator Daily D            -> Assignment Daily  D-1 (generated on D)
        Coordinator Weekly / Monthly / Manual -> Assignment report of the SAME period"""
    if kind == "Daily":
        d = end - timedelta(days=1)
        return CP.make_job("Daily", d, d)
    return CP.make_job(kind, start, end)


def _num_cell(v):
    """'71.6%' / 71.6 / '—' -> 71.6 / 71.6 / None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("%", "").replace(",", "")
    try:
        return float(t)
    except ValueError:
        return None


def parse_assignment_summary(wb) -> dict | None:
    """The KPI strip of an Assignment Submission Performance report's Summary tab
    (label row, value row underneath): Submission %, Submissions Expected,
    Submitted, Eligible Assignments, plus the title line. Located by LABEL, not
    by cell position. None when the tab / the Submission % KPI is not there."""
    ws = wb["Summary"] if "Summary" in wb.sheetnames else None
    if ws is None:
        return None
    want = {"submission %": "pct", "submissions expected": "expected",
            "submitted": "submitted", "eligible assignments": "assignments"}
    out = {"title": str(ws.cell(row=1, column=1).value or "").strip()}
    for r in range(1, min(ws.max_row or 0, 15) + 1):
        for c in range(1, (ws.max_column or 0) + 1):
            label = _n(re.sub(r"^[^A-Za-z]+", "", str(ws.cell(row=r, column=c).value or "")))
            key = want.get(label)
            if key and key not in out:
                out[key] = ws.cell(row=r + 1, column=c).value
        if "pct" in out:
            break
    if "pct" not in out:
        return None
    pct_v = _num_cell(out["pct"])
    return {"pct": pct_v, "no_assignments": pct_v is None,
            "expected": int(_num_cell(out.get("expected")) or 0),
            "submitted": int(_num_cell(out.get("submitted")) or 0),
            "assignments": int(_num_cell(out.get("assignments")) or 0),
            "title": out["title"]}


def assignment_from_report(rep, label):
    """Outcome from a generated Assignment report (parse_assignment_summary)."""
    gen = ""
    m = re.search(r"Generated On:\s*(.+)$", rep.get("title", ""))
    if m:
        gen = f", generated {m.group(1).strip()}"
    if rep["no_assignments"]:
        return outcome("assignment", 0, 0, state=ST_NO_DATA,
                       note=f"Assignment report {label}{gen}: no assignment deadline had passed")
    return outcome("assignment", rep["submitted"], rep["expected"], rep["pct"],
                   note=f"from Assignment report {label}{gen}",
                   extra={"assignments": rep["assignments"], "source": "report"})


def assignment_outcome(ASR, subs_df, start: date, end: date, now: datetime, daily: bool = False,
                       reason=""):
    """PROVISIONAL figure, used only when the corresponding Assignment report is
    not available: the Assignment report's own functions (build_dataset +
    summarise) for the same period, deadlines passed by now."""
    ds = ASR.build_dataset(subs_df, start, end, now)
    s = ASR.summarise(ds)
    note = (f"PROVISIONAL — {reason}; calculated now with the Assignment report's rules"
            if reason else "calculated with the Assignment report's rules")
    if s.get("not_yet_due"):
        note += f"; {s['not_yet_due']} assignment(s) not yet due excluded"
    return outcome("assignment", s["submitted"], s["expected"], s["rate"], note=note,
                   extra={"assignments": s["assignments"], "source": "provisional"})


# =============================================================================
#  3. ADMISSION FORMALITIES  (evening re-check of the morning list)
# =============================================================================
def admission_check(day, tasks, evening_rows, checked_at, norm_email, norm):
    """Morning admission tasks (ledger tasks of the day) -> detail rows.
    evening_rows = pyAdmissionFormalitiesReport.build_rows(current_only=False)
    on refreshed signer data (every ACTIVE student, signed or not), or None
    when the sources could not be read."""
    by_email, by_name = {}, {}
    for r in evening_rows or []:
        e = norm_email(r.get("Email ID"))
        if e:
            by_email.setdefault(e, r)
        nm = norm(r.get("Student Name"))
        if nm:
            by_name.setdefault(nm, r)
    out = []
    for t in tasks:
        c = t.get("cells") or {}
        name, email = _blank(c.get("Student Name")), _blank(c.get("Email ID"))
        morning = _blank(c.get("Recipient Status")) or "Not signed"
        item = "  ·  ".join(x for x in (name, email) if x) or t.get("label", "")
        if evening_rows is None:
            res, eve = RES_NOT_CHECKED, "admission sources unreadable"
        else:
            r = by_email.get(norm_email(email)) if email else None
            if r is None and name:
                r = by_name.get(norm(name))
            if r is None:
                res, eve = RES_NA, "no longer an active student"
            elif norm(r.get("Recipient Status")) == "signed":
                signed = _blank(r.get("Signed Date"))
                res, eve = RES_RESOLVED, "Signed" + (f" ({signed})" if signed else "")
            else:
                res, eve = RES_PENDING, _blank(r.get("Recipient Status")) or "Form not sent"
        out.append(detail(day, "admission", "Admission Formalities", item,
                          "Admission form not signed", morning, eve, res, checked_at, t["key"]))
    return out


# =============================================================================
#  4. WISE & INTERVIEW FEEDBACK VALIDATION  (issue-level evening re-check)
# =============================================================================
def _wise_section_of(task) -> str | None:
    sec = _n(task.get("section", ""))
    for k, (title, _f) in WISE_SECTIONS.items():
        if sec.startswith(title.casefold()):
            return k
    # earlier layout without banners: infer from the columns present
    cells = {_n(k) for k in (task.get("cells") or {})}
    if "interviewer name" in cells:
        return "interview"
    if "instructor id" in cells:
        return "instructor"
    if "course title" in cells:
        return "course"
    if "student name" in cells:
        return "student"
    return None


def _cell(cells, name):
    for k, v in (cells or {}).items():
        if _n(k) == _n(name):
            return _blank(v)
    return ""


def wise_morning_issues(tasks) -> list:
    """Morning Wise tasks -> one issue per failing field (student / course), per
    failed check (instructor), per missing interview feedback."""
    out = []
    for t in tasks:
        sec = _wise_section_of(t)
        c = t.get("cells") or {}
        if sec == "student":
            name, batch = _cell(c, "Student Name"), _cell(c, "Batch Name")
            match = {"name": _n(name), "batch": _n(batch), "joined": _n(_cell(c, "Joined On"))}
            for f in WISE_SECTIONS["student"][1]:
                st = _cell(c, f)
                if _n(st) in WISE_STATUS_VALUES:
                    out.append({"section": sec, "item": f"{name}  ·  {batch}", "field": f,
                                "issue": f"{f}: {st}", "morning": st, "match": match,
                                "key": f"{t['key']}|{f}", "task_key": t["key"]})
        elif sec == "course":
            title, sub = _cell(c, "Course Title"), _cell(c, "Course Subtitle")
            match = {"title": _n(title), "subtitle": _n(sub), "created": _n(_cell(c, "Created On"))}
            for f in WISE_SECTIONS["course"][1]:
                st = _cell(c, f)
                if _n(st) in WISE_STATUS_VALUES:
                    out.append({"section": sec, "item": f"{title}  ·  {sub}", "field": f,
                                "issue": f"{f}: {st}", "morning": st, "match": match,
                                "key": f"{t['key']}|{f}", "task_key": t["key"]})
        elif sec == "instructor":
            iid, iname = _cell(c, "Instructor ID"), _cell(c, "Instructor Name")
            try:
                n = max(1, int(float(_cell(c, "Failed Checks") or 1)))
            except ValueError:
                n = 1
            for i in range(1, n + 1):
                out.append({"section": sec, "item": f"{iname}  ·  {iid}", "field": str(i),
                            "issue": f"Failed check {i} of {n}", "morning": f"{n} failed check(s)",
                            "match": {"id": _n(iid), "name": _n(iname), "n": n, "i": i},
                            "key": f"{t['key']}|check{i}", "task_key": t["key"]})
        elif sec == "interview":
            start = _cell(c, "Interview Start Date")
            batch, title = _cell(c, "Batch Name"), _cell(c, "Batch Title / Duration")
            out.append({"section": sec, "item": f"{batch}  ·  {title}  ·  {start}",
                        "field": "feedback", "issue": "Interview feedback not completed",
                        "morning": "Not in Consolidate Sheet",
                        "match": {"start": start, "batch": batch, "title": title},
                        "key": t["key"], "task_key": t["key"]})
    return out


def wise_evening_state(TL, data, interview_index, WISE=None):
    """Evening state from the re-run Wise validation (TL.load_wise_validation
    with return_sources=True, or None when unavailable) + the Interview
    Consolidate index (TL._iv_consolidate_index, None when unreadable)."""
    st = {"available": data is not None, "student": {}, "student_by_joined": {},
          "student_applicable": None, "course": {}, "course_by_created": {},
          "course_applicable": None, "instructor": {}, "instructor_applicable": None,
          "interview_index": interview_index}
    if data is None:
        return st
    # Joined On / Created On: the fallback identity when the NAME itself was
    # corrected during the day — used only when the timestamp is unique.
    stu_fields = WISE_SECTIONS["student"][1]
    joined = defaultdict(list)
    for d in TL._wise_student_display(data.get("student", [])):
        k = (_n(d["cells"][1]), _n(d["cells"][2]))
        statuses = {stu_fields[i - 3]: d["status_cells"].get(i, "") for i in range(3, 9)}
        st["student"][k] = statuses
        j = _n(d["cells"][9]) if len(d["cells"]) > 9 else ""
        if j:
            joined[j].append(k)
    st["student_by_joined"] = {j: ks[0] for j, ks in joined.items() if len(ks) == 1}
    crs_fields = WISE_SECTIONS["course"][1]
    created = defaultdict(list)
    for d in TL._wise_course_display(data.get("course", [])):
        k = (_n(d["cells"][1]), _n(d["cells"][2]))
        st["course"][k] = {crs_fields[i - 4]: d["status_cells"].get(i, "") for i in range(4, 7)}
        cr = _n(d["cells"][7]) if len(d["cells"]) > 7 else ""
        if cr:
            created[cr].append(k)
    st["course_by_created"] = {c: ks[0] for c, ks in created.items() if len(ks) == 1}
    for d in TL._wise_instructor_display(data.get("instructor", [])):
        st["instructor"][_n(d["cells"][1]) or _n(d["cells"][2])] = int(d["cells"][3] or 0)
    src = data.get("_sources") or {}
    if src.get("student") is not None:
        app, jset = set(), set()
        for r in src["student"]:
            if str(r.get("Is_Deleted", "")).strip().upper() == "Y":
                continue
            if str(r.get("is_candidate_active", "")).strip().upper() == "N":
                continue
            if WISE is not None:
                try:
                    if WISE._is_excluded_student(r.get("student_name", ""), r.get("candidate_name", "")):
                        continue
                    if WISE._is_excluded_batch(r.get("batch_name", "")):
                        continue
                except Exception:                                   # pragma: no cover
                    pass
            app.add((_n(r.get("student_name", "")), _n(r.get("batch_name", ""))))
            j = r.get("joined_on", "")
            if j:                     # (Joined On as the tab shows it, batch)
                try:
                    jset.add((_n(TL._wise_joined_on_ist(j)), _n(r.get("batch_name", ""))))
                except Exception:                                   # pragma: no cover
                    jset.add((_n(j), _n(r.get("batch_name", ""))))
        st["student_applicable"] = app
        st["student_joined_applicable"] = jset
    if src.get("combined") is not None:
        st["course_applicable"] = {(_n(r.get("class_name", "")), _n(r.get("class_subject", "")))
                                   for r in src["combined"]}
    if src.get("instructor") is not None:
        st["instructor_applicable"] = {_n(r.get("instructor_id", "")) for r in src["instructor"]
                                       if str(r.get("Is_Active", "Y")).strip().upper() == "Y"}
    return st


def _iv_dates(start_txt):
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(start_txt).strip(), fmt).date()
        except ValueError:
            continue
    return None


def wise_check(day, issues, evening, checked_at, norm_name=None):
    """Morning issues + evening state -> detail rows."""
    norm_name = norm_name or _n
    out = []
    for iss in issues:
        sec, m = iss["section"], iss["match"]
        res, eve = RES_NOT_CHECKED, "Wise validation unavailable"
        if sec == "interview":
            idx = evening.get("interview_index")
            if idx is None:
                res, eve = RES_NOT_CHECKED, "Interview Consolidate Sheet unreadable"
            else:
                d0 = _iv_dates(m["start"])
                found = False
                if d0 is not None:
                    d = d0
                    while d <= day and not found:
                        found = (norm_name(m["batch"]), norm_name(m["title"]), d) in idx
                        d += timedelta(days=1)
                res, eve = ((RES_RESOLVED, "Feedback now in Consolidate Sheet") if found
                            else (RES_PENDING, "Still missing"))
        elif evening.get("available"):
            if sec == "student":
                k = (m["name"], m["batch"])
                cur = evening["student"].get(k)
                app = evening.get("student_applicable")
                # the name was corrected (old key gone from the source): follow Joined On
                if cur is None and m.get("joined") and (app is None or k not in app):
                    k2 = evening["student_by_joined"].get(m["joined"])
                    cur = evening["student"].get(k2) if k2 else None
                if cur is not None:
                    s = _blank(cur.get(iss["field"], ""))
                    res, eve = ((RES_PENDING, s) if _n(s) in WISE_STATUS_VALUES
                                else (RES_RESOLVED, s or "Valid"))
                elif app is not None and k not in app:
                    # gone under this name: renamed (same Joined On, still listed) = fixed;
                    # otherwise deleted / inactive = no longer applicable
                    if m.get("joined") and (m["joined"], m["batch"]) in evening.get(
                            "student_joined_applicable", set()):
                        res, eve = RES_RESOLVED, "record corrected"
                    else:
                        res, eve = RES_NA, "student no longer active / listed"
                else:
                    res, eve = RES_RESOLVED, "Valid"
            elif sec == "course":
                k = (m["title"], m["subtitle"])
                cur = evening["course"].get(k)
                capp = evening.get("course_applicable")
                if cur is None and m.get("created") and (capp is None or k not in capp):
                    k2 = evening["course_by_created"].get(m["created"])
                    cur = evening["course"].get(k2) if k2 else None
                if cur is not None:
                    s = _blank(cur.get(iss["field"], ""))
                    res, eve = ((RES_PENDING, s) if _n(s) in WISE_STATUS_VALUES
                                else (RES_RESOLVED, s or "Valid"))
                else:
                    app = evening.get("course_applicable")
                    res, eve = ((RES_NA, "course no longer listed")
                                if app is not None and k not in app else (RES_RESOLVED, "Valid"))
            elif sec == "instructor":
                k = m["id"] or m["name"]
                n_now = evening["instructor"].get(k)
                if n_now is None:
                    app = evening.get("instructor_applicable")
                    if app is not None and m["id"] and m["id"] not in app:
                        res, eve = RES_NA, "instructor no longer active"
                    else:
                        res, eve = RES_RESOLVED, "no failed checks"
                else:
                    # n_now checks still fail: the first n_now of the morning's n stay pending
                    res = RES_PENDING if m["i"] <= n_now else RES_RESOLVED
                    eve = f"{n_now} failed check(s)"
        out.append(detail(day, "wise", WISE_SECTION_LABEL[sec], iss["item"], iss["issue"],
                          iss["morning"], eve, res, checked_at, iss["key"]))
    return out


# =============================================================================
#  5. INSTRUCTOR  (session level — one session counts once)
# =============================================================================
def instructor_reason_codes(why_text: str) -> list:
    t = _n(why_text)
    return [code for txt, code in INSTR_REASON_TEXT if txt in t]


def instructor_day(day, tasks, reviews, checked_at, keyfn=None):
    """Held sessions of the day's window (reviews = TL.review_instructor_sessions,
    type != 'scheduled') vs the day's Instructor Follow-Ups tasks (each = one
    flagged session). Returns (outcome, detail rows)."""
    held = [r for r in (reviews or []) if r.get("type") != "scheduled"]
    flagged = list(tasks or [])
    den = max(len(held), len(flagged))
    reasons = defaultdict(int)
    out = []
    flagged_keys = set()
    for t in flagged:
        codes = instructor_reason_codes((t.get("cells") or {}).get("Why Flagged", "") or t.get("what", ""))
        for c in codes:
            reasons[c] += 1
        flagged_keys.add(t["key"])
        out.append(detail(day, "instructor", "Instructor Follow-Ups", t.get("label", ""),
                          "; ".join(INSTR_REASON_LABELS[c] for c in codes) or (t.get("what") or ""),
                          "Flagged", "", RES_ESCALATED, checked_at, t["key"]))
    if keyfn is not None:
        for r in held:
            k = keyfn(r)
            if k in flagged_keys:
                continue
            item = "  ·  ".join(x for x in (str(r.get("instr", "")),
                                            str(r["sess"].get("course_name", "")),
                                            r["actual_start_dt"].strftime("%Y-%m-%d %H:%M")
                                            if r.get("actual_start_dt") is not None else "") if x)
            out.append(detail(day, "instructor", "Instructor Follow-Ups", item, "",
                              "Held", "", RES_CLEAN, checked_at, k))
    br = [(INSTR_REASON_LABELS[c], reasons[c]) for c in INSTR_REASON_LABELS if reasons[c]]
    o = outcome("instructor", den - len(flagged), den, breakdown=br,
                extra={"held": den, "flagged": len(flagged), "reasons": dict(reasons)})
    return o, out


# =============================================================================
#  6. INTERVIEW ATTENDANCE  (IntelliBI Interview Consolidated Report rules)
# =============================================================================
# Same header resolution / attended rule as Reports/pyInterviewConsolidatedReport.py
IV_HEADER_ROW = 2
IV_COLUMN_ALIASES = OrderedDict([
    ("batch_id", ["batch id"]),
    ("interview_date", ["interview date", "date"]),
    ("interviewer", ["interviewer", "instructor", "interviewer name"]),
    ("tech_stack", ["tech stack", "technology", "tech"]),
    ("batch_name", ["batch name"]),
    ("batch_title", ["batch title"]),
    ("slot", ["slot"]),
    ("candidate", ["candidate name", "student name", "candidate", "student"]),
    ("email", ["email"]),
    ("phone", ["phone", "mobile"]),
    ("communication", ["communication"]),
    ("tech_scores", ["tech scores", "tech score"]),
    ("total_score", ["total score"]),
    ("avg_score", ["average score", "avg score", "feedback rating", "rating"]),
    ("zone", ["zone", "status", "result"]),
    ("comments", ["comments", "remark", "feedback"]),
    ("published_on", ["published on", "published"]),
    ("start_time", ["start time", "start"]),
    ("end_time", ["end time", "end"]),
    ("duration", ["duration", "interview duration"]),
])
_IV_DATE_FORMATS = ("%Y-%m-%d", "%d-%b-%Y", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y",
                    "%Y-%m-%d %H:%M:%S", "%d-%b-%Y %H:%M", "%d %b %Y", "%b %d, %Y",
                    "%Y/%m/%d", "%d-%b-%y")


def _iv_resolve_columns(headers):
    norm = {h: str(h).strip().lower() for h in headers if str(h).strip()}
    resolved, used = {}, set()
    for canon, aliases in IV_COLUMN_ALIASES.items():
        for alias in aliases:
            hit = next((h for h, hl in norm.items() if hl == alias and h not in used), None)
            if hit:
                resolved[canon] = hit
                used.add(hit)
                break
        if canon in resolved:
            continue
        for alias in aliases:
            hit = next((h for h, hl in norm.items() if alias in hl and h not in used), None)
            if hit:
                resolved[canon] = hit
                used.add(hit)
                break
    return resolved


def _iv_float(v):
    try:
        s = str(v).strip().replace("%", "")
        if s in ("", "-", "—", "nan", "None"):
            return 0.0
        if "/" in s:
            s = s.split("/")[0].strip()
        return float(s)
    except Exception:
        return 0.0


def _iv_norm_zone(z):
    zl = str(z).strip().lower()
    if not zl:
        return ""
    if "absent" in zl or "skip" in zl or "no show" in zl or "noshow" in zl:
        return "Absent"
    return "Other"


def _iv_parse_date(s):
    s = str(s).strip()
    if not s:
        return None
    s = s.split("T")[0].strip()
    for fmt in _IV_DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    try:
        import pandas as pd
        d = pd.to_datetime(s, errors="coerce", dayfirst=True)
        return None if pd.isna(d) else d.date()
    except Exception:
        return None


def interview_rows(grid):
    """Consolidated Report grid (row 1 title, row 2 headers, data from row 3) ->
    [{"date", "attended"}] with the source report's own rules."""
    if not grid or len(grid) < IV_HEADER_ROW:
        return []
    headers = [str(h).strip() for h in grid[IV_HEADER_ROW - 1]]
    col = _iv_resolve_columns(headers)
    idx = {c: headers.index(h) for c, h in col.items()}

    def g(row, canon):
        i = idx.get(canon)
        return str(row[i]).strip() if i is not None and i < len(row) else ""
    out = []
    for row in grid[IV_HEADER_ROW:]:
        if not (g(row, "batch_id") or g(row, "candidate") or g(row, "slot")):
            continue                                    # fully empty row
        has_score = _iv_float(g(row, "avg_score")) > 0 or _iv_float(g(row, "total_score")) > 0
        published = g(row, "published_on") != ""
        attended = _iv_norm_zone(g(row, "zone")) != "Absent" and (has_score or published)
        out.append({"date": _iv_parse_date(g(row, "interview_date")), "attended": attended})
    return out


def interview_outcome(rows, start: date, end: date, note=""):
    sel = [r for r in rows if r["date"] is not None and start <= r["date"] <= end]
    att = sum(1 for r in sel if r["attended"])
    return outcome("interview", att, len(sel), pct(att, len(sel)), note=note,
                   extra={"skipped": len(sel) - att})


def interview_window(kind: str, start: date, end: date, today: date):
    """Daily: yesterday .. the day; Weekly: Monday .. today (or Sunday);
    Monthly / Manual: first day .. today (or the last day)."""
    if kind == "Daily":
        return end - timedelta(days=1), end
    return start, min(end, today)


# =============================================================================
#  STATUS / QUADRANT
# =============================================================================
QUAD_PAYING = "Effort is paying off"
QUAD_NOT_CONVERTING = "Effort not converting"
QUAD_CHECK_TASKS = "Check if tasks needed"
QUAD_ATTENTION = "Needs management attention"
QUAD_NA = "—"
QUAD_NO_TASKS_OK = "On target, no tasks needed"
QUAD_NO_TASKS_LOW = "Below target, no tasks raised"


def outcome_status(p, target, near_band=15.0):
    if p is None:
        return "Not measured"
    if p >= target:
        return "On target"
    if p >= target - near_band:
        return "Near target"
    return "Below target"


def quadrant(completion_pct, actual_pct, target, effort_high):
    """High/Low Effort × Good/Low Performance. Effort = Task Completion %.
    No tasks generated in the period: on target = nothing was needed; below
    target = the result is poor but no follow-up was raised (task rules?)."""
    if actual_pct is None:
        return QUAD_NA
    if completion_pct is None:                      # no Coordinator task in the period
        return QUAD_NO_TASKS_OK if actual_pct >= target else QUAD_NO_TASKS_LOW
    hi_eff = completion_pct >= effort_high
    good = actual_pct >= target
    if hi_eff and good:
        return QUAD_PAYING
    if hi_eff:
        return QUAD_NOT_CONVERTING
    if good:
        return QUAD_CHECK_TASKS
    return QUAD_ATTENTION


QUAD_LEVEL = {QUAD_PAYING: "ok", QUAD_NOT_CONVERTING: "medium", QUAD_CHECK_TASKS: "info",
              QUAD_ATTENTION: "high", QUAD_NA: "muted", QUAD_NO_TASKS_OK: "ok",
              QUAD_NO_TASKS_LOW: "medium"}


# =============================================================================
#  PREVIOUS vs CURRENT PERIOD  (Yesterday → Today for the follow-up groups)
# =============================================================================
# Attendance, Assignment and Instructor outcomes are measured over a window that
# STARTS BEFORE the day's task list is worked (the Daily window is yesterday
# 12:00 PM → the task-list time; assignments due the day before). The day's own
# tasks are raised FROM that window, so the result that follows a day's
# follow-ups is the NEXT day's figure. The comparison therefore shows the
# previous period's effort and outcome next to the current one, each labelled
# with the window it measures. Wording is deliberately observational: a change
# that follows follow-ups is not proof that the follow-ups caused it.
FR_IMPROVED_AFTER = "Improved after follow-ups"
FR_IMPROVED = "Improved"
FR_IMPROVED_NO_TASKS = "Improved · no follow-ups"
FR_STEADY = "Steady"
FR_DECLINED_DESPITE = "Declined despite follow-ups"
FR_DECLINED_INCOMPLETE = "Declined · follow-ups incomplete"
FR_DECLINED_NO_TASKS = "Declined · no follow-ups raised"
FR_NA = "Not comparable"
FR_LEVEL = {FR_IMPROVED_AFTER: "ok", FR_IMPROVED: "ok", FR_IMPROVED_NO_TASKS: "ok",
            FR_STEADY: "info", FR_DECLINED_DESPITE: "medium", FR_DECLINED_INCOMPLETE: "high",
            FR_DECLINED_NO_TASKS: "medium", FR_NA: "muted"}


def followup_result(effort_pct, change, steady_band=1.0, effort_high=75.0):
    """Verdict for 'are completed follow-ups followed by a better result?'.
    effort_pct = Task Completion % of the follow-ups that PRECEDED the newer
    outcome (Daily: the previous day's; period reports: the period's own), None
    when no task was raised. change = newer outcome − older outcome (points),
    None when either side is not measured. |change| < steady_band = Steady."""
    if change is None:
        return FR_NA
    if abs(change) < steady_band:
        return FR_STEADY
    up = change > 0
    if effort_pct is None:
        return FR_IMPROVED_NO_TASKS if up else FR_DECLINED_NO_TASKS
    if up:
        return FR_IMPROVED_AFTER if effort_pct >= effort_high else FR_IMPROVED
    return FR_DECLINED_DESPITE if effort_pct >= effort_high else FR_DECLINED_INCOMPLETE


def measurement_window(gk, kind, start, end, cutoff_fn=None) -> dict:
    """What a group's Actual Performance % measures for one period — the SAME
    window group_outcome() uses — as {"text", "start", "end"} (start / end =
    datetime or date, None when not applicable). cutoff_fn(day) = the time of
    that day's Coordinator task list (OutcomeEngine.cutoff)."""
    if gk in ("attendance", "instructor"):
        if kind == "Daily":
            cut = cutoff_fn(end) if cutoff_fn else datetime.combine(end, time(23, 59, 59))
            ws, we = attendance_daily_window(end, cut)
            return {"text": f"sessions {ws:%d-%b %I:%M %p} – {we:%d-%b %I:%M %p}",
                    "start": ws, "end": we}
        if gk == "attendance":
            return {"text": f"sessions dated {start:%d-%b} – {end:%d-%b}", "start": start, "end": end}
        return {"text": f"each report day's session window, {start:%d-%b} – {end:%d-%b}",
                "start": start, "end": end}
    if gk == "assignment":
        if kind == "Daily":
            d = end - timedelta(days=1)
            return {"text": f"assignments due {d:%d-%b} (Assignment Daily {d:%d-%b})",
                    "start": d, "end": d}
        return {"text": f"assignments due {start:%d-%b} – {end:%d-%b}", "start": start, "end": end}
    return {"text": "", "start": None, "end": None}


# =============================================================================
#  SOURCES  (all Google access; tests pass a fake with the same methods)
# =============================================================================
class GoogleOutcomeSources:
    """Lazy, cached readers for the outcome sources. Every method returns None
    (never raises) when its source cannot be read; the measure is then reported
    as 'not checked'."""

    def __init__(self, TL, cache_dir=None):
        self.TL = TL
        self.AR = TL.AR
        self.cache_dir = cache_dir
        self._cache = {}

    def _once(self, key, fn):
        if key not in self._cache:
            try:
                self._cache[key] = fn()
            except (Exception, SystemExit) as exc:              # noqa: BLE001 (some readers sys.exit)
                log.warning("Outcome source '%s' unavailable (%s).", key, exc)
                self._cache[key] = None
        return self._cache[key]

    def _sheets_service(self):
        from utils import get_sheets_service
        return self._once("svc", lambda: get_sheets_service(self.AR.SERVICE_ACCOUNT_FILE))

    def attendance_data(self, end: date):
        """(sess, att, fb, tf) of every date up to `end` (load_all_data)."""
        def _load():
            sess, att, fb, tf, _susp = self.AR.load_all_data(self._sheets_service(),
                                                             date(2000, 1, 1), end)
            return sess, att, fb, tf
        return self._once(("att", end), _load)

    def submissions(self):
        import pyAssignmentSubmissionPerformanceReport as ASR
        return self._once("subs", lambda: ASR.load_submissions(self._sheets_service()))

    def assignment_report(self, job):
        """The generated Assignment Submission Performance report of `job`
        (Drive: <Assignment Submission Report>/<type folder>/<period folder>/
        <name>, the names the Assignment script itself uses). Returns
        {"found": True, **parse_assignment_summary} / {"found": False} when it
        has not been generated, or None when Drive could not be read. Read-only:
        no folder is ever created. Cached per file revision."""
        return self._once(("asgrep", job["kind"], job["start"], job["end"]),
                          lambda: self._read_assignment_report(job))

    def _read_assignment_report(self, job):
        import io as _io
        import json as _json
        import openpyxl
        import pyAssignmentSubmissionPerformanceReport as ASR
        drive = self._once("asg_drive", ASR._drive_client)
        if drive is None:
            raise RuntimeError("Drive client unavailable")

        def _child(parent, name, folder):
            q = (f"'{parent}' in parents and name='{ASR._q(name)}' and trashed=false"
                 + (" and mimeType='application/vnd.google-apps.folder'" if folder else ""))
            return ASR._call(lambda: drive.files().list(
                q=q, fields="files(id,name,mimeType,modifiedTime)", supportsAllDrives=True,
                includeItemsFromAllDrives=True).execute(), "Drive: find " + name).get("files", [])
        parent = ASR.PERFORMANCE_ROOT_FOLDER_ID
        for name in ASR.period_folders(job):
            hit = _child(parent, name, True)
            if not hit:
                return {"found": False}
            parent = hit[0]["id"]
        base = ASR.report_name(ASR.PERFORMANCE_BASENAME, job)
        files = [f for f in ASR._existing(drive, parent, base)
                 if f.get("mimeType") == ASR.SHEET_MIME
                 and re.fullmatch(re.escape(base) + r"(\s*-\s*Version\s*\d+)?", f.get("name", ""))]
        if not files:
            return {"found": False}
        f = ASR._pick_existing(files, base)
        meta = ASR._call(lambda: drive.files().get(fileId=f["id"], fields="id,modifiedTime",
                                                  supportsAllDrives=True).execute(),
                         "Drive: report metadata")
        cache = None
        if self.cache_dir:
            safe = re.sub(r"[^0-9A-Za-z]", "", meta.get("modifiedTime", ""))
            cache = os.path.join(self.cache_dir, f"assignment_{f['id']}_{safe}.json")
            if os.path.exists(cache):
                try:
                    with open(cache, "r", encoding="utf-8") as fh:
                        return _json.load(fh)
                except Exception:                                   # noqa: BLE001
                    pass
        data = ASR._call(lambda: drive.files().export(
            fileId=f["id"], mimeType="application/vnd.openxmlformats-officedocument."
                                     "spreadsheetml.sheet").execute(), "Drive: export " + base)
        parsed = parse_assignment_summary(openpyxl.load_workbook(_io.BytesIO(data), data_only=True))
        if parsed is None:
            raise RuntimeError(f"{base}: Summary KPI 'Submission %' not found")
        res = {"found": True, "name": f.get("name", base), **parsed}
        if cache:
            try:
                os.makedirs(self.cache_dir, exist_ok=True)
                with open(cache, "w", encoding="utf-8") as fh:
                    _json.dump(res, fh)
            except Exception:                                       # noqa: BLE001
                pass
        return res

    def admission_rows(self):
        def _load():
            import pyAdmissionFormalitiesReport as ADM
            svc = self._sheets_service()
            students = ADM.read_records(svc, ADM.STUDENT_SHEET_ID, ADM.STUDENTS_TAB)
            signers = ADM.read_records(svc, ADM.SIGNERS_SHEET_ID, ADM.SIGNERS_TAB)
            by_email, by_name = ADM.build_signer_indexes(signers)
            rows, _st = ADM.build_rows(students, by_email, by_name, current_only=False,
                                       start_date=date(2000, 1, 1))
            return rows
        return self._once("adm", _load)

    def wise_data(self):
        return self._once("wise", lambda: self.TL.load_wise_validation(return_sources=True))

    def _iv_sheets(self):
        return self._once("ivs", lambda: self.TL._iv_services()[0])

    def interview_index(self):
        s = self._iv_sheets()
        return None if s is None else self._once("ividx", lambda: self.TL._iv_consolidate_index(s))

    def interview_grid(self):
        s = self._iv_sheets()
        if s is None:
            return None

        def _load():
            meta = s.spreadsheets().get(spreadsheetId=self.TL.INTERVIEW_CONSOLIDATE_SHEET_ID,
                                        fields="sheets(properties(title))").execute()
            tab = meta["sheets"][0]["properties"]["title"]
            resp = s.spreadsheets().values().get(
                spreadsheetId=self.TL.INTERVIEW_CONSOLIDATE_SHEET_ID,
                range=f"'{tab}'!A1:ZZ").execute()
            return resp.get("values", [])
        return self._once("ivgrid", _load)


# =============================================================================
#  ENGINE  (periods, days, trend)
# =============================================================================
class OutcomeEngine:
    """Applies the calculators to report days and periods.

    sources   GoogleOutcomeSources (or a fake)
    snapshot  callable(day) -> [detail dicts] stored by that day's Daily report,
              or None when no stored check exists
    cutoffs   {day: datetime of that day's latest Coordinator task list}
    keyfn     instructor review -> task identity key (the performance report's)
    """

    def __init__(self, TL, sources, now: datetime, snapshot=None, cutoffs=None, keyfn=None,
                 ASR=None, ADM=None, WISE=None, data_end=None):
        self.TL, self.AR, self.src, self.now = TL, TL.AR, sources, now
        self.snapshot = snapshot or (lambda day: None)
        self.cutoffs = dict(cutoffs or {})
        self.keyfn = keyfn
        self.ASR, self.ADM, self.WISE = ASR, ADM, WISE
        self._day_details = {}
        self._instr = {}
        # ONE attendance load serves every day / period of the run
        self.data_end = max(data_end or now.date(), now.date())

    # ── helpers ─────────────────────────────────────────────────────────────
    def cutoff(self, day: date) -> datetime:
        c = self.cutoffs.get(day)
        if c is not None:
            return c
        end = datetime.combine(day, time(23, 59, 59))
        return min(end, self.now) if day == self.now.date() else end

    def _att(self):
        return self.src.attendance_data(self.data_end)

    def _is_live(self, day):
        return day == self.now.date()

    # ── per-day evening checks (admission + wise) ───────────────────────────
    def day_details(self, day: date, tasks_by_group: dict) -> dict:
        """{'admission': [...], 'wise': [...], 'source': 'live'|'stored'|'none'}.
        The report day itself is checked live; earlier days use the checks
        stored by that day's Daily report."""
        if day in self._day_details:
            return self._day_details[day]
        res = {"admission": [], "wise": [], "source": "none"}
        adm_tasks = tasks_by_group.get("admission", [])
        wise_tasks = tasks_by_group.get("wise", [])
        if self._is_live(day):
            res["source"] = "live"
            if adm_tasks:
                rows = self.src.admission_rows()
                ADM = self.ADM
                res["admission"] = admission_check(day, adm_tasks, rows, self.now,
                                                   ADM.norm_email, ADM.norm)
            if wise_tasks:
                issues = wise_morning_issues(wise_tasks)
                need_iv = any(i["section"] == "interview" for i in issues)
                need_wise = any(i["section"] != "interview" for i in issues)
                ev = wise_evening_state(self.TL, self.src.wise_data() if need_wise else None,
                                        self.src.interview_index() if need_iv else None,
                                        self.WISE)
                res["wise"] = wise_check(day, issues, ev, self.now, self.TL._norm_name)
        else:
            stored = self.snapshot(day)
            if stored is not None:
                res["source"] = "stored"
                res["admission"] = [d for d in stored if d["group"] == "admission"]
                res["wise"] = [d for d in stored if d["group"] == "wise"]
            else:
                # no stored evening check for that day: its items are 'not checked'
                res["admission"] = [detail(day, "admission", "Admission Formalities",
                                           t.get("label", ""), "Admission form not signed", "",
                                           "no evening check stored", RES_NOT_CHECKED, "", t["key"])
                                    for t in adm_tasks]
                res["wise"] = [detail(day, "wise", WISE_SECTION_LABEL.get(i["section"], ""),
                                      i["item"], i["issue"], i["morning"],
                                      "no evening check stored", RES_NOT_CHECKED, "", i["key"])
                               for i in wise_morning_issues(wise_tasks)]
        self._day_details[day] = res
        return res

    # ── instructor (any day: recomputed from the attendance data) ───────────
    def instructor_for_day(self, day: date, tasks: list):
        if day in self._instr:
            return self._instr[day]
        data = self._att()
        if data is None:
            r = (not_checked("instructor", "attendance data unreadable"), [])
        else:
            sess, att, fb, tf = data
            ws, we = attendance_daily_window(day, self.cutoff(day))
            reviews = self.TL.review_instructor_sessions(
                filter_window(sess, ws, we), filter_window(att, ws, we),
                filter_window(fb, ws, we), filter_window(tf, ws, we))
            r = instructor_day(day, tasks, reviews, self.now, self.keyfn)
        self._instr[day] = r
        return r

    # ── one group, one period ───────────────────────────────────────────────
    def group_outcome(self, gk, kind, start, end, tasks, task_days):
        """tasks = the period's ledger tasks; task_days = report days that had a
        Coordinator task list (any group)."""
        daily = kind == "Daily"
        if gk == "attendance":
            data = self._att()
            if data is None:
                return not_checked(gk, "attendance data unreadable")
            att = data[1]
            if daily:
                ws, we = attendance_daily_window(end, self.cutoff(end))
                o = attendance_outcome(self.AR, filter_window(att, ws, we), period=False)
                o["note"] = f"{ws:%d-%b %I:%M %p} – {we:%d-%b %I:%M %p}"
                return o
            return attendance_outcome(self.AR, filter_dates(att, start, end), period=True)
        if gk == "assignment":
            import coordinator_periods as CP
            job = assignment_report_job(CP, kind, start, end)
            label = f"{job['kind']} {job['label']}"
            rep = self.src.assignment_report(job)
            if rep and rep.get("found"):
                return assignment_from_report(rep, label)
            reason = (f"Assignment report {label} not generated yet" if rep is not None
                      else f"Assignment report {label} could not be read")
            subs = self.src.submissions()
            if subs is None:
                return not_checked(gk, reason + "; Submissions sheet unreadable")
            return assignment_outcome(self.ASR, subs, job["start"], job["end"], self.now,
                                      reason=reason)
        if gk == "interview":
            grid = self.src.interview_grid()
            if grid is None:
                return not_checked(gk, "Interview Consolidated Report unreadable")
            a, b = interview_window(kind, start, end, self.now.date())
            if getattr(self, "_iv_rows", None) is None:
                self._iv_rows = interview_rows(grid)
            return interview_outcome(self._iv_rows, a, b,
                                     note=f"interviews {a:%d-%b} – {b:%d-%b}")
        by_day = defaultdict(lambda: defaultdict(list))
        for t in tasks:
            by_day[t["day"]][t["group"]].append(t)
        if gk in ("admission", "wise"):
            dets, sources = [], defaultdict(int)
            for d in sorted(by_day):
                if not by_day[d].get(gk):
                    continue
                dd = self.day_details(d, by_day[d])
                dets += dd[gk]
                sources[dd["source"]] += 1
            note = ""
            if sources.get("none"):
                note = f"{sources['none']} day(s) without a stored evening check"
            return resolution_outcome(gk, dets, note=note)
        if gk == "instructor":
            num = den = 0
            reasons = defaultdict(int)
            missing = 0
            for d in sorted(set(task_days)):
                if not (start <= d <= end):
                    continue
                o, _det = self.instructor_for_day(d, by_day[d].get("instructor", []))
                if o["state"] == ST_NOT_CHECKED:
                    missing += 1
                    continue
                num += o["num"]
                den += o["den"]
                for k, v in o.get("reasons", {}).items():
                    reasons[k] += v
            if missing and not den:
                return not_checked(gk, "attendance data unreadable")
            br = [(INSTR_REASON_LABELS[c], reasons[c]) for c in INSTR_REASON_LABELS if reasons[c]]
            return outcome(gk, num, den, breakdown=br,
                           extra={"held": den, "flagged": den - num, "reasons": dict(reasons)})
        return None

    def period_outcomes(self, kind, start, end, tasks, task_days) -> dict:
        out = OrderedDict()
        for gk in GROUP_ORDER:
            try:
                out[gk] = self.group_outcome(gk, kind, start, end, tasks, task_days)
            except Exception as exc:                                # noqa: BLE001
                log.exception("Outcome %s %s–%s failed: %s", gk, start, end, exc)
                out[gk] = not_checked(gk, f"error: {exc}")
            if out[gk] is not None and gk in ("attendance", "assignment", "instructor"):
                # what this figure measures (shown with the previous-vs-current comparison)
                out[gk]["window"] = measurement_window(gk, kind, start, end, self.cutoff)
        return out

    def detail_rows(self, day: date, tasks: list) -> list:
        """Every checked item of the report day (stored by the Daily report)."""
        by_g = defaultdict(list)
        for t in tasks:
            if t["day"] == day:
                by_g[t["group"]].append(t)
        dd = self.day_details(day, by_g)
        _o, instr = self.instructor_for_day(day, by_g.get("instructor", []))
        return dd["admission"] + dd["wise"] + instr
