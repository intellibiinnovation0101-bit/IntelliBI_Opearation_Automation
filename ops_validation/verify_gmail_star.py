"""
Verification of Gmail "Starred" for the Coordinator report e-mails
(run from the project root:  python ops_validation\\verify_gmail_star.py).

No network: Gmail SMTP and Gmail IMAP are replaced by an in-memory mailbox that
behaves like Gmail — a sent message lands in "All Mail" (here under a LOCALISED
folder name, found by its \\All flag), can take a few look-ups to appear, and
the IMAP \\Flagged flag is Gmail's Starred. Checks:
  * common/gmail_star.py: exact Message-ID match, waiting for a late message,
    localised All Mail, subject fallback (newest only), every failure mode
    returns False without raising
  * both Coordinator scripts (REAL generate()): every report e-mail is starred,
    exactly that message and nothing else; subject / recipients / body are the
    same as without starring; STAR_EMAIL_IN_GMAIL = False and SEND_EMAIL = False
    do not touch IMAP; an IMAP failure never changes the send result
"""
import email
import imaplib
import io
import os
import re
import sys
import types
from datetime import date, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.modules.setdefault("_bootstrap", types.ModuleType("_bootstrap"))
ec = types.ModuleType("email_config")
ec.GMAIL_SENDER = "info@intellibiinnovationstechnologies.in"
ec.GMAIL_APP_PASS = "test-only"
sys.modules["email_config"] = ec                     # never the real credentials
for p in ("common", "ops_reports_action", "co-ordinator reports", "scripts"):
    sys.path.insert(0, os.path.join(ROOT, p))

import smtplib                                        # noqa: E402
import openpyxl                                       # noqa: E402
import gmail_star as GS                               # noqa: E402

FAIL = []


def check(label, got, want):
    ok = got == want
    print(f"[{'pass' if ok else 'FAIL'}] {label}: {got!r}" + ("" if ok else f"  (expected {want!r})"))
    if not ok:
        FAIL.append(label)


# =============================================================================
#  In-memory Gmail
# =============================================================================
class Mailbox:
    def __init__(self):
        self.reset()

    def reset(self, visible_after=0, login_error=None, net_errors=0):
        GS._GIVE_UP.clear()
        self.msgs = []                  # {uid, msgid, subject, to, raw, flags, seen_at}
        self.visible_after = visible_after
        self.login_error = login_error
        self.net_errors = net_errors
        self.sessions = 0
        self.selected = []
        self.searches = 0

    def deliver(self, raw, to):
        m = email.message_from_string(raw)
        mid = m.get("Message-ID") or f"<gmail-generated-{len(self.msgs)}@mail.gmail.com>"
        self.msgs.append({"uid": str(100 + len(self.msgs)).encode(), "msgid": mid,
                          "subject": str(email.header.make_header(email.header.decode_header(m["Subject"]))),
                          "to": list(to), "raw": raw, "flags": set(), "found_after": self.visible_after})


BOX = Mailbox()
SENT = []


class FakeSMTP:
    def __init__(self, host, port, timeout=None):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def login(self, user, pw):
        self.user = user

    def sendmail(self, sender, recipients, msg):
        SENT.append({"from": sender, "to": list(recipients), "msg": msg})
        BOX.deliver(msg, recipients)


class FakeIMAP:
    def __init__(self, host, port, timeout=None):
        BOX.sessions += 1
        if BOX.net_errors:
            BOX.net_errors -= 1
            raise OSError("network unreachable")
        self.host = host

    def login(self, user, pw):
        if BOX.login_error:
            raise imaplib.IMAP4.error(BOX.login_error)
        if (user, pw) != (ec.GMAIL_SENDER, ec.GMAIL_APP_PASS):
            raise imaplib.IMAP4.error("[AUTHENTICATIONFAILED] Invalid credentials")
        return "OK", [b"logged in"]

    def list(self):
        return "OK", [b'(\\HasNoChildren) "/" "INBOX"',
                      b'(\\HasNoChildren \\Sent) "/" "[Gmail]/Gesendet"',
                      b'(\\HasNoChildren \\All) "/" "[Gmail]/Alle Nachrichten"']

    def select(self, mailbox, readonly=False):
        BOX.selected.append((mailbox, readonly))
        return "OK", [str(len(BOX.msgs)).encode()]

    def noop(self):
        return "OK", [b""]

    def _visible(self):
        return [m for m in BOX.msgs if BOX.searches > m["found_after"]]

    def uid(self, cmd, *args):
        if cmd == "SEARCH":
            BOX.searches += 1
            vis = self._visible()
            if args[0] == "X-GM-RAW":
                q = args[1].strip('"').replace('\\"', '"')
                m = re.match(r"rfc822msgid:(\S+)$", q)
                if m:
                    hit = [x["uid"] for x in vis if x["msgid"].strip("<>") == m.group(1)]
                else:                                     # subject fallback
                    words = re.search(r"subject:\((.*)\)", q).group(1).split()
                    hit = [x["uid"] for x in vis if all(w in x["subject"] for w in words)]
            else:                                          # HEADER Message-ID
                want = args[3].strip('"')
                hit = [x["uid"] for x in vis if x["msgid"] == want]
            return "OK", [b" ".join(hit)]
        if cmd == "STORE":
            uid, op, flags = args
            for x in BOX.msgs:
                if x["uid"] == uid and op == "+FLAGS":
                    x["flags"].add(flags.strip("()"))
            return "OK", [b""]
        raise AssertionError(cmd)

    def logout(self):
        return "BYE", [b""]


smtplib.SMTP_SSL = FakeSMTP
imaplib.IMAP4_SSL = FakeIMAP
GS.FIND_DELAYS = (0, 0, 0, 0, 0)                     # no real waiting in the test
LOG = []


def log(s):
    LOG.append(s)
    print("      " + s)


def starred():
    return [x["subject"] for x in BOX.msgs if "\\Flagged" in x["flags"]]


# =============================================================================
print("\n== 1. common/gmail_star.py ==")
mid = GS.new_message_id(ec.GMAIL_SENDER)
check("Message-ID is unique and on the sender's domain",
      (bool(re.fullmatch(r"<[^<>@\s]+@intellibiinnovationstechnologies\.in>", mid)),
       mid != GS.new_message_id(ec.GMAIL_SENDER)), (True, True))


def plant(subject, msgid=None, visible_after=0):
    BOX.visible_after = visible_after
    raw = f"Subject: {subject}\r\nMessage-ID: {msgid}\r\n\r\nbody" if msgid else f"Subject: {subject}\r\n\r\nbody"
    BOX.deliver(raw, [ec.GMAIL_SENDER])


BOX.reset()
plant("Daily Coordinator Task Performance Report - 01-Oct-2026", "<older@x>")
plant("Daily Coordinator Task Performance Report - 01-Oct-2026", mid)
check("exact message starred (found by its Message-ID)",
      (GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, mid, log=log),
       [x["msgid"] for x in BOX.msgs if x["flags"]]), (True, [mid]))
check("… in All Mail, found by its \\All flag (localised name), read-write",
      BOX.selected[-1], ('"[Gmail]/Alle Nachrichten"', False))

BOX.reset()
m2 = GS.new_message_id(ec.GMAIL_SENDER)
plant("Weekly Batch Coordinator Report - x", m2, visible_after=3)
check("waits for a message that appears a few seconds later",
      (GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, m2, log=log), len(starred())), (True, 1))

BOX.reset()
plant("Daily Lead Report - 01-Oct-2026")               # Gmail replaced the Message-ID
plant("Daily Lead Report - 01-Oct-2026")
check("fallback by subject stars the NEWEST matching sent message only",
      (GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, "<lost@x>",
                            subject="Daily Lead Report - 01-Oct-2026", log=log),
       [x["uid"] for x in BOX.msgs if x["flags"]]), (True, [b"101"]))

BOX.reset()
check("message never appears → False, no exception",
      GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, "<nowhere@x>", log=log), False)
BOX.reset(login_error="[ALERT] IMAP access is disabled for your account")
LOG.clear()
check("IMAP disabled / login refused → False, with the fix in the message",
      (GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, mid, log=log),
       "Forwarding and POP/IMAP" in " ".join(LOG)), (False, True))
BOX.reset(net_errors=1)
plant("Daily Batch Coordinator Report - x", mid)
check("dropped connection → reconnects once and still stars",
      (GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, mid, log=log), BOX.sessions), (True, 2))
BOX.reset(net_errors=5)
check("network down → False after one retry, no exception",
      (GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, mid, log=log), BOX.sessions), (False, 2))
BOX.reset(net_errors=5)
GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, mid, log=log)
_n = BOX.sessions
check("after IMAP failed once in a run, later e-mails skip starring at once (no waiting)",
      (GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, mid, log=log), BOX.sessions - _n), (False, 0))
BOX.reset()
check("missing Message-ID → False, no IMAP session",
      (GS.star_sent_message(ec.GMAIL_SENDER, ec.GMAIL_APP_PASS, "", log=log), BOX.sessions), (False, 0))

# =============================================================================
print("\n== 2. Coordinator Task Performance report (REAL generate()) ==")
import pyBatchCoordinatorDailyAttendanceReport as BC  # noqa: E402
import pyCoordinatorTaskPerformanceReport as P        # noqa: E402

check("STAR_EMAIL_IN_GMAIL default True in both scripts", (P.STAR_EMAIL_IN_GMAIL, BC.STAR_EMAIL_IN_GMAIL),
      (True, True))
D = date(2026, 10, 5)


def adm(name):
    return {"Student Name": name, "Email ID": name.split()[0].lower() + "@x.com",
            "Phone Number": "+919000000001", "Batch Name": "DAAI1026", "Joined On": "2026-10-01",
            "Request Form Name": "", "Recipient Status": "Form Not Sent", "Request Status": "",
            "Sent Date": "", "Signed Date": "", "Expiry Date": ""}


wb = openpyxl.Workbook()
wb.remove(wb.active)
BC.build_admission_formalities(wb.create_sheet("Learner Admission Formalities"), [adm("Asha Rao")], D)
buf = io.BytesIO()
wb.save(buf)
parsed = P.parse_report_workbook(openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=True))
versions = [{"id": "v1", "name": "v1", "day": D, "version": 1,
             "created": datetime(2026, 10, 5, 9, 0), "modified_raw": "", "folder": "Daily 05-Oct-2026"}]
P._drive_service = lambda: None
P.discover_report_versions = lambda drive, ranges: versions
P.load_version = lambda drive, v: parsed
P._now_ist = lambda: datetime(2026, 10, 5, 19, 0)
P.BC.upload_report = lambda folders, filename, b, base: f"https://docs.google.com/spreadsheets/d/{base}"
P.LEGACY_OUTPUT_DIR = os.path.join(HERE, "__no_such_dir__")
P.GENERATE_AUTO = True
P.SEND_EMAIL = True


def run_perf(star=True):
    SENT.clear()
    BOX.reset()
    P.STAR_EMAIL_IN_GMAIL = star
    res, _e = P.generate()
    return res


res = run_perf()
check("every report e-mail (Monday = Daily + Weekly) sent and Starred",
      (len(SENT), starred()),
      (2, ["Daily Coordinator Task Performance Report - 05-Oct-2026",
           "Weekly Coordinator Task Performance Report - 28-Sep-2026 to 04-Oct-2026"]))
check("… each starred message is exactly the one sent (its own Message-ID)",
      [x["msgid"] == email.message_from_string(s["msg"])["Message-ID"]
       for x, s in zip(BOX.msgs, SENT)], [True, True])
_with = [(s["from"], s["to"], email.message_from_string(s["msg"])["Subject"],
          [p.get_payload(decode=True) for p in email.message_from_string(s["msg"]).walk()
           if p.get_content_maintype() == "text"]) for s in SENT]
res_off = run_perf(star=False)
_without = [(s["from"], s["to"], email.message_from_string(s["msg"])["Subject"],
             [p.get_payload(decode=True) for p in email.message_from_string(s["msg"]).walk()
              if p.get_content_maintype() == "text"]) for s in SENT]
_strip = lambda rows: [(f, t, sub, [re.sub(rb"Generated [^<]*", b"", b) for b in bodies])
                       for f, t, sub, bodies in rows]
check("sender, recipients, subject and body identical with and without starring",
      (len(_with), _strip(_with) == _strip(_without)), (2, True))
check("STAR_EMAIL_IN_GMAIL = False → no IMAP session, nothing starred", (BOX.sessions, starred()), (0, []))
P.STAR_EMAIL_IN_GMAIL = True
P.SEND_EMAIL = False
SENT.clear()
BOX.reset()
P.generate()
check("SEND_EMAIL = False → nothing sent, no IMAP session", (len(SENT), BOX.sessions), (0, 0))
P.SEND_EMAIL = True
SENT.clear()
BOX.reset(login_error="[AUTHENTICATIONFAILED] Invalid credentials")
P.STAR_EMAIL_IN_GMAIL = True
res_fail = P.generate()[0]
check("IMAP failure: e-mails still sent and the run still succeeds",
      (len(SENT), [bool(r.get("failed")) for r in res_fail], P.email_results(res_fail) is not False),
      (2, [False, False], True))

# =============================================================================
print("\n== 3. Batch Coordinator report (REAL generate()) ==")
import utils as _utils                                                    # noqa: E402
import pandas as pd                                                        # noqa: E402
_utils.get_sheets_service = lambda *a, **k: None
BC.AR.load_all_data = lambda *a, **k: tuple(pd.DataFrame() for _ in range(5))
BC.AR._ist_today = lambda: date(2026, 10, 5)
BC._generate_daily = lambda service, report_date, *a: {
    "link": "https://docs.google.com/spreadsheets/d/x", "report_date": report_date.isoformat(),
    "daily_rows": 1, "agg_groups": 1, "task_counts": {"Learner Admission Formalities": 1}}
BC._generate_period = lambda service, report_type, start, end, plabel, *a: {
    "link": "https://docs.google.com/spreadsheets/d/w", "report_type": report_type,
    "period": plabel, "start": start.isoformat(), "end": end.isoformat()}
BC.GENERATE_AUTO = True
BC.SEND_EMAIL = True
SENT.clear()
BOX.reset(visible_after=2)
r1 = BC.generate()
check("every Batch Coordinator report e-mail sent and Starred",
      (len(SENT), starred()),
      (2, ["Daily Batch Coordinator Report - 05-Oct-2026",
           "Weekly Batch Coordinator Report - 28 Sep – 04 Oct 2026"]))
BC.STAR_EMAIL_IN_GMAIL = False
SENT.clear()
BOX.reset()
BC.generate()
check("STAR_EMAIL_IN_GMAIL = False → sent, not starred, no IMAP", (len(SENT), starred(), BOX.sessions), (2, [], 0))

# only the report e-mails of these scripts are starred: the shared send()
# defaults to star=False for any other caller
import coordinator_email as CE                                             # noqa: E402
SENT.clear()
BOX.reset()
CE.send("Some other e-mail", "<p>x</p>", [ec.GMAIL_SENDER], sender=ec.GMAIL_SENDER)
check("any other use of the shared sender is NOT starred (star defaults to off)",
      (len(SENT), starred(), BOX.sessions), (1, [], 0))

print("\nALL CHECKS PASSED" if not FAIL else f"\n{len(FAIL)} CHECK(S) FAILED: {FAIL}")
sys.exit(1 if FAIL else 0)
