"""
Verification of the Effort -> Outcome (actual performance) part of the
Coordinator Task Performance report (run from the project root:
  python ops_validation\\verify_coordinator_outcomes.py).

Synthetic data only — no Google access, no real learner data. Checks:
  1. Attendance  : the shared AR.overall_att_kpi reproduces the Session Summary's
                   daily AND period formulas exactly (random frames)
  2. Instructor  : TL.review_instructor_sessions drives the Instructor Follow-Ups
                   tab (same flagged sessions); session-level % with no double
                   counting; reasons counted separately
  3. Admission   : evening re-check of the morning list (4 -> 1 pending = 75%),
                   signed / still pending / no longer applicable, not checked
  4. Wise        : issue level (per field / failed check / interview feedback),
                   parsed from the REAL Wise tab layout
  5. Interview   : Consolidated Report attended rule + Daily / Weekly / Monthly windows
  6. Assignment  : the Assignment report's own Total Submission %
  7. Engine + workbook: Daily live check stored on the PC (no report tab), read back by
                   the Weekly report (Weekly = sum of its Dailies), Dashboard /
                   e-mail figures reconcile, quadrants, headline, six questions
  8. Evening batch: source refresh before the report, cache cap
  9. Yesterday → Today (Attendance / Assignment / Instructor): labels by actual
                   dates, measurement windows, availability, Follow-up → Result,
                   Dashboard block + charts, trend data, e-mail cards
"""
import io
import os
import random
import sys
import types
from datetime import date, datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.modules["_bootstrap"] = types.ModuleType("_bootstrap")
paths = types.ModuleType("paths")
paths.CREDENTIALS_DIR = paths.CONFIG_DIR = paths.LOGS_DIR = paths.CACHE_DIR = HERE
sys.modules["paths"] = paths
ec = types.ModuleType("email_config"); ec.GMAIL_SENDER = "x@x"; ec.GMAIL_APP_PASS = "x"
sys.modules["email_config"] = ec
for p in ("common", "ops_reports_action", "co-ordinator reports"):
    sys.path.insert(0, os.path.join(ROOT, p))

import pandas as pd                                      # noqa: E402
import openpyxl                                          # noqa: E402
import pyCoordinatorTaskListReport as BC                 # noqa: E402
import pyCoordinatorTaskPerformanceReport as P           # noqa: E402
import coordinator_outcomes as CO                        # noqa: E402
import coordinator_email as CE                           # noqa: E402

AR = BC.AR
FAIL = []


def check(label, got, want):
    ok = got == want
    print(f"[{'pass' if ok else 'FAIL'}] {label}: {got!r}" + ("" if ok else f"  (expected {want!r})"))
    if not ok:
        FAIL.append(label)


def norm(v):
    return " ".join(str(v or "").split()).strip().lower()


# =============================================================================
print("\n== 1. Attendance: Overall Att % ==")
mism = []
for seed in range(200):
    r = random.Random(seed)
    rows = [{"session_id": r.choice(["s1", "s2", "s3", ""]), "student_id": r.choice(["a", "b", "c", "", "nan"]),
             "status": r.choice(["Present", "Present", "Absent"]), "_attn_applicable": r.random() < .8}
            for _ in range(r.randint(1, 30))]
    df = pd.DataFrame(rows)
    app = df[df["_attn_applicable"]]
    old_daily = round((app.status == "Present").sum() / len(app) * 100, 2) if len(app) else 0.0
    d = df.drop_duplicates(subset=["session_id", "student_id"])
    d = d[~d.student_id.astype(str).str.strip().isin(["", "nan", "None", "NaN"])]
    d = d[~d.session_id.astype(str).str.strip().isin(["", "nan", "None", "NaN"])]
    ap = d[d["_attn_applicable"]]
    old_period = round((ap.status == "Present").sum() / len(ap) * 100, 1) if len(ap) else 0.0
    if (AR.overall_att_kpi(df, False)["pct"], AR.overall_att_kpi(df, True)["pct"]) != (old_daily, old_period):
        mism.append(seed)
check("overall_att_kpi == the Session Summary's daily (2 dp) and period (1 dp) formulas", mism, [])
src = open(os.path.join(ROOT, "ops_reports_action", "pyAttendaceFeedbackReport.py"), encoding="utf-8").read()
check("both Session Summary builders call the shared function",
      (src.count("overall_att_kpi(att_f, period=False)"), src.count("overall_att_kpi(att_f, period=True)")),
      (1, 1))
o = CO.attendance_outcome(AR, pd.DataFrame([{"session_id": "s", "student_id": x, "status": st,
                                             "_attn_applicable": ap}
                                            for x, st, ap in (("a", "Present", True), ("b", "Absent", True),
                                                              ("c", "Present", True), ("d", "Present", False))]),
                          period=False)
check("attendance outcome: 2 present of 3 applicable = 66.67%", (o["num"], o["den"], o["pct"]), (2, 3, 66.67))

# =============================================================================
print("\n== 2. Instructor: session level ==")
F = "%Y-%m-%d %H:%M:%S"
D = date(2026, 10, 6)


def mk_session(sid, start, late=0, dur=120, kind="held", course="SQL", tutor="Instructor A", day=None):
    st = datetime.combine((day or D) - timedelta(days=1), datetime.min.time()) + timedelta(hours=start)
    act = st + timedelta(minutes=late)
    return {"session_id": sid, "course_name": course, "course_title": "01-Oct-2026 To Current Date",
            "tutor_name": tutor, "start_time_ist": (act if kind != "sched" else st).strftime(F),
            "end_time_ist": (act + timedelta(minutes=dur)).strftime(F) if kind == "held" else "",
            "Session Scheduled Start": st.strftime(F),
            "Session Scheduled End": (st + timedelta(minutes=120)).strftime(F),
            "_date": act.date(), "_dt": act}


def day_sessions(day, tag=""):
    """10 held sessions on the afternoon/evening before `day` (inside its Daily
    window): S0 late + underrun, S1 instructor feedback missing; + 1 future one."""
    sess, att, tfb, fbk = [], [], [], []
    for i in range(10):
        s = mk_session(f"{tag}S{i}", 13 + i, late=12 if i == 0 else -10, dur=60 if i == 0 else 120,
                       day=day)
        sess.append(s)
        for k in range(4):
            att.append({"session_id": s["session_id"], "student_id": f"st{k}", "status": "Present",
                        "_attn_applicable": True, "_dur_min": 60.0,
                        "session_start_ist": s["start_time_ist"], "session_end_ist": s["end_time_ist"]})
            fbk.append({"session_id": s["session_id"], "student_id": f"st{k}", "_rating": 9.0})
        if i != 1:
            tfb.append({"session_id": s["session_id"], "topics_covered": "Topic", "comments": ""})
    sch = mk_session(f"{tag}S10", 20, kind="sched", day=day)     # not yet held: never counted
    sch.update({"start_time_ist": "2099-01-05 10:00:00", "Session Scheduled Start": "2099-01-05 10:00:00",
                "Session Scheduled End": "2099-01-05 12:00:00"})
    sess.append(sch)
    return pd.DataFrame(sess), pd.DataFrame(att), pd.DataFrame(fbk), pd.DataFrame(tfb)


SESS, ATT, fb, TF = day_sessions(D)
reviews = BC.review_instructor_sessions(SESS, ATT, fb, TF)
flagged = [r for r in reviews if r["reasons"]]
check("review: S0 (late + underrun) and S1 (feedback missing) flagged; the future session is not",
      sorted((r["sid"], tuple(r["codes"])) for r in flagged),
      [("S0", ("started_late", "underrun")), ("S1", ("feedback_missing",))])
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Instructor Follow-Ups"
BC.build_instructor_followups(ws, SESS, ATT, fb, TF, {}, D)
parsed = P.parse_report_workbook(wb)
itasks = parsed["instructor"]["tasks"]
check("the tab lists exactly the reviewed follow-ups", len(itasks), len(flagged))
for t in itasks:
    t["day"], t["key"] = D, t["key"]
o, det = CO.instructor_day(D, itasks, reviews, datetime(2026, 10, 6, 19, 0))
check("10 sessions held, 2 escalated -> 80% (one session counted once)",
      (o["den"], o["num"], o["pct"]), (10, 8, 80.0))
check("reasons shown separately (3 reasons on 2 sessions)",
      sorted(o["breakdown"]), sorted([("Started late", 1), ("Session underrun (>= 30 min short)", 1),
                                      ("Instructor feedback missing", 1)]))
_keyfn = lambda r: "|".join(P._key_value(f, v) for f, v in zip(
    ["Tech Name", "Duration", "Instructor", "Session Date"],
    [r["sess"]["course_name"], r["sess"]["course_title"], r["instr"],
     r["actual_start_dt"].strftime(F) if r["actual_start_dt"] is not None else ""]))
_o2, det2 = CO.instructor_day(D, itasks, reviews, datetime(2026, 10, 6, 19, 0), _keyfn)
check("detail: 2 escalated + 8 clean sessions (tab key == review key)",
      (sum(d["result"] == CO.RES_ESCALATED for d in det2), sum(d["result"] == CO.RES_CLEAN for d in det2)),
      (2, 8))

# =============================================================================
print("\n== 3. Admission: evening re-check ==")


def adm(name, email, status="Form Not Sent"):
    return {"Student Name": name, "Email ID": email, "Phone Number": "+919000000001",
            "Batch Name": "DAAI1026", "Joined On": "2026-10-01", "Request Form Name": "",
            "Recipient Status": status, "Request Status": "", "Sent Date": "", "Signed Date": "",
            "Expiry Date": ""}


morning = [adm("Learner A", "a@x.com"), adm("Learner B", "b@x.com", "Sent"),
           adm("Learner C", "c@x.com"), adm("Learner D", "d@x.com")]
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Learner Admission Formalities"
BC.build_admission_formalities(ws, morning, D)
atasks = P.parse_report_workbook(wb)["admission"]["tasks"]
for t in atasks:
    t["day"] = D
check("morning list parsed: 4 learners with their cells",
      sorted(t["cells"]["Student Name"] for t in atasks), ["Learner A", "Learner B", "Learner C", "Learner D"])
evening = [dict(adm("Learner A", "A@X.com"), **{"Recipient Status": "Signed", "Signed Date": "2026-10-06"}),
           dict(adm("Learner B", "b@x.com"), **{"Recipient Status": "Signed"}),
           dict(adm("Learner C", ""), **{"Recipient Status": "Signed"}),       # matched by name
           adm("Learner D", "d@x.com", "Sent")]
det = CO.admission_check(D, atasks, evening, datetime(2026, 10, 6, 19), lambda v: norm(v), norm)
o = CO.resolution_outcome("admission", det)
check("4 pending in the morning -> 1 pending in the evening = 75% (as specified)",
      (o["num"], o["den"], o["pct"]), (3, 4, 75.0))
check("… resolved = SIGNED, not 'Done? = Yes'", sorted(d["result"] for d in det),
      [CO.RES_RESOLVED] * 3 + [CO.RES_PENDING])
det = CO.admission_check(D, atasks, evening[:3], datetime(2026, 10, 6, 19), lambda v: norm(v), norm)
o = CO.resolution_outcome("admission", det)
check("learner no longer active -> 'No longer applicable', out of the denominator",
      (o["num"], o["den"], dict(o["breakdown"])[CO.RES_NA]), (3, 3, 1))
det = CO.admission_check(D, atasks, None, datetime(2026, 10, 6, 19), lambda v: norm(v), norm)
o = CO.resolution_outcome("admission", det)
check("sources unreadable -> not checked (never counted as pending)", (o["state"], o["pct"]),
      (CO.ST_NOT_CHECKED, None))

# =============================================================================
print("\n== 4. Wise & Interview Feedback: issue level ==")


def stu(name, joined="2026-10-01T05:00:01Z", **st):
    rec = {"Student Name": name, "Batch Name": "DAAI1026", "Joined On": joined}
    for f in ("Student Name", "Email ID", "Phone Number", "Tag Name", "Private Note", "Profile Picture"):
        rec[f + " Status"] = st.get(f.split()[0].lower(), "Valid")
    return rec


def crs(course_title, **st):
    rec = {"Course Title": course_title, "Course Subtitle": "Sub", "Instructor_Name": "IA", "Created On": "2026-09-01"}
    for f in ("Course Title", "Course Subtitle", "Course Tag Name"):
        rec[f + " Status"] = st.get(f.split()[1].lower(), "Valid")
    return rec


wdata = {"student": [stu("Learner E", email="Missing", phone="Invalid"),
                     stu("Learner F", "2026-10-01T05:00:02Z", profile="Missing"),
                     stu("Learner G", "2026-10-01T05:00:03Z", tag="Missing"),
                     stu("learner h", "2026-10-01T05:00:04Z", student="Invalid")],
         "course": [crs("Course One", title="Invalid")],
         "instructor": [{"Instructor ID": "I1", "Instructor Name": "Instructor B", "Validation Type": t,
                         "Detailed Validation Message": "x"} for t in ("Missing phone", "Invalid email",
                                                                       "Missing photo")]}
ivrows = [{"cells": ["03-Oct-2026", "Interviewer A", "SQL", "Batch IV", "Title IV"], "status_cells": {},
           "severity": 2, "why_reasons": [[("Interview feedback missing", False)]]}]
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Wise & Interview Feedback Validation"
BC.build_wise_validation(ws, wdata, D, interview_rows=ivrows)
wtasks = P.parse_report_workbook(wb)["wise"]["tasks"]
for t in wtasks:
    t["day"] = D
issues = CO.wise_morning_issues(wtasks)
check("issues: E=2, F=1, G=1, H=1, course=1, instructor=3 checks, interview=1",
      sorted((i["section"], i["issue"]) for i in issues),
      sorted([("student", "Email: Missing"), ("student", "Phone: Invalid"), ("student", "Picture: Missing"),
              ("student", "Name: Invalid"),
              ("student", "Tag: Missing"), ("course", "Title: Invalid"),
              ("instructor", "Failed check 1 of 3"), ("instructor", "Failed check 2 of 3"),
              ("instructor", "Failed check 3 of 3"), ("interview", "Interview feedback not completed")]))
eve_data = {"student": [stu("Learner E", phone="Invalid")],          # email fixed, phone still invalid
            "course": [],                                             # course fixed (or gone: see sources)
            "instructor": [wdata["instructor"][0]],                   # 1 of 3 checks still failing
            "_sources": {"student": [{"student_name": "Learner E", "batch_name": "DAAI1026",
                                      "joined_on": "2026-10-01T05:00:01Z"},
                                     {"student_name": "Learner F", "batch_name": "DAAI1026",
                                      "joined_on": "2026-10-01T05:00:02Z"},
                                     # 'learner h' renamed to 'Learner H' (same Joined On) -> fixed
                                     {"student_name": "Learner H", "batch_name": "DAAI1026",
                                      "joined_on": "2026-10-01T05:00:04Z"}],
                         # Learner G no longer in the student source -> not applicable
                         "combined": [{"class_name": "Course One", "class_subject": "Sub"}],
                         "instructor": [{"instructor_id": "I1", "Is_Active": "Y"}]}}
iv_index = {(BC._norm_name("Batch IV"), BC._norm_name("Title IV"), date(2026, 10, 4))}
ev = CO.wise_evening_state(BC, eve_data, iv_index)
det = CO.wise_check(D, issues, ev, datetime(2026, 10, 6, 19), BC._norm_name)
res = {(d["section"], d["issue"]): d["result"] for d in det}
check("per-issue results", res, {
    ("Student", "Email: Missing"): CO.RES_RESOLVED, ("Student", "Phone: Invalid"): CO.RES_PENDING,
    ("Student", "Picture: Missing"): CO.RES_RESOLVED, ("Student", "Tag: Missing"): CO.RES_NA,
    ("Student", "Name: Invalid"): CO.RES_RESOLVED,
    ("Course", "Title: Invalid"): CO.RES_RESOLVED,
    ("Instructor", "Failed check 1 of 3"): CO.RES_PENDING,
    ("Instructor", "Failed check 2 of 3"): CO.RES_RESOLVED,
    ("Instructor", "Failed check 3 of 3"): CO.RES_RESOLVED,
    ("Interview feedback", "Interview feedback not completed"): CO.RES_RESOLVED})
o = CO.resolution_outcome("wise", det)
check("Resolved Issues ÷ Initial Applicable Issues = 7 ÷ 9 = 77.8% (deleted student excluded)",
      (o["num"], o["den"], o["pct"]), (7, 9, 77.8))
det = CO.wise_check(D, issues, CO.wise_evening_state(BC, eve_data, None), datetime(2026, 10, 6, 19))
check("Consolidate Sheet unreadable -> that issue 'not checked', the rest still measured",
      [d["result"] for d in det if d["section"] == "Interview feedback"], [CO.RES_NOT_CHECKED])
det = CO.wise_check(D, issues, CO.wise_evening_state(BC, None, iv_index), datetime(2026, 10, 6, 19))
check("Wise unavailable -> Wise issues not checked (never pending)",
      sorted({d["result"] for d in det if d["section"] != "Interview feedback"}), [CO.RES_NOT_CHECKED])

# =============================================================================
print("\n== 5. Interview attendance ==")
grid = [["IntelliBI Interview Consolidated Report"],
        ["Batch ID", "Interview Date", "Interview Time", "Interviewer", "Tech Stack", "Batch Name",
         "Batch Title", "Slot #", "Candidate Name", "Email", "Phone", "Communication", "Tech Scores",
         "Total Score", "Average Score", "Zone", "Comments", "Instructor Name", "Co-Ordinator Name",
         "Published On", "Start Time"]]


def iv(day, zone="Good", avg="7", pub=""):
    return ["B1", day, "10:00", "Interviewer A", "SQL", "Batch", "Title", "1", "Candidate", "c@x.com",
            "+910000000000", "", "", "", avg, zone, "", "", "", pub, ""]


grid += [iv("05-Oct-2026"), iv("05-Oct-2026", zone="Absent"), iv("06-Oct-2026", avg="", pub=""),
         iv("06-Oct-2026", avg="", pub="06-Oct-2026"), iv("01-Oct-2026"), iv("29-Sep-2026", zone="Skipped"),
         ["", "", "", "", "", "", "", "", "", "", ""]]
rows = CO.interview_rows(grid)
check("attended rule: not absent/skipped AND (scored OR published)",
      [r["attended"] for r in rows], [True, False, False, True, True, False])
TODAY = date(2026, 10, 6)
win = {k: CO.interview_window(k, s, e, TODAY) for k, s, e in
       (("Daily", D, D), ("Weekly", date(2026, 10, 5), date(2026, 10, 11)),
        ("Monthly", date(2026, 10, 1), date(2026, 10, 31)))}
check("windows: Daily D-1..D, Weekly Monday..today, Monthly 1st..today", win,
      {"Daily": (date(2026, 10, 5), D), "Weekly": (date(2026, 10, 5), TODAY),
       "Monthly": (date(2026, 10, 1), TODAY)})
check("Daily 05..06-Oct: 2 of 4 interviewed = 50%",
      (lambda o: (o["num"], o["den"], o["pct"]))(CO.interview_outcome(rows, *win["Daily"])), (2, 4, 50.0))
check("Monthly 01..06-Oct: 3 of 5 = 60% (29-Sep excluded)",
      (lambda o: (o["num"], o["den"], o["pct"]))(CO.interview_outcome(rows, *win["Monthly"])), (3, 5, 60.0))

# =============================================================================
print("\n== 6. Assignment: Total Submission % ==")
import pyAssignmentSubmissionPerformanceReport as ASR    # noqa: E402


def sub(aid, deadline, sid, status):
    return {"assessment_id": aid, "assessment_title": aid, "class_id": "c" + aid, "class_name": "SQL",
            "class_subject": "B", "submission_start_date": "01/10/2026 10:00:00 IST",
            "submission_deadline": deadline, "student_id": sid, "student_name": sid,
            "student_email": f"{sid}@x.com", "student_phone": "", "submitted_at": "",
            "submission_status": status, "evaluation_marks": "", "maximum_marks": "10"}


SUBS = pd.DataFrame([sub("A1", "06/10/2026 23:45:00 IST", "s1", "Submitted"),
                     sub("A1", "06/10/2026 23:45:00 IST", "s2", "Not Submitted"),
                     sub("A2", "05/10/2026 12:00:00 IST", "s1", "Submitted"),
                     sub("A2", "05/10/2026 12:00:00 IST", "s3", "Submitted"),
                     sub("A3", "08/10/2026 23:45:00 IST", "s1", "Not Submitted")])
NOW = datetime(2026, 10, 7, 19, 0)
import coordinator_periods as CPm                       # noqa: E402
check("corresponding Assignment report: Daily D -> Daily D-1 (made on D); Weekly / Monthly -> same period",
      [(j["kind"], j["start"], j["end"]) for j in (
          CO.assignment_report_job(CPm, "Daily", date(2026, 10, 7), date(2026, 10, 7)),
          CO.assignment_report_job(CPm, "Weekly", date(2026, 9, 28), date(2026, 10, 4)),
          CO.assignment_report_job(CPm, "Monthly", date(2026, 9, 1), date(2026, 9, 30)))],
      [("Daily", date(2026, 10, 6), date(2026, 10, 6)), ("Weekly", date(2026, 9, 28), date(2026, 10, 4)),
       ("Monthly", date(2026, 9, 1), date(2026, 9, 30))])
# the Summary tab exactly as the Assignment script writes it (its own builder)
_job = CPm.make_job("Daily", D, D)
_ds = ASR.build_dataset(SUBS, D, D, NOW)
_src = ASR.summarise(_ds)
_rep = CO.parse_assignment_summary(ASR.build_performance_workbook(_job, _ds)[0])
check("Summary KPIs read by label from the Assignment report's own layout",
      (_rep["pct"], _rep["submitted"], _rep["expected"], _rep["assignments"]),
      (_src["rate"], _src["submitted"], _src["expected"], _src["assignments"]))
o = CO.assignment_from_report(dict(_rep, found=True), "Daily 06-Oct-2026")
check("outcome = the report's own Total Submission % (no recalculation)",
      (o["pct"], o["num"], o["den"], o["note"].startswith("from Assignment report Daily 06-Oct-2026")),
      (_src["rate"], _src["submitted"], _src["expected"], True))
_empty = CO.parse_assignment_summary(ASR.build_performance_workbook(
    CPm.make_job("Daily", date(2026, 10, 1), date(2026, 10, 1)),
    ASR.build_dataset(SUBS, date(2026, 10, 1), date(2026, 10, 1), NOW))[0])
check("report with no passed deadline ('—') -> nothing to measure, said clearly",
      (lambda o: (o["state"], o["pct"], "no assignment deadline" in o["note"]))(
          CO.assignment_from_report(dict(_empty, found=True), "Daily 01-Oct-2026")),
      (CO.ST_NO_DATA, None, True))
o = CO.assignment_outcome(ASR, SUBS, date(2026, 10, 5), date(2026, 10, 11), NOW,
                          reason="Assignment report Weekly … not generated yet")
s_src = ASR.summarise(ASR.build_dataset(SUBS, date(2026, 10, 5), date(2026, 10, 11), NOW))
check("not generated yet -> PROVISIONAL figure with the Assignment report's own functions",
      (o["num"], o["den"], o["pct"], o["note"].startswith("PROVISIONAL — Assignment report")),
      (s_src["submitted"], s_src["expected"], s_src["rate"], True))


class _AsgSrc:
    def __init__(self, rep):
        self.rep = rep
        self.asked = []

    def assignment_report(self, job):
        self.asked.append((job["kind"], job["start"]))
        return self.rep

    def submissions(self):
        return SUBS


for rep, want in ((dict(_rep, found=True), ("report", _src["rate"])),
                  ({"found": False}, ("provisional", None)),
                  (None, ("provisional", None))):
    src = _AsgSrc(rep)
    eng = CO.OutcomeEngine(BC, src, NOW, ASR=ASR)
    o = eng.group_outcome("assignment", "Daily", date(2026, 10, 7), date(2026, 10, 7), [], [])
    got = (o.get("source"), o["pct"] if o.get("source") == "report" else None)
    tag = "found" if rep and rep.get("found") else ("missing" if rep is not None else "unreadable")
    check(f"engine, Daily 07-Oct, Assignment report {tag} -> {want[0]} (asks for Daily 06-Oct)",
          (got, src.asked), (want, [("Daily", date(2026, 10, 6))]))
    if want[0] == "provisional":
        check("… and says why", ("not generated yet" if rep is not None else "could not be read") in o["note"],
              True)

# =============================================================================
print("\n== 7. Engine, stored checks, workbook, e-mail ==")
D0, D1 = date(2026, 10, 5), date(2026, 10, 6)      # Monday, Tuesday


def day_report(day):
    """A Coordinator day report built with the real builders (admission + wise + instructor)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    BC.build_admission_formalities(wb.create_sheet("Learner Admission Formalities"), morning, day)
    BC.build_wise_validation(wb.create_sheet("Wise & Interview Feedback Validation"), wdata, day,
                             interview_rows=ivrows)
    s, a, f, t = DAYDATA[day]
    BC.build_instructor_followups(wb.create_sheet("Instructor Follow-Ups"), s, a, f, t, {}, day)
    return P.parse_report_workbook(wb)


DAYDATA = {d: day_sessions(d, tag=f"{d.day}") for d in (date(2026, 10, 5), date(2026, 10, 6))}


versions = [{"id": f"v{d.day}", "name": f"r{d.day}", "day": d, "version": 1,
             "created": datetime.combine(d, datetime.min.time()) + timedelta(hours=10, minutes=30),
             "modified_raw": ""} for d in (D0, D1)]
PARSED = {v["id"]: day_report(v["day"]) for v in versions}


class FakeSources:
    def __init__(self):
        self.calls = []

    def attendance_data(self, end):
        self.calls.append("att")
        parts = list(zip(*DAYDATA.values()))
        s, a, f2, t2 = (pd.concat(x, ignore_index=True) for x in parts)
        smap = s.set_index("session_id")["_dt"].to_dict()
        for df in (a, f2, t2):
            df["_dt"] = df["session_id"].map(smap)
            df["_date"] = df["_dt"].apply(lambda x: x.date() if pd.notna(x) else None)
        return s, a, f2, t2

    def submissions(self):
        return SUBS

    def assignment_report(self, job):
        return {"found": False}                      # -> provisional, labelled

    def admission_rows(self):
        self.calls.append("adm")
        return evening

    def wise_data(self):
        self.calls.append("wise")
        return eve_data

    def interview_index(self):
        return iv_index

    def interview_grid(self):
        return grid


ADMm = types.SimpleNamespace(norm_email=lambda v: norm(v), norm=norm)
STORE = {}


def engine(now, snapshot=None):
    e = CO.OutcomeEngine(BC, FakeSources(), now, snapshot=snapshot or (lambda d: STORE.get(d)),
                         keyfn=_keyfn, ASR=ASR, ADM=ADMm, WISE=None)
    return e


P.TASK_GROUPS = P.TASK_GROUPS                     # unchanged registry
loader = lambda v: PARSED[v["id"]]
# Monday evening: Daily of 05-Oct, checked live
now0 = datetime(2026, 10, 5, 19, 0)
e0 = engine(now0)
jobs0 = [{"kind": "Daily", "start": D0, "end": D0, "label": "05-Oct-2026"}]
UP = {}
import tempfile                                          # noqa: E402
P.OUTCOME_CHECKS_DIR = tempfile.mkdtemp(prefix="outcome_checks_")   # never the real store


def upload(folders, name, buf):
    UP[name] = buf.getvalue()
    return f"https://docs.google.com/spreadsheets/d/{name}"


r0 = P.run_jobs(jobs0, versions[:1], loader, now0, upload=upload, engine=e0)[0]
check("Daily job delivered with an Effort -> Outcome view", (r0.get("failed"), bool(r0.get("outcome_rows"))),
      (None, True))
wb0 = openpyxl.load_workbook(io.BytesIO(UP[r0["name"]]))
check("Daily tabs (no Outcome Detail tab)", wb0.sheetnames,
      ["Dashboard", "Effort vs Outcome Trend", "Progress Trend", "Task Register", "Data Coverage & Rules"])
check("every tab's header ends with '  |  Generated On: <run time, IST>' (merge kept)",
      [(ws.title, str(ws["A1"].value).endswith("  |  Generated On: 05-Oct-2026 07:00 PM"),
        any(m.min_row == 1 and m.min_col == 1 and m.max_col > 1 for m in ws.merged_cells.ranges))
       for ws in wb0.worksheets],
      [(ws.title, True, True) for ws in wb0.worksheets])
stored = P.load_outcome_checks(D0)
check("evening checks stored on the PC and read back (admission 4 + wise 10 + instructor 10 rows)",
      (sum(d["group"] == "admission" for d in stored), sum(d["group"] == "wise" for d in stored),
       sum(d["group"] == "instructor" for d in stored)), (4, 10, 10))
STORE[D0] = [d for d in stored]
rows0 = {r["gk"]: r for r in r0["outcome_rows"]}
check("Daily actuals: admission 75%, wise 77.8%, instructor 80%",
      (rows0["admission"]["actual"], rows0["wise"]["actual"], rows0["instructor"]["actual"]),
      (75.0, 77.8, 80.0))
dash = wb0["Dashboard"]
sc = {r[0]: r for r in dash.iter_rows(values_only=True) if r and r[0] in
      [g["name"] for g in P.TASK_GROUPS]}
_card = {k: r for k, r in rows0.items() if k not in P.COMPARE_GROUPS}
check("Dashboard scorecard: the Yesterday → Today groups are not repeated; the others stay, in order",
      [n for n in sc], [g["name"] for g in P.TASK_GROUPS if g["key"] not in P.COMPARE_GROUPS])
check("Dashboard Actual Performance % == the outcome figures",
      {k: sc[r["name"]][7] for k, r in _card.items()},
      {k: (round(r["actual"], 1) if r["actual"] is not None else "—") for k, r in _card.items()})
check("Dashboard Target column == OUTCOME_TARGETS",
      {k: sc[r["name"]][8] for k, r in _card.items()},
      {k: v for k, v in P.OUTCOME_TARGETS.items() if k not in P.COMPARE_GROUPS})
_tot0 = [r for r in dash.iter_rows(values_only=True) if r and r[0] == "All task groups"][0]
check("… the total row still covers EVERY group (13 tasks, 0 done) and says so",
      (_tot0[1], _tot0[2], _tot0[10]),
      (13, 0, "Incl. the 3 follow-up groups below · Actual % = simple average"))
cells = [str(c.value) for row in dash.iter_rows() for c in row if c.value is not None]
check("headline: 'Coordinator completed X% of required actions — actual result …'",
      any(c.startswith("Coordinator completed 0.0% of required actions (0 of 13) — actual result:")
          for c in cells), True)
check("quadrant: no follow-up done but admission on 75% target -> 'Check if tasks needed'",
      rows0["admission"]["quadrant"], CO.QUAD_CHECK_TASKS)
check("quadrant: instructor 80% < 90% target with 0% effort -> 'Needs management attention'",
      rows0["instructor"]["quadrant"], CO.QUAD_ATTENTION)
check("attention list names the instructor area first",
      r0["attention"][0][:2], ("Instructor", "high"))
html = CE.performance_html("Daily", "05-Oct-2026", r0["summary"], r0["groups"], "https://x", "—",
                           benchmark=95.0, group_labels=P.EMAIL_GROUP_LABELS,
                           outcome_rows=r0["outcome_rows"], headline=r0["headline"],
                           attention=r0["attention"])
g = html.split("Performance vs Goals", 1)[1]
check("e-mail: Overall Completion % first, then effort + outcome bars per area",
      (g.index("Overall Completion %") < g.index("Effort &middot; Task Completion %")
       < g.index("Outcome &middot; Formalities Resolved %")), True)
check("e-mail outcome figures == Dashboard", all(f"{r['actual']:.1f}%" in g
                                                  for r in r0["outcome_rows"] if r["actual"] is not None), True)
check("e-mail carries the headline and the attention list",
      (r0["headline"] in html.replace("&mdash;", "—").replace("&#x27;", "'"), "Needs Management Attention" in html),
      (True, True))
check("e-mail labels: Wise is no longer shown as a second 'Learner Admission Formalities'",
      P.EMAIL_GROUP_LABELS["Wise & Interview Feedback Validation"], "Wise & Interview Feedback")

# Tuesday evening: Daily 06-Oct live + Weekly-to-date (Manual Mon..Tue) reading Monday's stored check
now1 = datetime(2026, 10, 6, 19, 0)
evening = evening[:1] + [adm("Learner B", "b@x.com", "Sent")] + evening[2:]   # B unsigned again on Tuesday
e1 = engine(now1)
jobs1 = [{"kind": "Daily", "start": D1, "end": D1, "label": "06-Oct-2026"},
         {"kind": "Manual", "start": D0, "end": D1, "label": "05-Oct-2026 to 06-Oct-2026"}]
r1, rw = P.run_jobs(jobs1, versions, loader, now1, upload=upload, engine=e1)
a1 = {r["gk"]: r["o"] for r in r1["outcome_rows"]}
aw = {r["gk"]: r["o"] for r in rw["outcome_rows"]}
check("Tuesday Daily admission: 2 of 4 = 50%", (a1["admission"]["num"], a1["admission"]["den"]), (2, 4))
check("period = sum of its Dailies (Monday from the STORED check, Tuesday live): 5 of 8",
      (aw["admission"]["num"], aw["admission"]["den"], aw["admission"]["pct"]), (5, 8, 62.5))
check("… wise 14 of 18, instructor 16 of 20",
      ((aw["wise"]["num"], aw["wise"]["den"]), (aw["instructor"]["num"], aw["instructor"]["den"])),
      ((14, 18), (16, 20)))
check("Δ vs previous day on the Daily (admission 50 - 75 = -25 pts)",
      {r["gk"]: r["d_actual"] for r in r1["outcome_rows"]}["admission"], -25.0)
e2 = engine(now1, snapshot=lambda d: None)
rw2 = P.run_jobs(jobs1[1:], versions, loader, now1, upload=upload, engine=e2)[0]
aw2 = {r["gk"]: r["o"] for r in rw2["outcome_rows"]}
check("a day without a stored check is NOT guessed: Monday's items 'not checked', Tuesday measured",
      (aw2["admission"]["den"], dict(aw2["admission"]["breakdown"])[CO.RES_NOT_CHECKED]), (4, 4))
wbw = openpyxl.load_workbook(io.BytesIO(UP[rw["name"]]))
check("period report: no Outcome Detail tab", "Outcome Detail" in wbw.sheetnames, False)
check("a past day is never overwritten by a re-run (only LIVE checks are stored)",
      sorted(os.listdir(P.OUTCOME_CHECKS_DIR)), ["2026-10-05.json", "2026-10-06.json"])
_dry = P.run_jobs(jobs0, versions[:1], loader, now0, upload=None, engine=engine(now0))[0]
check("dry run (no upload) stores nothing new", sorted(os.listdir(P.OUTCOME_CHECKS_DIR)),
      ["2026-10-05.json", "2026-10-06.json"])
tr = wbw["Effort vs Outcome Trend"]
texts = [str(c.value) for row in tr.iter_rows() for c in row if c.value is not None]
_secs = [t.strip() for t in texts if t.strip().isupper() and "DAY BY DAY" in t or "PREVIOUS PERIOD" in t]
check("trend tab: only OVERALL — DAY BY DAY and its chart (no follow-up groups table)", _secs,
      ["OVERALL — DAY BY DAY", "TASK COMPLETION % VS AVERAGE ACTUAL % — DAY BY DAY"])
check("… nothing below the chart (no leftover rows / follow-up table)",
      [str(c) for c in texts if "FOLLOW-UP" in str(c).upper()], [])
_th = [r for r in tr.iter_rows(values_only=True) if r and r[0] == "Day"][0]
check("… table columns ('Areas On Target' removed)", [h for h in _th if h],
      ["Day", "Tasks Generated", "Tasks Completed", "Task Completion %", "Areas Measured", "Average Actual %"])
_tr_rows = [r for r in tr.iter_rows(values_only=True) if r and isinstance(r[0], str) and r[0][:2].isdigit()]
check("… one row per day of the period, Average Actual % = the day's measured areas averaged",
      [r[0] for r in _tr_rows], ["05-Oct (Mon)", "06-Oct (Tue)"])
_lc = tr._charts
check("… one line chart: Task Completion % (D) vs Average Actual % (F), labelled, 0–110% scale",
      (len(_lc), [x.val.numRef.f.split("!")[1][:2] for x in _lc[0].series],
       [bool(x.dLbls and x.dLbls.showVal) for x in _lc[0].series],
       _lc[0].y_axis.scaling.min, _lc[0].y_axis.scaling.max),
      (1, ["$D", "$F"], [True, True], 0, 110))
cov = [str(c.value) for row in wbw["Data Coverage & Rules"].iter_rows() for c in row if c.value is not None]
check("Data Coverage & Rules: grouped sections in reading order",
      [x for x in ("AT A GLANCE", "A.  WHAT DATA IS COVERED", "B.  WHERE THE ACTUAL PERFORMANCE",
                   "C.  RULES & DEFINITIONS", "1.  DATA SOURCE", "2.  REPORTING PERIODS", "3.  EFFORT",
                   "4.  ACTUAL PERFORMANCE", "5.  TARGETS", "6.  IMPORTANT LIMITATIONS")
       if not any(c.strip().startswith(x) for c in cov)], [])
check("… every rule shown exactly once (none lost by the grouping)",
      sorted(k for k, _v in P.RULES + P.OUTCOME_RULES if cov.count(k) < 1), [])
check("Data Coverage & Rules explains every outcome + targets + quadrants",
      [k for k in ("Attendance", "Assignment", "Admission", "Wise & IV Feedback", "Instructor", "Interview",
                   "Targets", "Quadrants", "Stored checks") if k not in cov], [])
_dcells = [str(c.value) for row in wbw["Dashboard"].iter_rows() for c in row if c.value is not None]
check("Dashboard: question and attention sections removed",
      [c for c in _dcells if "WHAT THIS REPORT ANSWERS" in c or "NEEDS MANAGEMENT ATTENTION" in c
       or c[:3] in ("1. ", "6. ")], [])
_dw = wbw["Dashboard"]
_tile = next((r, c) for r in range(1, _dw.max_row + 1) for c in range(1, _dw.max_column + 1)
             if _dw.cell(row=r, column=c).value == "Average Actual %")
_tot = [r for r in _dw.iter_rows(values_only=True) if r and r[0] == "All task groups"][0]
_avg = P.outcome_totals(rw2["outcome_rows"])["avg_actual"]   # wbw = the last upload (rw2)
check("Dashboard: 'Average Actual %' tile = scorecard total = average of the measured areas",
      (_dw.cell(row=_tile[0] + 1, column=_tile[1]).value, _tot[7]), (_avg, _avg))
check("Dashboard: Effort vs Actual Outcome chart still reads the scorecard columns E and H",
      [x.val.numRef.f.split("!")[1][:2] for x in _dw._charts[0].series], ["$E", "$H"])
check("… plus the previous-vs-current chart beside it (2 charts)", len(_dw._charts), 2)

# effort-only fallback still works (CHECK_OUTCOMES off / engine unavailable)
r_off = P.run_jobs(jobs0, versions[:1], loader, now0, upload=upload, engine=None)[0]
check("engine=None -> effort-only report, no outcome rows, no failure",
      (r_off.get("failed"), r_off.get("outcome_rows")), (None, None))

# =============================================================================
print("\n== 8. Evening batch ==")
ev_src = open(os.path.join(ROOT, "scripts", "run_evening_reports.py"), encoding="utf-8").read()
order = [ev_src.index(x) for x in ("pyZohoSignatureStatusRefresh", "pyStudentPaymentClassesStudentEnrolled",
                                   "pyAssignmentSubmissions", '("pyCoordinatorTaskPerformanceReport"')]
check("sources refreshed BEFORE the performance report", order == sorted(order), True)
check("refresh status handed to the report; cache capped",
      ("INTELLIBI_EVENING_REFRESH" in ev_src, "INTELLIBI_CACHE_MAX_AGE_SECONDS" in ev_src), (True, True))
os.environ["INTELLIBI_EVENING_REFRESH"] = "pyZohoSignatureStatusRefresh=SUCCESS;pyAssignmentSubmissions=FAILED"
check("report says which source was not refreshed", P._refresh_note().startswith(
    "NOT refreshed: pyAssignmentSubmissions"), True)
del os.environ["INTELLIBI_EVENING_REFRESH"]
col = open(os.path.join(ROOT, "ops_data_collection", "pyStudentPaymentClassesStudentEnrolled.py"),
           encoding="utf-8").read()
ns = {"os": os}
exec(col[col.index("_CACHE_CAP_EXEMPT ="):col.index("def _cache_get(")], ns)
os.environ["INTELLIBI_CACHE_MAX_AGE_SECONDS"] = "1800"
check("collector cache cap: students 12h -> 30 min; participant identities keep 24h",
      (ns["_effective_ttl"]("sp_students", 43200), ns["_effective_ttl"]("sp_participant", 86400)), (1800, 86400))
del os.environ["INTELLIBI_CACHE_MAX_AGE_SECONDS"]
check("… unset -> normal TTL", ns["_effective_ttl"]("sp_students", 43200), 43200)

# =============================================================================
print("\n== 9. Previous vs current period (Yesterday → Today) for the follow-up groups ==")
# pure rules
check("verdicts: up after high effort / down despite high effort / down with low effort / steady / n.a.",
      [CO.followup_result(100, 5.0), CO.followup_result(80, -21.0), CO.followup_result(40, -3.0),
       CO.followup_result(100, 0.6), CO.followup_result(None, 4.0), CO.followup_result(None, -4.0),
       CO.followup_result(60, 2.0), CO.followup_result(100, None)],
      [CO.FR_IMPROVED_AFTER, CO.FR_DECLINED_DESPITE, CO.FR_DECLINED_INCOMPLETE, CO.FR_STEADY,
       CO.FR_IMPROVED_NO_TASKS, CO.FR_DECLINED_NO_TASKS, CO.FR_IMPROVED, CO.FR_NA])
_cut = lambda d: datetime.combine(d, datetime.min.time()) + timedelta(hours=9, minutes=12)
check("Daily windows: attendance / instructor = yesterday 12:00 PM → task-list time; assignment = due D-1",
      [CO.measurement_window(g, "Daily", date(2026, 10, 8), date(2026, 10, 8), _cut)["text"]
       for g in ("attendance", "instructor", "assignment")],
      ["sessions 07-Oct 12:00 PM – 08-Oct 09:12 AM", "sessions 07-Oct 12:00 PM – 08-Oct 09:12 AM",
       "assignments due 07-Oct (Assignment Daily 07-Oct)"])
check("period windows: the period's own sessions / deadlines",
      [CO.measurement_window(g, "Weekly", date(2026, 9, 28), date(2026, 10, 4))["text"]
       for g in ("attendance", "assignment")],
      ["sessions dated 28-Sep – 04-Oct", "assignments due 28-Sep – 04-Oct"])
_n8 = datetime(2026, 10, 8, 19, 5)
_lab = lambda k, s, e, n=_n8: P.compare_labels(k, s, e, *P.previous_period(k, s, e), n)
check("labels by actual reporting dates (never the generation time)",
      [tuple(_lab("Daily", date(2026, 10, 8), date(2026, 10, 8))[k] for k in ("prev_tag", "cur_tag", "prev_dates", "cur_dates", "as_of", "observation")),
       ],
      [("Yesterday", "Today", "Wed 07-Oct-2026", "Thu 08-Oct-2026", "07:05 PM", "next_day")])
check("pinned past Daily -> 'Previous Day' / 'Report Day', no 'as of'",
      tuple(_lab("Daily", date(2026, 10, 2), date(2026, 10, 2))[k] for k in ("prev_tag", "cur_tag", "prev_dates", "as_of")),
      ("Previous Day", "Report Day", "Thu 01-Oct-2026", ""))
check("Weekly / Monthly / Manual -> previous period vs report period (same-period comparison)",
      [(_lab(k, s, e)["prev_tag"], _lab(k, s, e)["prev_dates"], _lab(k, s, e)["cur_dates"], _lab(k, s, e)["observation"])
       for k, s, e in (("Weekly", date(2026, 9, 28), date(2026, 10, 4)),
                       ("Monthly", date(2026, 9, 1), date(2026, 9, 30)),
                       ("Manual", date(2026, 9, 21), date(2026, 9, 25)))],
      [("Previous Week", "21-Sep – 27-Sep-2026", "28-Sep – 04-Oct-2026", "same_period"),
       ("Previous Month", "Aug-2026", "Sep-2026", "same_period"),
       ("Previous Period", "16-Sep – 20-Sep-2026", "21-Sep – 25-Sep-2026", "same_period")])
_o = lambda **kw: dict({"key": "attendance", "state": CO.ST_MEASURED, "pct": 70.0, "note": ""}, **kw)
check("availability: final / provisional / not available / nothing / period still running",
      [P.outcome_availability(_o(), "Daily", date(2026, 10, 8), date(2026, 10, 8))[0],
       P.outcome_availability(_o(source="provisional"), "Daily", date(2026, 10, 8), date(2026, 10, 8))[0],
       P.outcome_availability(_o(state=CO.ST_NOT_CHECKED, pct=None), "Daily", date(2026, 10, 8), date(2026, 10, 8))[0],
       P.outcome_availability(_o(state=CO.ST_NO_DATA, pct=None), "Daily", date(2026, 10, 8), date(2026, 10, 8))[0],
       P.outcome_availability(_o(), "Weekly", date(2026, 10, 11), date(2026, 10, 8))[0],
       P.outcome_availability(_o(), "Weekly", date(2026, 10, 4), date(2026, 10, 8))[0]],
      ["Final", "Provisional", "Not available", "Nothing to measure", "In progress", "Final"])
# the business example (illustrative figures): 72/72 → 71.6 %, 20/20 → 50.6 %
_s = lambda d, n: {"tasks": n, "completed": d, "completion_pct": d / n * 100 if n else None}
_v = {"enabled": True, "kind": "Daily", "now": _n8,
      "compare": _lab("Daily", date(2026, 10, 8), date(2026, 10, 8))}
_ex = P.compare_row("assignment", _s(20, 20), _o(key="assignment", pct=50.6), _s(72, 72),
                    _o(key="assignment", pct=71.6), _v, _n8)
check("example: Yesterday 72/72 · 71.6 % → Today 20/20 · 50.6 % = ▼ 21.0 pts, declined despite follow-ups",
      (_ex["prev_comp"], _ex["prev_actual"], _ex["cur_comp"], _ex["cur_actual"], _ex["change"], _ex["verdict"]),
      (100.0, 71.6, 100.0, 50.6, -21.0, CO.FR_DECLINED_DESPITE))
_na = P.compare_row("assignment", _s(20, 20), _o(key="assignment", state=CO.ST_NOT_CHECKED, pct=None,
                    note="Assignment report unreadable"), _s(72, 72), _o(key="assignment", pct=71.6), _v, _n8)
check("today's outcome not available -> status, no change, no verdict guessed",
      (_na["cur_actual"], _na["change"], _na["verdict"], _na["why"]),
      (None, None, CO.FR_NA, "Today's outcome: Not available — Assignment report unreadable"))

# end to end (Tuesday 06-Oct Daily, Monday as yesterday)
_cr = {r["gk"]: r for r in r1["outcome_rows"]}
check("only the three follow-up groups carry a comparison",
      sorted(k for k, r in _cr.items() if r.get("compare")), ["assignment", "attendance", "instructor"])
check("Daily labels: Yesterday (Mon 05-Oct) → Today (Tue 06-Oct), as of the run time",
      tuple(r1["compare"][k] for k in ("prev_tag", "prev_dates", "cur_tag", "cur_dates", "as_of")),
      ("Yesterday", "Mon 05-Oct-2026", "Today", "Tue 06-Oct-2026", "07:00 PM"))
_ic = _cr["instructor"]["compare"]
check("instructor: same figures as the report's own period outcomes and Δ (80 % → 80 %, 0 pts, Steady)",
      (_ic["prev_actual"], _ic["cur_actual"], _ic["change"], _cr["instructor"]["d_actual"], _ic["verdict"]),
      (80.0, 80.0, 0.0, 0.0, CO.FR_STEADY))
check("windows are aligned: today's starts at yesterday 12:00 PM, after yesterday's 10:30 AM list",
      (_ic["prev_window"], _ic["cur_window"]),
      ("sessions 04-Oct 12:00 PM – 05-Oct 10:30 AM", "sessions 05-Oct 12:00 PM – 06-Oct 10:30 AM"))
_ac = _cr["assignment"]["compare"]
check("assignment: today provisional (report not generated), yesterday nothing to measure -> not comparable",
      (_ac["cur_av"][0], _ac["prev_av"][0], _ac["verdict"]), ("Provisional", "Nothing to measure", CO.FR_NA))
_d1 = openpyxl.load_workbook(io.BytesIO(UP[r1["name"]]))
_dd = _d1["Dashboard"]
_dvals = [[c.value for c in row] for row in _dd.iter_rows()]
_sec = [v[0].strip() for v in _dvals if v and isinstance(v[0], str) and "FOLLOW-UP GROUPS: COORDINATOR EFFORT" in v[0]]
_si = lambda txt: [i for i, v in enumerate(_dvals) if v and isinstance(v[0], str) and v[0].strip() == txt][0]
check("Dashboard: Yesterday → Today block inside the scorecard section, before the charts",
      (len(_sec), _sec[0].startswith("YESTERDAY  →  TODAY") if _sec else None,
       _si("TASK GROUP SCORECARD") < _si(_sec[0]) < _si("EFFORT VS OUTCOME BY TASK GROUP")),
      (1, True, True))
_hi = [i for i, v in enumerate(_dvals) if v and v[0] == "Task Group  ·  Outcome Measure"][0]
check("… super-headers carry the actual dates", [_dvals[_hi - 1][c] for c in (1, 4, 7)],
      ["YESTERDAY  ·  Mon 05-Oct-2026", "TODAY  ·  Tue 06-Oct-2026  ·  as of 07:00 PM",
       "CHANGE  ·  FOLLOW-UP → RESULT"])
_brow = {str(v[0]).split("\n")[0]: (v, _hi + 1 + k) for k, v in enumerate(_dvals[_hi + 1:_hi + 4])}
_inst, _irow = _brow["Instructor Follow-Ups"]
check("… instructor row: Yesterday 0/5 · 0 % · 80 % | Today 0/5 · 0 % · 80 % | ● 0.0 | Steady",
      (_inst[1], _inst[3], _inst[4], _inst[6], _inst[7], _inst[10]),
      (_inst[1], 80.0, _inst[4], 80.0, 0.0, CO.FR_STEADY))
_asg, _arow = _brow["Learner Assignment Follow-Ups"]
check("… a figure that is not final is labelled: today provisional ('% · prov.'), yesterday shown as its status",
      (_dd.cell(row=_arow + 1, column=7).number_format, _asg[3]), ('0.0"% · prov."', "Nothing to measure"))
check("… change cells use ▲ / ▼ / ● formatting", _dd.cell(row=_irow + 1, column=8).number_format, P.CHANGE_FMT)
_ch = _dd._charts
_bf, _bl = _hi + 2, _hi + 4                      # the block's 3 data rows (1-based)
check("Dashboard: compare chart = 4 series (prev effort, prev outcome, cur effort, cur outcome) "
      "read from the Dashboard's own Yesterday → Today block (C, D, F, G)",
      ([s.tx.v for s in _ch[1].series], [s.val.numRef.f.split("!")[1] for s in _ch[1].series],
       all(s.val.numRef.f.startswith("'Dashboard'!") or s.val.numRef.f.startswith("Dashboard!")
           for s in _ch[1].series)),
      (["Yesterday · Effort", "Yesterday · Outcome", "Today · Effort", "Today · Outcome"],
       [f"${c}${_bf}:${c}${_bl}" for c in "CDFG"], True))
_lc0 = _ch[0]
_cats = _lc0.series[0].cat.numRef.f if _lc0.series[0].cat.numRef is not None else _lc0.series[0].cat.strRef.f
_r0, _r1 = [int(x.split("$")[-1]) for x in _cats.split("!")[1].split(":")]
check("left chart = only the scorecard's remaining groups (no blank categories), titled accordingly",
      ([_dd.cell(row=r, column=1).value for r in range(_r0, _r1 + 1)],
       _lc0.title.tx.rich.p[0].r[0].t if _lc0.title and _lc0.title.tx and _lc0.title.tx.rich else None),
      (["Learner Admission Formalities", "Wise & Interview Feedback Validation",
        "Learner Instructor Interview Reminder"], "Other task groups · Today"))
check("both charts share one height (no ragged blank space)", _ch[0].height == _ch[1].height, True)
_eo = openpyxl.load_workbook(io.BytesIO(UP[r_off["name"]]))["Dashboard"]
check("effort-only report (no comparison): scorecard keeps all six groups",
      len([r for r in _eo.iter_rows(values_only=True) if r and r[0] in [g["name"] for g in P.TASK_GROUPS]]), 6)
check("… previous = lighter tints of the same blue / orange",
      [str(getattr(s.graphicalProperties.solidFill.srgbClr, "val", s.graphicalProperties.solidFill.srgbClr)) for s in _ch[1].series],
      [P.CHART_PREV_EFFORT_HEX, P.CHART_PREV_OUTCOME_HEX, P.CHART_EFFORT_HEX, P.CHART_OUTCOME_HEX])
check("Effort vs Outcome Trend tab: no follow-up groups table",
      any("FOLLOW-UP GROUPS" in str(c.value or "") for row in _d1["Effort vs Outcome Trend"].iter_rows()
          for c in row), False)
_h1 = CE.performance_html("Daily", "06-Oct-2026", r1["summary"], r1["groups"], "https://x", "—",
                          benchmark=95.0, group_labels=P.EMAIL_GROUP_LABELS,
                          outcome_rows=r1["outcome_rows"], headline=r1["headline"],
                          attention=r1["attention"], compare=r1["compare"])
_g1 = _h1.split("Performance vs Goals", 1)[1]
_card = _g1.split("Instructor Instructions", 1)[1].split("</table></td></tr></table>", 1)[0]
check("e-mail: follow-up group card = Yesterday | Today | Change table + paired outcome bars + verdict",
      [x in _card for x in ("Yesterday", "Mon 05-Oct-2026", "Today", "Tue 06-Oct-2026",
                            "Tasks Completed", "Effort &middot; Task Completion %",
                            "Outcome &middot; Sessions Without Escalation %", "&#9679;&nbsp;0.0</span>",
                            "Follow-up &rarr; Result", "Steady", "sessions 05-Oct 12:00 PM")],
      [True] * 11)
check("e-mail: the three groups are named in one reading note, with the 'not proof of cause' caveat",
      ("Learner Attendance, Learner Assignment, Instructor Instructions</b> are shown Yesterday &rarr; Today" in _g1,
       "not proof that the follow-ups caused them" in _g1), (True, True))
check("e-mail: provisional / unavailable outcomes shown as status, not as a plain number",
      ("prov.</span>" in _g1 and "Nothing to measure" in _g1), True)
check("e-mail: other groups keep their single-period card (Admission effort + outcome bars)",
      "Outcome &middot; Formalities Resolved %" in _g1, True)
check("e-mail without comparison labels (effort-only / old caller) is unchanged",
      "&#9679;" in CE.performance_html("Daily", "06-Oct-2026", r1["summary"], r1["groups"], "https://x", "—",
                                        outcome_rows=r1["outcome_rows"]), False)
_rwc = {r["gk"]: r.get("compare") for r in rw["outcome_rows"]}
check("Manual report: previous period vs report period, same-period wording",
      (rw["compare"]["prev_tag"], rw["compare"]["cur_tag"], rw["compare"]["observation"],
       _rwc["instructor"]["verdict"] in CO.FR_LEVEL), ("Previous Period", "Report Period", "same_period", True))
check("rules tab explains the comparison", any(k == "Previous vs current" for k, _v in P.OUTCOME_RULES), True)

print("\nALL CHECKS PASSED" if not FAIL else f"\n{len(FAIL)} CHECK(S) FAILED: {FAIL}")
sys.exit(1 if FAIL else 0)
