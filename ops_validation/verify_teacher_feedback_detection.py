"""
Verification of the instructor-feedback detection fix (run from the project root:
  python ops_validation\\verify_teacher_feedback_detection.py).

Reproduces, generically, the reported case: the LMS creates a Teacher_Feedback
record the moment the instructor marks a session Completed (session_status =
COMPLETED, created_at = session end) BEFORE any feedback is written - the row has
blank "topics_covered" and "comments" and the LMS shows "Instructor Feedback:
Pending". Such a session must appear in Teacher_No_Feedback (AR) and count as
"Feedback Given = No" in the Batch Coordinator Instructor Follow-Ups, while a
session whose row carries feedback content must NOT.

Runs the REAL AR builders (daily + grouped + period Teacher_No_Feedback) on an
in-memory workbook with synthetic sessions - no Google access needed.
"""
import os
import sys
import types
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

# ── stub the project bootstrap / credential modules the report imports ──────
sys.modules["_bootstrap"] = types.ModuleType("_bootstrap")
paths = types.ModuleType("paths")
paths.CREDENTIALS_DIR = paths.CONFIG_DIR = paths.LOGS_DIR = paths.CACHE_DIR = HERE
sys.modules["paths"] = paths
ec = types.ModuleType("email_config"); ec.GMAIL_SENDER = "x@x"; ec.GMAIL_APP_PASS = "x"
sys.modules["email_config"] = ec

sys.path.insert(0, os.path.join(ROOT, "ops_reports_action"))
import pandas as pd                                   # noqa: E402
import openpyxl                                       # noqa: E402
import pyAttendaceFeedbackReport as AR                # noqa: E402  (the live script)


def sess(sid, course, tutor, start, end, d):
    return {"session_id": sid, "course_name": course, "course_title": "23-Sep-2026 To Current Date",
            "tutor_name": tutor, "start_time_ist": start, "end_time_ist": end, "_date": d}


D = date(2026, 9, 30)
sessions = pd.DataFrame([
    sess("S_TOPICS",   "Power BI", "Instructor 1", "2026-09-30 06:56:20", "2026-09-30 09:01:45", D),  # feedback: topics
    sess("S_COMMENTS", "SQL",      "Instructor 2", "2026-09-30 18:50:00", "2026-09-30 21:00:00", D),  # feedback: comments only
    sess("S_PLACEHOLDER", "Power BI", "Instructor 1", "2026-09-30 06:56:20", "2026-09-30 09:01:45", D),  # completion stamp only
    sess("S_NOROW",    "Python",   "Instructor 3", "2026-09-30 06:50:54", "2026-09-30 09:02:42", D),  # no TF row at all
    sess("S_SCHEDULED", "SQL",     "Instructor 2", "2026-10-01 19:00:00", "",                    date(2026, 10, 1)),  # not yet held
])
tf = pd.DataFrame([
    {"session_id": "S_TOPICS",      "topics_covered": "Power Query – ETL", "comments": "",  "session_status": "COMPLETED"},
    {"session_id": "S_COMMENTS",    "topics_covered": "",                  "comments": "Ran long", "session_status": "COMPLETED"},
    {"session_id": "S_PLACEHOLDER", "topics_covered": "",                  "comments": "",  "session_status": "COMPLETED"},
    {"session_id": "S_PLACEHOLDER", "topics_covered": "nan",               "comments": "None", "session_status": "COMPLETED"},  # stringified blanks
])
att = pd.DataFrame([{"session_id": s, "student_id": "st1", "duration": "7200", "attendance_percent": "90"}
                    for s in ("S_TOPICS", "S_COMMENTS", "S_PLACEHOLDER", "S_NOROW")])

# ── 1. the shared rule ───────────────────────────────────────────────────────
given = AR.teacher_feedback_session_ids(tf)
assert given == {"S_TOPICS", "S_COMMENTS"}, given
assert AR.teacher_feedback_session_ids(pd.DataFrame()) == set()
assert AR.teacher_feedback_session_ids(None) == set()
# older layout without content columns -> presence-only (previous behaviour)
legacy = AR.teacher_feedback_session_ids(pd.DataFrame([{"session_id": "X"}]))
assert legacy == {"X"}
print("[pass] feedback 'given' = Teacher_Feedback row with topics or comments; placeholder rows excluded")


def sheet_rows(ws, first_data_row):
    out = []
    for r in ws.iter_rows(min_row=first_data_row, values_only=True):
        if any(v not in (None, "") for v in r):
            out.append(list(r))
    return out


# ── 2. AR daily Teacher_No_Feedback ──────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.create_sheet("Teacher_No_Feedback")
AR.build_teacher_no_feedback(ws, sessions.copy(), tf.copy(), "daily", "30-Sep-2026", att_f=att, yest_date=None)
rows = sheet_rows(ws, 4)
by_sid = {r[0] + "|" + r[3]: r for r in rows}
listed = {(r[0], r[2], r[3], r[5]) for r in rows}
assert ("Power BI", "Instructor 1", "2026-09-30 06:56:20", "No feedback received") in listed, listed
assert ("Python", "Instructor 3", "2026-09-30 06:50:54", "No feedback received") in listed
assert ("SQL", "Instructor 2", "2026-10-01 19:00:00", "Scheduled") in listed
assert not any(r[3] == "2026-09-30 18:50:00" for r in rows), "SQL session with comments-only feedback must NOT be listed"
# the Power BI session with real feedback shares course/tutor/time with the placeholder one -> exactly ONE Power BI row
assert sum(1 for r in rows if r[0] == "Power BI") == 1
# every other column populated for the flagged row (course title, tutor, start, end, remark)
pb = next(r for r in rows if r[0] == "Power BI")
assert pb[1] == "23-Sep-2026 To Current Date" and pb[2] == "Instructor 1" and pb[4] == "2026-09-30 09:01:45", pb
print("[pass] daily Teacher_No_Feedback lists the placeholder-only and no-row sessions with all columns filled")

# ── 3. AR grouped (weekly-style) Teacher_No_Feedback ─────────────────────────
ws = wb.create_sheet("TNF_weekly")
AR.build_teacher_no_feedback(ws, sessions.copy(), tf.copy(), "weekly", "wk", att_f=att)
rows = {(r[0], r[2]): r for r in sheet_rows(ws, 4)}
assert rows[("Power BI", "Instructor 1")][3:6] == [2, 1, 1], rows[("Power BI", "Instructor 1")]
assert rows[("Python", "Instructor 3")][3:6] == [1, 0, 1]
assert rows[("SQL", "Instructor 2")][3:6] == [2, 1, 1]
print("[pass] grouped Teacher_No_Feedback counts: with/without feedback follow the same rule")

# ── 4. AR period Teacher Feedback Status ─────────────────────────────────────
ws = wb.create_sheet("TNF_period")
AR.build_period_teacher_no_feedback(ws, sessions.copy(), tf.copy(), "monthly", "Sep-2026")
rows = {(r[0], r[2]): r for r in sheet_rows(ws, 3)}
assert rows[("Power BI", "Instructor 1")][3:6] == [2, 1, 1], rows[("Power BI", "Instructor 1")]
assert rows[("Python", "Instructor 3")][3:6] == [1, 0, 1]
print("[pass] period Teacher Feedback Status uses the same rule")

# ── 5. Batch Coordinator report uses the SAME helper (no second definition) ──
bc = open(os.path.join(ROOT, "co-ordinator reports", "pyCoordinatorTaskListReport.py"),
          encoding="utf-8").read()
assert bc.count("AR.teacher_feedback_session_ids(") == 2, "daily + period Instructor Follow-Ups must both use the shared rule"
assert 'set(tf_daily["session_id"]' not in bc and 'set(tf_period["session_id"]' not in bc
print("[pass] Instructor Follow-Ups (daily + period) share AR.teacher_feedback_session_ids")

print("ALL CHECKS PASSED")
