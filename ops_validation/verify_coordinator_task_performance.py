"""
Verification of pyCoordinatorTaskPerformanceReport.py (run from the project root:
  python ops_validation\\verify_coordinator_task_performance.py).

Builds REAL Coordinator task tabs with pyCoordinatorTaskListReport's own
builders (Admission Formalities, Assignment Follow-Ups, Wise Validation), fills the
follow-up columns the way the Coordinator would, across two report days and several
report versions, and checks the performance ledger against hand-computed answers:
  * one task per day however many versions list it (no double counting)
  * a task actioned in an early version and dropped later is still counted;
    an un-actioned task that dropped off is not
  * Done? = Yes in ANY version completes the task; earliest valid stamp is used
  * #REF! stamp -> completed, excluded from timing
  * on time / late / missed / open, Timely %, Days on List, median time
  * Progress Trend layout (Daily and multi-day): WITHIN-DAY PROGRESS first with the
    hour-by-hour table + chart under it, the two removed sections gone, charts inside
    the grid and clear of cells, and period hour figures = the days summed by hour
No Google access needed.
"""
import os
import re
import sys
import types
from datetime import date, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.modules["_bootstrap"] = types.ModuleType("_bootstrap")
paths = types.ModuleType("paths")
paths.CREDENTIALS_DIR = paths.CONFIG_DIR = paths.LOGS_DIR = paths.CACHE_DIR = HERE
sys.modules["paths"] = paths
ec = types.ModuleType("email_config"); ec.GMAIL_SENDER = "x@x"; ec.GMAIL_APP_PASS = "x"
sys.modules["email_config"] = ec
for p in ("common", "ops_reports_action", "co-ordinator reports"):
    sys.path.insert(0, os.path.join(ROOT, p))

import openpyxl                                          # noqa: E402
import pyCoordinatorTaskListReport as BC     # noqa: E402  (the live task report)
import pyCoordinatorTaskPerformanceReport as P           # noqa: E402  (the new report)

FAIL = []


def check(label, got, want):
    ok = got == want
    print(f"[{'pass' if ok else 'FAIL'}] {label}: {got!r}" + ("" if ok else f"  (expected {want!r})"))
    if not ok:
        FAIL.append(label)


def adm(name, email):
    return {"Student Name": name, "Email ID": email, "Phone Number": "+919000000001",
            "Batch Name": "DAAI1026", "Joined On": "2026-10-01", "Request Form Name": "",
            "Recipient Status": "Form Not Sent", "Request Status": "", "Sent Date": "",
            "Signed Date": "", "Expiry Date": ""}


def asg(name, lvl="1st"):
    return {"student_name": name, "student_email": name.lower().replace(" ", ".") + "@x.com",
            "student_phone": "+919000000002", "deadline_str": "08-Oct-2026", "reminder_level": lvl,
            "reminder_label": "1st Reminder", "assigned_date_str": "04-Oct-2026",
            "class_subject": "01-Oct-2026 To Current Date", "maximum_marks": ""}


def report(day, adm_rows, asg_map, wise_students=None):
    """A Coordinator daily workbook built by the real builders."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    BC.build_admission_formalities(wb.create_sheet("Learner Admission Formalities"), adm_rows, day)
    by_tech = {"Power BI": [(title, f"id-{title}", [asg(n) for n in names])
                            for title, names in asg_map.items()]}
    BC.build_assignment_followups(wb.create_sheet("Learner Assignment Follow-Ups"), by_tech, day)
    if wise_students is not None:
        data = {"student": [{"Student Name": n, "Batch Name": "DAAI1026", "Joined On": "",
                             "Student Name Status": "Valid", "Email ID Status": "Valid",
                             "Phone Number Status": "Valid", "Tag Name Status": "Missing",
                             "Private Note Status": "Valid", "Profile Picture Status": "Valid"}
                            for n in wise_students], "course": [], "instructor": []}
        BC.build_wise_validation(wb.create_sheet("Wise & Interview Feedback Validation"), data, day)
    return wb


def fill(wb, tab, who, done=None, stamp=None, action=None, section=None):
    """Type a follow-up into the row whose cells contain `who` (below the banner
    containing `section`, when given) — what the Coordinator does in the sheet."""
    ws = wb[tab]
    hdr, in_section = None, section is None
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        txt = [str(v or "") for v in vals]
        if section and any(section in t for t in txt[:1]):
            in_section = True
        if any(t.strip() == "Why Flagged" for t in txt):
            hdr = {t.replace("✎", "").strip(): i + 1 for i, t in enumerate(txt) if t}
            continue
        if hdr and in_section and any(t == who for t in txt):
            if done is not None:
                ws.cell(row=r, column=hdr["Follow-Up Done?"]).value = done
            if stamp is not None:
                ws.cell(row=r, column=hdr["Follow-Up DateTime"]).value = stamp
            if action is not None:
                ws.cell(row=r, column=hdr["Action Taken"]).value = action
            return
    raise AssertionError(f"row {who!r} not found in {tab}")


def clear_formulas(wb):
    """Google stores the stamp's VALUE; an untouched row exports blank."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=IF("):
                    c.value = None


D1, D2 = date(2026, 10, 5), date(2026, 10, 6)
ADM, ASG = "Learner Admission Formalities", "Learner Assignment Follow-Ups"

# ── Day 1, Version 1 (09:00) ────────────────────────────────────────────────
d1v1 = report(D1, [adm("Asha Rao", "asha@x.com"), adm("Bala K", "bala@x.com"), adm("Chitra M", "chitra@x.com")],
              {"Project X": ["Lina P", "Mohan S"], "Project Y": ["Lina P"]})
clear_formulas(d1v1)
fill(d1v1, ADM, "Asha Rao", "Yes", datetime(2026, 10, 5, 10, 0), "Call")
fill(d1v1, ADM, "Bala K", "No", datetime(2026, 10, 5, 10, 30), "WhatsApp")
fill(d1v1, ASG, "Lina P", "Yes", "#REF!", "WhatsApp Group", section="Project X")
# ── Day 1, Version 2 (12:00): Asha dropped off, Dev added, Mohan (X) dropped off ─
d1v2 = report(D1, [adm("Bala K", "bala@x.com"), adm("Chitra M", "chitra@x.com"), adm("Dev N", "dev@x.com")],
              {"Project X": ["Lina P"], "Project Y": ["Lina P"]})
clear_formulas(d1v2)
fill(d1v2, ADM, "Bala K", "Yes", datetime(2026, 10, 5, 13, 0), "Call")
fill(d1v2, ADM, "Chitra M", "Yes", datetime(2026, 10, 6, 10, 0), "Call")      # done next day -> late
# ── Day 2, Version 1 (09:00) ────────────────────────────────────────────────
d2v1 = report(D2, [adm("Chitra M", "chitra@x.com"), adm("Dev N", "dev@x.com")],
              {"Project Y": ["Lina P"]}, wise_students=["Wasim Q"])
clear_formulas(d2v1)
fill(d2v1, ADM, "Dev N", "Yes", datetime(2026, 10, 6, 11, 0), "Form Send")

BOOKS = {"d1v1": d1v1, "d1v2": d1v2, "d2v1": d2v1}
versions = [
    {"id": "d1v1", "name": "d1v1", "day": D1, "version": 1, "created": datetime(2026, 10, 5, 9, 0), "modified_raw": ""},
    {"id": "d1v2", "name": "d1v2", "day": D1, "version": 2, "created": datetime(2026, 10, 5, 12, 0), "modified_raw": ""},
    {"id": "d2v1", "name": "d2v1", "day": D2, "version": 1, "created": datetime(2026, 10, 6, 9, 0), "modified_raw": ""},
]
now = datetime(2026, 10, 6, 18, 0)
ledger = P.build_ledger(versions, lambda v: P.parse_report_workbook(BOOKS[v["id"]]), now)
T = {(t["day"], t["group"], t["label"].split("  ·  ")[0], t.get("context", "")): t for t in ledger["tasks"]}


def status(day, group, who, ctx_has=""):
    m = [t for (d, g, w, c), t in T.items() if d == day and g == group and w == who and ctx_has in c]
    return m[0]["status"] if len(m) == 1 else f"{len(m)} matches"


check("day-1 tasks (no double counting across 2 versions; dropped un-actioned task excluded)",
      sum(1 for t in ledger["tasks"] if t["day"] == D1), 6)
check("Asha: actioned in V1, dropped in V2 -> still counted, on time", status(D1, "admission", "Asha Rao"), P.ST_ON_TIME)
check("Bala: No in V1, Yes in V2 -> completed on time", status(D1, "admission", "Bala K"), P.ST_ON_TIME)
check("Chitra day 1: Yes stamped next day -> late", status(D1, "admission", "Chitra M"), P.ST_LATE)
check("Dev day 1: never done on day 1 -> missed", status(D1, "admission", "Dev N"), P.ST_MISSED)
check("Lina / Project X: Yes with #REF! stamp -> completed, time not recorded",
      status(D1, "assignment", "Lina P", "Project X"), P.ST_UNKNOWN)
check("Lina / Project Y day 1 -> missed", status(D1, "assignment", "Lina P", "Project Y"), P.ST_MISSED)
check("Mohan / Project X (listed in V1 only, never actioned) -> not a task",
      status(D1, "assignment", "Mohan S", "Project X"), "0 matches")
check("Dev day 2 -> on time", status(D2, "admission", "Dev N"), P.ST_ON_TIME)
check("Chitra day 2 -> open (due today)", status(D2, "admission", "Chitra M"), P.ST_OPEN)
check("Wise record day 2 -> open", status(D2, "wise", "Wasim Q"), P.ST_OPEN)

lina_y2 = [t for t in ledger["tasks"] if t["day"] == D2 and t["group"] == "assignment"][0]
check("Days on List: Lina / Project Y pending on both days", lina_y2["days_on_list"], 2)
chitra2 = [t for t in ledger["tasks"] if t["day"] == D2 and t["label"].startswith("Chitra")][0]
check("Days on List: Chitra day 2 (day-1 task was completed) ", chitra2["days_on_list"], 1)

s = P.summarise(ledger["tasks"])
check("total tasks", s["tasks"], 10)
check("completed", s["completed"], 5)
check("pending (open + missed)", (s["pending"], s["open"], s["missed"]), (5, 3, 2))
check("completion %", round(s["completion_pct"], 1), 50.0)
check("timely % = on time / tasks with measurable timing = 3 / 9", round(s["timely_pct"], 1), 33.3)
check("median time to complete (1h, 4h, 2h)", s["median_ttc"], 2.0)
check("pending 2+ report days", s["carried"], 1)
check("unique items", s["unique_items"], 7)

cov = {c["day"]: c for c in ledger["coverage"]}
check("coverage day 1: versions / dropped / multi-version / no-time / Yes-No conflicts",
      (cov[D1]["versions"], cov[D1]["dropped"], cov[D1]["multi_version"], cov[D1]["yes_no_time"],
       cov[D1]["conflicts"]), (2, 1, 1, 1, 1))

# the workbook builds end-to-end for both a day and a multi-day period
wb, tasks, summ = P.build_workbook(ledger, D1, D2, "Test", False, now)
check("workbook tabs", wb.sheetnames,
      ["Dashboard", "Progress Trend", "Task Register", "Data Coverage & Rules"])
wb, tasks, summ = P.build_workbook(ledger, D2, D2, "Test", True, now)
check("daily workbook task count", len(tasks), 4)

# ── Progress Trend tab layout (Daily and multi-day periods) ─────────────────
def trend_layout(start, end, is_daily):
    wb_, _t, _s = P.build_workbook(ledger, start, end, "Test", is_daily, now)
    ws = wb_["Progress Trend"]
    texts = [str(c.value) for r in ws.iter_rows() for c in r if c.value is not None]
    names = ("WITHIN-DAY PROGRESS", "DAY-WISE PROGRESS", "COMPLETION % BY TASK GROUP AND DAY",
             "WHEN ARE TASKS COMPLETED", "HOW QUICKLY ARE TASKS COMPLETED")
    secs = [(r, str(ws.cell(row=r, column=1).value).strip()) for r in range(1, ws.max_row + 1)
            if str(ws.cell(row=r, column=1).value or "").strip().startswith(names)]
    return wb_, ws, texts, secs


def anchor_rc(c):
    """(row, column), 1-based, of a chart's top-left anchor cell."""
    from openpyxl.utils.cell import coordinate_to_tuple
    if isinstance(c.anchor, str):
        return coordinate_to_tuple(c.anchor)
    return c.anchor._from.row + 1, c.anchor._from.col + 1


def chart_checks(ws, label):
    charts = ws._charts
    beyond = [c.title.tx.rich.p[0].r[0].t for c in charts if anchor_rc(c)[1] > 10]
    check(f"{label}: every chart sits inside the 10-column grid (none off to the right)", beyond, [])
    clash = []
    for c in charts:
        r0 = anchor_rc(c)[0]
        n = int(-(-c.height * 28.35 // P.TREND_ROW_PT))
        for r in range(r0, r0 + n):
            if any(ws.cell(row=r, column=k).value not in (None, "") for k in range(1, 11)):
                clash.append((c.title.tx.rich.p[0].r[0].t, r))
                break
    check(f"{label}: no chart covers table / section cells", clash, [])
    seen, overlap = set(), False
    for mr in ws.merged_cells.ranges:
        for rr in range(mr.min_row, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1):
                overlap = overlap or (rr, cc) in seen
                seen.add((rr, cc))
    check(f"{label}: no overlapping merged cells", overlap, False)
    return charts


def title_of(c):
    return c.title.tx.rich.p[0].r[0].t


for start, end, daily, lab in ((D2, D2, True, "Daily"), (D1, D2, False, "Period")):
    wb_, ws, texts, secs = trend_layout(start, end, daily)
    when = start.strftime("%d-%b-%Y") if start == end else \
        f"{start.strftime('%d-%b-%Y')} to {end.strftime('%d-%b-%Y')}"
    check(f"{lab}: first section is WITHIN-DAY PROGRESS — <period>, right under the title",
          (secs[0][0], secs[0][1].startswith(f"WITHIN-DAY PROGRESS — {when}")), (4, True))
    check(f"{lab}: hour-by-hour table directly under it",
          ws.cell(row=secs[0][0] + 1, column=1).value, "Hour (IST)")
    check(f"{lab}: 'When are tasks completed' and 'How quickly' sections removed",
          [t for t in texts if "WHEN ARE TASKS COMPLETED" in t or "HOW QUICKLY" in t], [])
    check(f"{lab}: section order", [s.split(" — ")[0] for _r, s in secs],
          ["WITHIN-DAY PROGRESS"] if daily else
          ["WITHIN-DAY PROGRESS", "DAY-WISE PROGRESS", "COMPLETION % BY TASK GROUP AND DAY"])
    charts = chart_checks(ws, lab)
    check(f"{lab}: charts", [title_of(c) for c in charts],
          ["Generated vs Completed vs Open — hour by hour"] + ([] if daily else
          ["Each report day's tasks by outcome", "Completion % and Timely % by report day"]))
    hrow = secs[0][0] + 1
    last = max(r for r in range(hrow + 1, ws.max_row + 1)      # last hour row of the table
               if re.match(r"^\d{1,2} (AM|PM)$", str(ws.cell(row=r, column=1).value or ""))
               and r < (secs[1][0] if len(secs) > 1 else 10 ** 6))
    check(f"{lab}: hour chart anchored at column A just under the table (one spacer row)",
          anchor_rc(charts[0])[::-1], (1, last + 2))
    blank = run_ = 0
    for r in range(1, ws.max_row + 1):
        empty = all(ws.cell(row=r, column=k).value in (None, "") for k in range(1, 11))
        chart_row = any(anchor_rc(c)[0] <= r <= anchor_rc(c)[0] +
                        int(-(-c.height * 28.35 // P.TREND_ROW_PT)) for c in charts)
        run_ = run_ + 1 if empty and not chart_row else 0
        blank = max(blank, run_)
    check(f"{lab}: no blank gap longer than one row (outside chart areas)", blank <= 1, True)

# within-day figures: a single day is unchanged; a period = its days combined by hour
check("single-day hour figures identical whether given a day or [day]",
      P.hourly_progress(ledger["tasks"], D2, now), P.hourly_progress(ledger["tasks"], [D2], now))
pooled = P.hourly_progress(ledger["tasks"], [D1, D2], now)
ts12 = [t for t in ledger["tasks"] if t["day"] in (D1, D2)]
expect = []
for r in pooled:
    h = datetime.strptime(r["label"], "%I %p").hour
    g = sum(1 for t in ts12 if not t["first_seen"] or t["first_seen"].date() < t["day"]
            or (t["first_seen"].date() == t["day"] and t["first_seen"].hour <= h))
    d = sum(1 for t in ts12 if t["completed_at"] and t["completed_at"].date() == t["day"]
            and t["completed_at"].hour <= h)
    expect.append((r["label"], g, d, g - d))
check("period hour figures = every report day's own figures summed by hour (independent recount)",
      [(r["label"], r["generated"], r["completed"], r["open"]) for r in pooled], expect)
# ── status-based row colours (presentation only) ──────────────────────────
def _fill(c):
    return (c.fill.fgColor.rgb or "")[-6:].upper()


def _contrast(a, b):
    """WCAG contrast ratio of two hex colours."""
    def lum(h):
        ch = [int(h[-6:][i:i + 2], 16) / 255 for i in (0, 2, 4)]
        ch = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in ch]
        return 0.2126 * ch[0] + 0.7152 * ch[1] + 0.0722 * ch[2]
    hi, lo = sorted((lum(a), lum(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


_wb, _t, _s = P.build_workbook(ledger, D1, D2, "Test", False, now)
_reg = _wb["Task Register"]
_h = [c.value for c in _reg[3]]
_si, _di = _h.index("Status") + 1, _h.index("Days on List") + 1
# expected colour per task, re-derived from the ledger (register order = day, group, item)
_order = list(ledger["groups"])
_exp = []
for t in sorted(_t, key=lambda t: (t["day"], _order.index(t["group"]), t["label"])):
    if t["completed"]:
        _exp.append((t["status"], t.get("days_on_list", 0), "completed"))
    elif t["status"] == P.ST_MISSED or t.get("days_on_list", 0) >= 2:
        _exp.append((t["status"], t.get("days_on_list", 0), "urgent"))
    else:
        _exp.append((t["status"], t.get("days_on_list", 0), "attention"))
_name = {v[0]: k for k, v in P.TASK_ROW_COLORS.items()}         # row bg -> category
_got, _bad, _chip_bad, _accent_bad = [], [], [], []
for i, r in enumerate(range(4, 4 + len(_exp))):
    cat = _name.get(_fill(_reg.cell(row=r, column=1)), "?")
    _got.append((_reg.cell(row=r, column=_si).value, _exp[i][1], cat))
    row_bg, chip_bg, fg = P.TASK_ROW_COLORS.get(cat, ("", "", ""))
    for c in range(1, len(_h) + 1):
        if c != _si and _fill(_reg.cell(row=r, column=c)) != row_bg:
            _bad.append((r, _h[c - 1]))
    if (_fill(_reg.cell(row=r, column=_si)), (_reg.cell(row=r, column=_si).font.color.rgb or "")[-6:].upper()) \
            != (chip_bg, fg):
        _chip_bad.append(r)
    if (_reg.cell(row=r, column=1).border.left.color.rgb or "")[-6:].upper() != fg:
        _accent_bad.append(r)
check("Task Register: all five task statuses present in the test data",
      sorted({e[0] for e in _exp}),
      sorted([P.ST_ON_TIME, P.ST_LATE, P.ST_UNKNOWN, P.ST_OPEN, P.ST_MISSED]))
check("row colour by severity: completed = green, open = orange, missed or open 2+ days = red",
      _got, [(e[0], e[1], {"completed": "completed", "attention": "attention", "urgent": "urgent"}[e[2]])
             for e in _exp])
check("… covers an OPEN task already pending 2+ consecutive days (red) and a plain open one (orange)",
      ((P.ST_OPEN, "urgent") in {(g[0], g[2]) for g in _got}, (P.ST_OPEN, "attention") in {(g[0], g[2]) for g in _got}),
      (True, True))
check("… the whole row carries the colour (every cell)", _bad, [])
check("… Status chip: stronger shade of the same colour, dark matching text", _chip_bad, [])
check("… coloured left edge on the first cell", _accent_bad, [])
check("palette: green / orange / red hues, light backgrounds, dark text (contrast >= 7:1)",
      all(_contrast(bg, P.BC.DS_TEXT) >= 7 and _contrast(chip, fg) >= 4.5
          for bg, chip, fg in P.TASK_ROW_COLORS.values())
      and [max(range(3), key=lambda i: int(v[0][2 * i:2 * i + 2], 16)) for v in
           (P.TASK_ROW_COLORS["completed"], P.TASK_ROW_COLORS["attention"], P.TASK_ROW_COLORS["urgent"])]
      == [1, 0, 0], True)
check("Task Register guide carries the colour key", "Green = completed" in str(_reg["A2"].value)
      and "Orange = attention" in str(_reg["A2"].value) and "Red = overdue" in str(_reg["A2"].value), True)
_dl = [r for r in range(4, _reg.max_row + 1)
       if isinstance(_reg.cell(row=r, column=_di).value, int) and _reg.cell(row=r, column=_di).value >= 2]
check("'Days on List' 2+ highlight kept on top of the row colour",
      [(_reg.cell(row=r, column=_di).font.color.rgb or "")[-6:].upper() for r in _dl],
      [P.BC.DS_HIGH_FG.upper()[-6:]] * len(_dl))
check("header row unchanged (navy, white text)",
      (_fill(_reg.cell(row=3, column=1)), (_reg.cell(row=3, column=1).font.color.rgb or "")[-6:].upper()),
      (P.BC.DS_NAV2.upper()[-6:], "FFFFFF"))
_dash = _wb["Dashboard"]
_rows = {_dash.cell(row=r, column=1).value: r for r in range(1, _dash.max_row + 1)}
_bad = []
for gk, meta, gs in P._group_rows(P.scope_ledger(ledger, D1, D2), _t):
    r = _rows[meta["name"]]
    lvl = P._status_level(gs["status"]) if gs["tasks"] else "muted"
    for c in (2, 3, 4, 6):
        if _fill(_dash.cell(row=r, column=c)) != P.ROW_TINT[lvl]:
            _bad.append((meta["name"], c))
    if not gs["tasks"] and _fill(_dash.cell(row=r, column=5)) != P.ROW_TINT["muted"]:
        _bad.append((meta["name"], 5))
check("Dashboard scorecard: rows tinted by Status (On track / Watch / Behind / No tasks)", _bad, [])
check("… total row keeps its own total styling",
      _fill(_dash.cell(row=_rows["All task groups"], column=1)), P.BC.DS_SUB.upper())
check("removed display-only helpers are gone (TTC buckets, hour-of-day profile)",
      [n for n in ("ttc_profile", "completion_hour_profile", "TTC_BUCKETS") if hasattr(P, n)], [])
check("median time to complete still calculated (Dashboard / e-mail)", s["median_ttc"], 2.0)

# a report version WITHOUT follow-up columns is never scored
old = report(date(2026, 10, 4), [adm("Old L", "old@x.com")], {})
for ws in old.worksheets:            # drop the follow-up headers, as in pre-column reports
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("✎"):
                c.value = None
led2 = P.build_ledger([{"id": "o", "name": "o", "day": date(2026, 10, 4), "version": 1,
                        "created": datetime(2026, 10, 4, 9, 0), "modified_raw": ""}],
                      lambda v: P.parse_report_workbook(old), now)
check("pre-follow-up report: 0 scored tasks, listed as not trackable",
      (len(led2["tasks"]), led2["coverage"][0]["untracked_tasks"]), (0, 1))

print("\nALL CHECKS PASSED" if not FAIL else f"\n{len(FAIL)} CHECK(S) FAILED: {FAIL}")
sys.exit(1 if FAIL else 0)
