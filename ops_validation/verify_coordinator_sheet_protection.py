"""
Verification of the Google Sheet access & protection of pyCoordinatorTaskListReport.py
(run from the project root:  python ops_validation\\verify_coordinator_sheet_protection.py).

Builds REAL Coordinator task workbooks with the report's own builders (Admission
Formalities, Assignment Follow-Ups, Wise Validation with several sections), uploads
them through the real upload_report() into an in-memory Google Drive + Sheets that
applies a batchUpdate all-or-nothing and enforces protected ranges the way Google
does, and checks:
  1. the Coordinator can edit ONLY Action Taken / Follow-Up Comment / Follow-Up Done?
     on the follow-up rows of every tab (not headers, not Follow-Up DateTime, not data)
  2. info@ (owner) and the pipeline service account can still edit everything
  3. the Coordinator cannot insert / delete rows or columns, or rename / delete tabs
  4. the Coordinator gets "writer" on the file only AFTER protection is verified;
     editors cannot re-share; no notification e-mail
  5. every uploaded sheet (Daily, Version 2, Weekly roll-up) gets the protection
  6. failures (bad workbook, batchUpdate error, read-back mismatch) never share the
     sheet for editing and never stop the run; transient errors retry without
     duplicate protections; re-applying is idempotent
  7. the rename: old script name gone everywhere, scheduler / imports / docs use
     pyCoordinatorTaskListReport
No Google access needed.
"""
import io
import json
import os
import re
import sys
import tempfile
import types
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.modules["_bootstrap"] = types.ModuleType("_bootstrap")
paths = types.ModuleType("paths")
paths.CREDENTIALS_DIR = paths.CONFIG_DIR = paths.LOGS_DIR = paths.CACHE_DIR = HERE
sys.modules["paths"] = paths
ec = types.ModuleType("email_config"); ec.GMAIL_SENDER = "x@x"; ec.GMAIL_APP_PASS = "x"
sys.modules["email_config"] = ec
for p in ("common", "ops_reports_action", "co-ordinator reports"):
    sys.path.insert(0, os.path.join(ROOT, p))

import openpyxl                                          # noqa: E402
from openpyxl.utils import get_column_letter             # noqa: E402
import api_retry                                         # noqa: E402
import coordinator_periods as CP                         # noqa: E402
import pyCoordinatorTaskListReport as BC                 # noqa: E402

api_retry._sleep = lambda s: None                        # no real waiting in tests
FAIL = []


def check(label, got, want):
    ok = got == want
    print(f"[{'pass' if ok else 'FAIL'}] {label}: {got!r}" + ("" if ok else f"  (expected {want!r})"))
    if not ok:
        FAIL.append(label)


# fake service-account file (only client_email is read; no secret in it)
SA = "pipeline-test@example-project.iam.gserviceaccount.com"
_sa_file = os.path.join(tempfile.mkdtemp(prefix="sa_"), "sa.json")
with open(_sa_file, "w") as fh:
    json.dump({"client_email": SA}, fh)
BC.AR.SERVICE_ACCOUNT_FILE = _sa_file
OWNER, COORD = BC.IMPERSONATE_USER, BC.COORDINATOR_EDITOR
OTHER_EDITOR = "someone-else@example.com"


# =============================================================================
#  In-memory Google Drive + Sheets
# =============================================================================
class _Req:
    def __init__(self, fn):
        self.fn = fn

    def execute(self):
        return self.fn()


class Transient(Exception):
    """Looks like a 503 to api_retry.is_transient."""
    def __init__(self):
        super().__init__("503 backendError")
        self.resp = types.SimpleNamespace(status=503)
        self.status_code = 503


EVENTS = []          # ordered log of every mutating call: (kind, file id, detail)


class FakeDrive:
    FOLDER = "application/vnd.google-apps.folder"

    def __init__(self, root, sheets):
        self.items = {root: {"id": root, "name": "ROOT", "mimeType": self.FOLDER, "parents": []}}
        self.sheets, self.n = sheets, 0
        self.perms, self.props = {}, {}
        self.fail_permission = False

    def files(self):
        return self

    def permissions(self):
        return _Perms(self)

    def list(self, q=None, **kw):
        def match(it):
            for cl in [c.strip() for c in q.split(" and ")]:
                m = re.match(r"^'(.+)' in parents$", cl)
                if m and m.group(1) not in it["parents"]:
                    return False
                m = re.match(r"^name\s*=\s*'(.*)'$", cl)
                if m and it["name"] != m.group(1).replace("\\'", "'"):
                    return False
                m = re.match(r"^name contains '(.*)'$", cl)
                if m and m.group(1).replace("\\'", "'") not in it["name"]:
                    return False
                m = re.match(r"^mimeType\s*=\s*'(.*)'$", cl)
                if m and it["mimeType"] != m.group(1):
                    return False
            return True
        return _Req(lambda: {"files": [dict(it) for it in self.items.values() if match(it)]})

    def create(self, body=None, media_body=None, **kw):
        def go():
            self.n += 1
            fid = f"f{self.n}"
            self.items[fid] = {"id": fid, "name": body["name"], "mimeType": body["mimeType"],
                               "parents": list(body.get("parents", []))}
            if media_body is not None:            # xlsx converted to a native Google Sheet
                self.sheets.load(fid, media_body.getbytes(0, media_body.size()))
                self.perms[fid] = {OWNER: "owner", SA: "writer", OTHER_EDITOR: "writer",
                                   COORD: "reader"}        # inherited from the folder
                EVENTS.append(("create", fid, body["name"]))
            return {"id": fid, "webViewLink": f"https://docs.google.com/spreadsheets/d/{fid}/edit"}
        return _Req(go)

    def update(self, fileId=None, body=None, **kw):
        def go():
            self.props.setdefault(fileId, {}).update(body or {})
            EVENTS.append(("drive.update", fileId, dict(body or {})))
            return {"id": fileId}
        return _Req(go)


class _Perms:
    def __init__(self, d):
        self.d = d

    def create(self, fileId=None, body=None, sendNotificationEmail=True, **kw):
        def go():
            if self.d.fail_permission:
                raise RuntimeError("403 permission create refused")
            self.d.perms[fileId][body["emailAddress"]] = body["role"]
            EVENTS.append(("share", fileId, (body["emailAddress"], body["role"],
                                             sendNotificationEmail)))
            return {"id": "p1"}
        return _Req(go)


class FakeSheets:
    """Native Google Sheets with protected ranges, enforced like Google does."""

    def __init__(self):
        self.books, self.pid = {}, 100
        self.fail_batch = 0          # number of batchUpdates to fail (transient)
        self.lose_response = 0       # batchUpdates that APPLY but raise afterwards
        self.hard_fail_batch = False
        self.corrupt_readback = False

    def load(self, fid, xbytes):
        wb = openpyxl.load_workbook(io.BytesIO(xbytes), data_only=False)
        tabs = []
        for i, ws in enumerate(wb.worksheets):
            tabs.append({"sheetId": 1000 + i if i else 0, "title": ws.title,
                         "rows": ws.max_row, "cols": ws.max_column,
                         "cells": {(c.row - 1, c.column - 1): c.value
                                   for row in ws.iter_rows() for c in row if c.value is not None},
                         "protected": []})
        self.books[fid] = tabs

    def spreadsheets(self):
        return self

    def get(self, spreadsheetId=None, fields=None, **kw):
        def go():
            out = []
            for t in self.books[spreadsheetId]:
                prs = [json.loads(json.dumps(p)) for p in t["protected"]]
                if self.corrupt_readback:
                    for p in prs:
                        p.pop("unprotectedRanges", None)
                out.append({"properties": {"sheetId": t["sheetId"], "title": t["title"]},
                            "protectedRanges": prs})
            return {"sheets": out}
        return _Req(go)

    def batchUpdate(self, spreadsheetId=None, body=None):
        def go():
            if self.hard_fail_batch:
                raise RuntimeError("400 Invalid requests[0]")
            if self.fail_batch:
                self.fail_batch -= 1
                raise Transient()
            tabs = [dict(t, protected=[dict(p) for p in t["protected"]])   # all-or-nothing
                    for t in self.books[spreadsheetId]]
            by_id = {t["sheetId"]: t for t in tabs}
            for rq in body["requests"]:
                if "deleteProtectedRange" in rq:
                    pid = rq["deleteProtectedRange"]["protectedRangeId"]
                    for t in tabs:
                        t["protected"] = [p for p in t["protected"] if p["protectedRangeId"] != pid]
                elif "addProtectedRange" in rq:
                    pr = dict(rq["addProtectedRange"]["protectedRange"])
                    gid = pr["range"]["sheetId"]
                    for u in pr.get("unprotectedRanges", []):
                        assert u["sheetId"] == gid
                    self.pid += 1
                    pr["protectedRangeId"] = self.pid
                    by_id[gid]["protected"].append(pr)
                else:
                    raise AssertionError(f"unexpected request {rq}")
            self.books[spreadsheetId] = tabs
            EVENTS.append(("protect", spreadsheetId, len(body["requests"])))
            if self.lose_response:
                self.lose_response -= 1
                raise Transient()
            return {"replies": []}
        return _Req(go)

    # ── Google's enforcement ────────────────────────────────────────────────
    @staticmethod
    def _in(rng, r, c, t):
        return (rng.get("startRowIndex", 0) <= r < rng.get("endRowIndex", t["rows"] + 10**6)
                and rng.get("startColumnIndex", 0) <= c < rng.get("endColumnIndex", t["cols"] + 10**6))

    def _tab(self, fid, title):
        return next(t for t in self.books[fid] if t["title"] == title)

    def can_edit_cell(self, fid, user, title, r, c, drive):
        role = drive.perms[fid].get(user)
        if role not in ("owner", "writer"):
            return False
        if role == "owner":
            return True
        t = self._tab(fid, title)
        for p in t["protected"]:
            if p.get("warningOnly"):
                continue
            if self._in(p["range"], r, c, t) and \
                    not any(self._in(u, r, c, t) for u in p.get("unprotectedRanges", [])) and \
                    user not in p["editors"]["users"]:
                return False
        return True

    def can_change_structure(self, fid, user, title, drive):
        """insert/delete rows or columns, rename/delete the tab: blocked for a
        non-editor whenever the WHOLE tab is protected (the exceptions don't help)."""
        role = drive.perms[fid].get(user)
        if role not in ("owner", "writer"):
            return False
        if role == "owner":
            return True
        t = self._tab(fid, title)
        return not any(set(p["range"]) == {"sheetId"} and not p.get("warningOnly")
                       and user not in p["editors"]["users"] for p in t["protected"])


sheets = FakeSheets()
drive = FakeDrive(BC.PARENT_FOLDER_ID, sheets)
BC._drive_client = lambda: drive
BC._enable_followup_timestamps = lambda sid: None
_real_protect = BC.protect_and_share
BC.protect_and_share = lambda sid, xb, d=None, s=None: _real_protect(sid, xb, d, s or sheets)


# =============================================================================
#  REAL Coordinator task workbooks
# =============================================================================
def adm(name):
    return {"Student Name": name, "Email ID": name.split()[0].lower() + "@x.com",
            "Phone Number": "+919000000001", "Batch Name": "DAAI1026", "Joined On": "2026-10-01",
            "Request Form Name": "", "Recipient Status": "Form Not Sent", "Request Status": "",
            "Sent Date": "", "Signed Date": "", "Expiry Date": ""}


def asg(name):
    return {"student_name": name, "student_email": name.lower().replace(" ", ".") + "@x.com",
            "student_phone": "+919000000002", "deadline_str": "08-Oct-2026", "reminder_level": "1st",
            "reminder_label": "1st Reminder", "assigned_date_str": "04-Oct-2026",
            "class_subject": "01-Oct-2026 To Current Date", "maximum_marks": ""}


def task_workbook(day):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    BC.build_admission_formalities(wb.create_sheet("Learner Admission Formalities"),
                                   [adm("Asha Rao"), adm("Bala K"), adm("Chitra M")], day)
    BC.build_assignment_followups(wb.create_sheet("Learner Assignment Follow-Ups"),
                                  {"Power BI": [("Assignment 1", "id-1", [asg("Dev N"), asg("Esha P")]),
                                                ("Assignment 2", "id-2", [asg("Farah Q")])]}, day)
    st = lambda n, tag: {"Student Name": n, "Batch Name": "DAAI1026", "Joined On": "",
                         "Student Name Status": "Valid", "Email ID Status": "Valid",
                         "Phone Number Status": "Valid", "Tag Name Status": tag,
                         "Private Note Status": "Valid", "Profile Picture Status": "Valid"}
    BC.build_wise_validation(wb.create_sheet("Wise & Interview Feedback Validation"),
                             {"student": [st("Gita R", "Missing"), st("Hari S", "Missing")],
                              "course": [], "instructor": []}, day)
    ws = wb.create_sheet("Read Me")                         # a tab with no follow-ups
    ws["A1"] = "Information only"
    buf = io.BytesIO()
    wb.save(buf)
    return buf


def header_map(ws):
    """{(row, col): header text without ✎} for the follow-up header cells."""
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("✎ "):
                out[(c.row, c.column)] = c.value[2:]
    return out


def expected_cells(xbytes):
    """Ground truth from the workbook: the input cells of every follow-up row
    (rows that have the Follow-Up DateTime formula), found by HEADER NAME —
    independently of the code under test."""
    wb = openpyxl.load_workbook(io.BytesIO(xbytes))
    want, dt_cells, hdr_cells = {}, {}, {}
    for ws in wb.worksheets:
        hm = header_map(ws)
        hdr_cells[ws.title] = set(hm)
        cols_by_name = {}
        for (r, c), h in hm.items():
            cols_by_name.setdefault(h, set()).add(c)
        dt_cols = cols_by_name.get("Follow-Up DateTime", set())
        cells, dts = set(), set()
        for row in ws.iter_rows():
            for c in row:
                if c.column in dt_cols and isinstance(c.value, str) and "NOW()" in c.value:
                    dts.add((c.row, c.column))
                    for name in BC.COORDINATOR_EDITABLE:
                        col = c.column - 3 + BC.COORDINATOR_EDITABLE.index(name)
                        assert col in cols_by_name[name], (ws.title, name, col)
                        cells.add((c.row, col))
        want[ws.title], dt_cells[ws.title] = cells, dts
    return want, dt_cells, hdr_cells


# =============================================================================
print("\n== 1. Coordinator-editable cells are found from the workbook ==")
D = date(2026, 10, 1)
xb = task_workbook(D).getvalue()
want, dt_cells, hdr_cells = expected_cells(xb)
rng = BC.followup_input_ranges(xb)
got = {t: {(r0 + 1 + i, c0 + 1 + j) for r0, r1, c0, c1 in rs
           for i in range(r1 - r0) for j in range(c1 - c0)} for t, rs in rng.items()}
check("every tab returned", sorted(rng), sorted(want))
for t in want:
    check(f"[{t}] editable cells = Action Taken / Comment / Done? of each follow-up row",
          got[t] == want[t], True)
check("follow-up rows on the task tabs", {t: len(v) // 3 for t, v in want.items()},
      {"Learner Admission Formalities": 3, "Learner Assignment Follow-Ups": 3,
       "Wise & Interview Feedback Validation": 2, "Read Me": 0})
check("Follow-Up DateTime never editable", any(dt_cells[t] & got[t] for t in want), False)
check("header cells never editable", any(hdr_cells[t] & got[t] for t in want), False)
check("contiguous rows merged into one range per block",
      all(len(rs) <= 3 for rs in rng.values()), True)
# the dropdowns (Action Taken / Done?) sit exactly on editable cells
wb = openpyxl.load_workbook(io.BytesIO(xb))
dv_ok = True
for ws in wb.worksheets:
    for dv in ws.data_validations.dataValidation:
        for cr in str(dv.sqref).split():
            from openpyxl.utils.cell import range_boundaries
            c0, r0, c1, r1 = range_boundaries(cr)
            for r in range(r0, r1 + 1):
                for c in range(c0, c1 + 1):
                    dv_ok &= (r, c) in got[ws.title]
check("every dropdown cell (Action Taken, Done?) is editable", dv_ok, True)


# =============================================================================
print("\n== 2. upload_report → protected, then shared ==")
EVENTS.clear()
link = BC.upload_datewise(D, f"{BC.REPORT_BASENAME}_30-Sep-2026_12.00_PM_-_01-Oct-2026_09.30_AM.xlsx",
                          task_workbook(D))
fid = link.split("/d/")[1].split("/")[0]
check("upload still returns the sheet link", link.startswith("https://docs.google.com/spreadsheets/d/"), True)
check("order: create → protect → no re-sharing → Coordinator writer",
      [e[0] for e in EVENTS], ["create", "protect", "drive.update", "share"])
check("editors cannot re-share", drive.props[fid], {"writersCanShare": False})
check("Coordinator share: writer, no notification e-mail", EVENTS[-1][2], (COORD, "writer", False))
check("Coordinator role on the file", drive.perms[fid][COORD], "writer")
check("owner kept", drive.perms[fid][OWNER], "owner")
tabs = sheets.books[fid]
check("one non-warning protection covering each whole tab",
      [(len(t["protected"]), t["protected"][0]["range"], t["protected"][0]["warningOnly"])
       for t in tabs], [(1, {"sheetId": t["sheetId"]}, False) for t in tabs])
check("protection editors = owner + service account only",
      {tuple(t["protected"][0]["editors"]["users"]) for t in tabs}, {(OWNER, SA)})
check("domain users cannot edit", {t["protected"][0]["editors"]["domainUsersCanEdit"] for t in tabs}, {False})

print("\n== 3. What each account can do on the uploaded sheet ==")
bad = []
for t in tabs:
    title = t["title"]
    for r in range(0, t["rows"] + 2):
        for c in range(0, t["cols"] + 2):
            allowed = (r + 1, c + 1) in want[title]
            if sheets.can_edit_cell(fid, COORD, title, r, c, drive) != allowed:
                bad.append((title, r + 1, get_column_letter(c + 1)))
check("Coordinator: exactly the three input columns, nothing else (every cell checked)", bad[:5], [])
check("Coordinator: Follow-Up DateTime cells locked",
      any(sheets.can_edit_cell(fid, COORD, t, r - 1, c - 1, drive) for t in dt_cells for r, c in dt_cells[t]), False)
check("Coordinator: header cells locked",
      any(sheets.can_edit_cell(fid, COORD, t, r - 1, c - 1, drive) for t in hdr_cells for r, c in hdr_cells[t]), False)
for who in (OWNER, SA):
    check(f"{who.split('@')[0]}: every cell editable",
          all(sheets.can_edit_cell(fid, who, t["title"], r, c, drive)
              for t in tabs for r in range(t["rows"]) for c in range(t["cols"])), True)
check("Coordinator: cannot insert/delete rows or columns, rename or delete any tab",
      [t["title"] for t in tabs if sheets.can_change_structure(fid, COORD, t["title"], drive)], [])
check("owner + service account: structure changes allowed",
      all(sheets.can_change_structure(fid, w, t["title"], drive) for w in (OWNER, SA) for t in tabs), True)
check("another folder editor is limited exactly like the Coordinator",
      sum(sheets.can_edit_cell(fid, OTHER_EDITOR, t["title"], r, c, drive)
          for t in tabs for r in range(t["rows"]) for c in range(t["cols"])),
      sum(len(v) for v in want.values()))

print("\n== 4. Every uploaded sheet gets it (re-run = Version 2, Weekly roll-up) ==")
EVENTS.clear()
link2 = BC.upload_datewise(D, f"{BC.REPORT_BASENAME}_30-Sep-2026_12.00_PM_-_01-Oct-2026_09.30_AM.xlsx",
                           task_workbook(D))
fid2 = link2.split("/d/")[1].split("/")[0]
check("re-run saved as Version 2", drive.items[fid2]["name"].endswith("- Version 2"), True)
check("Version 2 protected then shared", [e[0] for e in EVENTS], ["create", "protect", "drive.update", "share"])
check("Version 1 untouched by the re-run", [e for e in EVENTS if e[1] == fid], [])
EVENTS.clear()
wk = CP.folder_path("Weekly", date(2026, 9, 28), date(2026, 10, 4))
BC.upload_report(wk, "IntelliBI_Batch_Coordinator_Weekly_Follow_Ups_28_Sep_to_04_Oct_2026.xlsx",
                 task_workbook(D), "IntelliBI_Batch_Coordinator_Weekly_Follow_Ups")
check("Weekly roll-up protected then shared", [e[0] for e in EVENTS], ["create", "protect", "drive.update", "share"])

print("\n== 5. Re-applying is idempotent; transient errors retry without duplicates ==")
check("re-apply on the same sheet", _real_protect(fid, xb, drive, sheets), True)
check("still exactly one protection per tab", [len(t["protected"]) for t in sheets.books[fid]],
      [1] * len(sheets.books[fid]))
EVENTS.clear()
sheets.fail_batch = 1
BC.upload_datewise(D, "x_retry.xlsx", task_workbook(D))
check("transient 503 → retried, then protected and shared",
      [e[0] for e in EVENTS], ["create", "protect", "drive.update", "share"])
EVENTS.clear()
sheets.lose_response = 1                     # applied server-side, response lost, retried
l3 = BC.upload_datewise(D, "x_lost.xlsx", task_workbook(D))
f3 = l3.split("/d/")[1].split("/")[0]
check("lost response + retry → no duplicate protection",
      [len(t["protected"]) for t in sheets.books[f3]], [1] * len(sheets.books[f3]))
check("…and shared once", [e[0] for e in EVENTS].count("share"), 1)

print("\n== 6. Failures never share for editing and never stop the run ==")
for label, setup, teardown in [
        ("batchUpdate rejected", lambda: setattr(sheets, "hard_fail_batch", True),
         lambda: setattr(sheets, "hard_fail_batch", False)),
        ("read-back mismatch", lambda: setattr(sheets, "corrupt_readback", True),
         lambda: setattr(sheets, "corrupt_readback", False))]:
    EVENTS.clear()
    setup()
    try:
        lk = BC.upload_datewise(D, f"x_{label}.xlsx", task_workbook(D))
        ok_run = lk.startswith("https://")
    except Exception as exc:                                   # noqa: BLE001
        ok_run = f"raised {exc}"
    teardown()
    f = lk.split("/d/")[1].split("/")[0]
    check(f"{label}: upload still returns its link", ok_run, True)
    check(f"{label}: Coordinator NOT given edit access", drive.perms[f][COORD], "reader")
    check(f"{label}: no share call", [e[0] for e in EVENTS if e[0] == "share"], [])
BC.PROTECT_SHEETS = False
EVENTS.clear()
lk = BC.upload_datewise(D, "x_switch_off.xlsx", task_workbook(D))
f = lk.split("/d/")[1].split("/")[0]
check("PROTECT_SHEETS = False: upload only (no protection, no share)", [e[0] for e in EVENTS], ["create"])
BC.PROTECT_SHEETS = True
check("unreadable workbook bytes: returns False", _real_protect(f, b"not a workbook", drive, sheets), False)
check("unreadable workbook bytes: nothing protected or shared",
      ([e[0] for e in EVENTS if e[1] == f], drive.perms[f][COORD]), (["create"], "reader"))
drive.fail_permission = True
EVENTS.clear()
lk = BC.upload_datewise(D, "x_share_refused.xlsx", task_workbook(D))
drive.fail_permission = False
f = lk.split("/d/")[1].split("/")[0]
check("share refused: sheet stays protected, Coordinator not an editor, run continues",
      ([len(t["protected"]) for t in sheets.books[f]] == [1] * len(sheets.books[f]),
       drive.perms[f][COORD]), (True, "reader"))


print("\n== 7. Run summary lines ==")
import contextlib                                                     # noqa: E402
import exec_summary                                                   # noqa: E402
out = io.StringIO()
with contextlib.redirect_stdout(out):
    _real_protect(fid, xb, drive, sheets)
    sheets.hard_fail_batch = True
    _real_protect(fid, xb, drive, sheets)
    sheets.hard_fail_batch = False
txt = out.getvalue()
summ = exec_summary.summarize("pyCoordinatorTaskListReport", txt)
check("summary counts protected sheets", dict(summ["kpis"]).get("Sheets protected & shared"), 1)
check("summary flags a failed protection", "could NOT be protected" in (summ["note"] or ""), True)


print("\n== 8. Rename to pyCoordinatorTaskListReport (formerly the Batch Coordinator Daily Attendance script) ==")
OLD = "pyBatchCoordinator" + "DailyAttendanceReport"
check("new script exists", os.path.isfile(os.path.join(ROOT, "co-ordinator reports",
                                                       "pyCoordinatorTaskListReport.py")), True)
check("old script gone", os.path.exists(os.path.join(ROOT, "co-ordinator reports", OLD + ".py")), False)
left = []
for base, dirs, files in os.walk(ROOT):
    dirs[:] = [d for d in dirs if d not in (".git", "__pycache__", "logs", "_to_delete", ".venv",
                                            "venv", "cache", "credentials")]
    for fn in files:
        if not fn.endswith((".py", ".md", ".ps1", ".bat", ".yaml", ".yml", ".txt", ".json", ".cfg")):
            continue
        fp = os.path.join(base, fn)
        try:
            for n, line in enumerate(open(fp, encoding="utf-8", errors="ignore"), 1):
                if OLD in line and "formerly" not in line.lower() and "renamed" not in line.lower():
                    left.append(f"{os.path.relpath(fp, ROOT)}:{n}")
        except OSError:
            pass
check("no reference to the old name left (except 'formerly …' notes)", left, [])
import importlib                                                      # noqa: E402
rra_src = open(os.path.join(ROOT, "scripts", "run_reports_action.py"), encoding="utf-8").read()
check("scheduled morning batch runs the renamed script", "pyCoordinatorTaskListReport" in rra_src, True)
perf_src = open(os.path.join(ROOT, "co-ordinator reports", "pyCoordinatorTaskPerformanceReport.py"),
                encoding="utf-8").read()
check("performance report imports the renamed module",
      "import pyCoordinatorTaskListReport as BC" in perf_src, True)
check("Drive report name unchanged (history & versioning keep working)",
      BC.REPORT_BASENAME, "IntelliBI_Batch_Coordinator_Daily_Attendance_Report")
check("logger renamed", BC.log.name, "CoordinatorTaskList")

print()
if FAIL:
    print(f"FAILED: {len(FAIL)} check(s): " + "; ".join(FAIL))
    sys.exit(1)
print("ALL CHECKS PASSED")
