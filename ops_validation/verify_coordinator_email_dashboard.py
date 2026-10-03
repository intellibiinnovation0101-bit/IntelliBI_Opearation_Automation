"""
Verification of the Coordinator report e-mails and the redesigned Task
Performance Dashboard (run from the project root:
  python ops_validation\\verify_coordinator_email_dashboard.py).

No Google / SMTP access: Gmail SMTP is replaced by an in-memory recorder, so
nothing is really sent. Checks, for BOTH reports:
  * SEND_EMAIL = True  -> report generated AND one e-mail from
    info@intellibiinnovationstechnologies.in to the two recipients
  * SEND_EMAIL = False -> report generated exactly the same, NO e-mail
  * a sender that is not the configured Gmail account is refused
  * reports are kept in Google Drive ONLY: no local copy, no temporary file,
    nothing written under output/; old local copies of the Task Performance
    report are removed (other files kept); the Batch Coordinator script has
    no file-writing code (built and uploaded from memory)
and the Dashboard: Missed tile, ASSESSMENT and the nine scorecard columns are
gone; 7-column grid; no overlapping merges; chart reads existing columns;
print area set; other tabs unchanged.
"""
import io
import re
import os
import sys
import types
import shutil
import tempfile
from datetime import date, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.modules.setdefault("_bootstrap", types.ModuleType("_bootstrap"))
ec = types.ModuleType("email_config")
ec.GMAIL_SENDER = "info@intellibiinnovationstechnologies.in"
ec.GMAIL_APP_PASS = "test-only"
sys.modules["email_config"] = ec                     # never the real credentials
for p in ("common", "ops_reports_action", "co-ordinator reports", "scripts"):
    sys.path.insert(0, os.path.join(ROOT, p))

import smtplib                                        # noqa: E402
import openpyxl                                       # noqa: E402
import pyCoordinatorTaskListReport as BC  # noqa: E402
import pyCoordinatorTaskPerformanceReport as P        # noqa: E402

FAIL = []
SENT = []
RECIPIENTS = ["info@intellibiinnovationstechnologies.in", "intellibihropsb2ch@gmail.com"]


def mail_text(m):
    """Decoded subject + bodies of a captured message."""
    import email
    from email.header import decode_header, make_header
    msg = email.message_from_string(m.get("msg", ""))
    parts = [str(make_header(decode_header(msg.get("Subject", ""))))]
    for part in msg.walk():
        if part.get_content_maintype() == "text":
            parts.append(part.get_payload(decode=True).decode("utf-8", "replace"))
    return "\n".join(parts)


def check(label, got, want):
    ok = got == want
    print(f"[{'pass' if ok else 'FAIL'}] {label}: {got!r}" + ("" if ok else f"  (expected {want!r})"))
    if not ok:
        FAIL.append(label)


class FakeSMTP:
    def __init__(self, host, port, timeout=None):
        self.host = host

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def login(self, user, pw):
        self.user = user

    def sendmail(self, sender, recipients, msg):
        SENT.append({"login": self.user, "from": sender, "to": list(recipients), "msg": msg})


smtplib.SMTP_SSL = FakeSMTP                           # nothing leaves this process

# =============================================================================
print("\n== 1. Configuration ==")
for mod, name in ((BC, "Batch Coordinator"), (P, "Task Performance")):
    check(f"{name}: SEND_EMAIL default True", mod.SEND_EMAIL, True)
    check(f"{name}: sender", mod.EMAIL_SENDER, "info@intellibiinnovationstechnologies.in")
    check(f"{name}: recipients", mod.EMAIL_RECIPIENTS, RECIPIENTS)

# =============================================================================
#  Synthetic Coordinator task report (real builders) → performance report
# =============================================================================
D = date(2026, 10, 5)


def adm(name):
    return {"Student Name": name, "Email ID": name.split()[0].lower() + "@x.com",
            "Phone Number": "+919000000001", "Batch Name": "DAAI1026", "Joined On": "2026-10-01",
            "Request Form Name": "", "Recipient Status": "Form Not Sent", "Request Status": "",
            "Sent Date": "", "Signed Date": "", "Expiry Date": ""}


wb = openpyxl.Workbook()
wb.remove(wb.active)
BC.build_admission_formalities(wb.create_sheet("Learner Admission Formalities"),
                               [adm("Asha Rao"), adm("Bala K"), adm("Chitra M")], D)
ws = wb.worksheets[0]
hdr = None
for r in range(1, ws.max_row + 1):
    txt = [str(ws.cell(row=r, column=c).value or "") for c in range(1, ws.max_column + 1)]
    if "Why Flagged" in txt:
        hdr = {t.replace("✎", "").strip(): i + 1 for i, t in enumerate(txt) if t}
        continue
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if isinstance(v, str) and v.startswith("=IF("):
            ws.cell(row=r, column=c).value = None
    if hdr and txt[0] == "Asha Rao":
        ws.cell(row=r, column=hdr["Follow-Up Done?"]).value = "Yes"
        ws.cell(row=r, column=hdr["Follow-Up DateTime"]).value = datetime(2026, 10, 5, 11, 0)
        ws.cell(row=r, column=hdr["Action Taken"]).value = "Call"
buf = io.BytesIO()
wb.save(buf)
versions = [{"id": "v1", "name": "v1", "day": D, "version": 1,
             "created": datetime(2026, 10, 5, 9, 0), "modified_raw": "", "folder": "Daily 05-Oct-2026"}]
parsed = P.parse_report_workbook(openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=True))

# Run the REAL generate() with Drive replaced: discovery → our version, and the
# upload captured in memory (exactly what would reach Google Drive).
UPLOADED = {}


def _fake_upload(folder_name, filename, buf, base_prefix):
    UPLOADED[filename] = buf.getvalue()
    return f"https://docs.google.com/spreadsheets/d/{base_prefix}"


def uploaded_wb(r):
    return openpyxl.load_workbook(io.BytesIO(UPLOADED[r["name"] + ".xlsx"]))


def tree(path):
    """Every file under `path` (to prove the run writes nothing locally)."""
    return sorted(os.path.join(d, f) for d, _s, fs in os.walk(path) for f in fs)


# earlier versions kept local copies: seed a fake legacy folder to prove cleanup
legacy = tempfile.mkdtemp(prefix="coord_legacy_")
_old_dir = os.path.join(legacy, "Daily Coordinator Reports", "Daily 01-Oct-2026")
os.makedirs(_old_dir)
for _f in (os.path.join(_old_dir, P.REPORT_BASENAME + "_Daily_01_Oct_2026.xlsx"),
           os.path.join(legacy, P.REPORT_BASENAME + "_Daily_01_Oct_2026.xlsx")):
    open(_f, "wb").write(b"old")
_keep_dir = os.path.join(legacy, "Notes")
os.makedirs(_keep_dir)
open(os.path.join(_keep_dir, "someone_elses_file.txt"), "w").write("keep me")

P._drive_service = lambda: None
P.discover_report_versions = lambda drive, ranges: versions
P.load_version = lambda drive, v: parsed
P._now_ist = lambda: datetime(2026, 10, 5, 19, 0)
P.BC.upload_report = _fake_upload
P.UPLOAD_TO_DRIVE = True
P.LEGACY_OUTPUT_DIR = legacy
P.GENERATE_AUTO = True
P.SEND_EMAIL = True
P.STAR_EMAIL_IN_GMAIL = False      # starring has its own test: verify_gmail_star.py
_out_before = tree(os.path.join(ROOT, "output"))
_tmp_before = set(os.listdir(tempfile.gettempdir()))

print("\n== 2. Task Performance report ==")
SENT.clear()
res, errs = P.generate()
check("SEND_EMAIL=True: report generated and uploaded to Drive",
      [(r["name"] + ".xlsx" in UPLOADED, bool(r.get("link"))) for r in res], [(True, True)] * 2)
check("no local report copy: result carries no local path", [r.get("local") for r in res], [None, None])
check("no file written under the project's output/ folder", tree(os.path.join(ROOT, "output")), _out_before)
check("no temporary file left in the temp folder",
      sorted(x for x in set(os.listdir(tempfile.gettempdir())) - _tmp_before
             if x.endswith(".xlsx") or x.startswith(P.REPORT_BASENAME)), [])
check("old local copies removed, empty period folders removed, other files kept",
      [os.path.relpath(f, legacy).replace(os.sep, "/") for f in tree(legacy)]
      + sorted(os.path.relpath(d, legacy).replace(os.sep, "/") for d, _s, _f in os.walk(legacy)),
      ["Notes/someone_elses_file.txt", ".", "Notes"])
check("… a second run finds nothing more to remove", P.cleanup_legacy_local_copies(), 0)
shutil.rmtree(legacy, ignore_errors=True)
check("… and copes with the legacy folder not existing", P.cleanup_legacy_local_copies(), 0)
check("SEND_EMAIL=True: one e-mail per report (Monday = Daily + Weekly)", len(SENT), 2)
m = SENT[0] if SENT else {}
check("… logged in and sent as info@", (m.get("login"), m.get("from")),
      ("info@intellibiinnovationstechnologies.in",) * 2)
check("… to both recipients", m.get("to"), RECIPIENTS)
check("… subjects", [mail_text(x).split("\n")[0] for x in SENT],
      ["Daily Coordinator Task Performance Report - 05-Oct-2026",
       "Weekly Coordinator Task Performance Report - 28-Sep-2026 to 04-Oct-2026"])
_body = mail_text(m)
check("… body: header, greeting, KPI cards, goal bars, task groups, button, signature",
      all(x in _body for x in ("Reporting Period: 05-Oct-2026", "Hello Team,", "Tasks Generated",
                               "Timely<br>Completion %", "Performance vs Goals", "Task Groups",
                               "Learner Admission Formalities", "IntelliBI Automation Team",
                               "Automated report &middot; Generated")), True)

# ---------------------------------------------------------------------------
print("\n== 2b. Performance vs Goals (e-mail) ==")
import coordinator_email as CE                        # noqa: E402
check("benchmark = 95% task completion", P.COMPLETION_BENCHMARK, 95.0)
check("e-mail display names (as specified)", P.EMAIL_GROUP_LABELS, {
    "Learner Attendance Follow-Ups": "Learner Attendance",
    "Learner Assignment Follow-Ups": "Learner Assignment",
    "Learner Admission Formalities": "Learner Admission Formalities",
    "Wise & Interview Feedback Validation": "Learner Admission Formalities",
    "Instructor Follow-Ups": "Instructor Instructions",
    "Learner Instructor Interview Reminder": "Interview Reminder"})
check("every registered task group has an e-mail display name",
      [g["name"] for g in P.TASK_GROUPS if g["name"] not in P.EMAIL_GROUP_LABELS], [])
check("registry (task-group identity) names unchanged", [g["name"] for g in P.TASK_GROUPS],
      ["Learner Attendance Follow-Ups", "Learner Assignment Follow-Ups", "Learner Admission Formalities",
       "Wise & Interview Feedback Validation", "Instructor Follow-Ups",
       "Learner Instructor Interview Reminder"])
_goals = _body.split("Performance vs Goals", 1)[-1].split("Task Groups", 1)[0]
check("section shows the benchmark", "95% Task Completion" in _goals, True)
check("applicable group shown with its actual figures (1 of 3 done)",
      "Learner Admission Formalities" in _goals and "1 / 3 &middot; 33.3%" in _goals, True)
check("below benchmark → red bar, benchmark marker at 95%",
      (f"#{CE.RED}" in _goals, f"#{CE.GREEN};width" in _goals, "left:95%" in _goals), (True, False, True))
check("groups with no tasks in the period are not shown",
      [n for n in ("Learner Attendance", "Learner Assignment", "Instructor Instructions",
                   "Interview Reminder") if n in _goals], [])
# reconciliation: the e-mail figure IS the Dashboard scorecard figure
_dash = uploaded_wb(res[0])["Dashboard"]
_sc = {r[0]: r for r in _dash.iter_rows(values_only=True) if r and r[0] in P.EMAIL_GROUP_LABELS}
_rows = CE.goal_rows(res[0]["groups"], P.COMPLETION_BENCHMARK, P.EMAIL_GROUP_LABELS)
check("e-mail Completion % == Dashboard scorecard Completion % (every group shown)",
      [(name, done, n, pct) for _d, name, done, n, pct, _m in _rows],
      [(name, _sc[name][2], _sc[name][1], _sc[name][4]) for _d, name, *_x in _rows])
# Weekly = 28-Sep → 04-Oct: the 05-Oct tasks belong to the NEXT week, so no
# group is applicable to this period (strict period scoping, no leakage)
_wk = mail_text(SENT[1]).split("Performance vs Goals", 1)[-1].split("Task Groups", 1)[0]
check("Weekly e-mail: only its own period's groups (05-Oct tasks not leaked in)",
      ("No tasks were generated in this period." in _wk, "33.3%" in _wk), (True, False))


def _g(n, done):
    return {"tasks": n, "completed": done, "completion_pct": (100.0 * done / n) if n else None}


_unit = CE.goal_rows([("Learner Attendance Follow-Ups", _g(20, 19)),         # 95.0  → met
                      ("Learner Assignment Follow-Ups", _g(100000, 94940)),  # 94.94 → 94.9 → below
                      ("Learner Admission Formalities", _g(100000, 94960)),  # 94.96 → 95.0 → met
                      ("Wise & Interview Feedback Validation", _g(0, 0)),    # not applicable
                      ("Instructor Follow-Ups", _g(4, 4)),                   # 100   → met
                      ("Some Future Tab", _g(10, 1))],                       # unmapped → own name
                     95.0, P.EMAIL_GROUP_LABELS)
check("threshold: >= 95% green, < 95% red (judged on the shown 0.1% value)",
      [(d, p, met) for d, _n, _c, _t, p, met in _unit],
      [("Learner Attendance", 95.0, True), ("Learner Assignment", 94.9, False),
       ("Learner Admission Formalities", 95.0, True), ("Instructor Instructions", 100.0, True),
       ("Some Future Tab", 10.0, False)])
_html = CE.goals_section([("Learner Attendance Follow-Ups", _g(20, 19)),
                          ("Learner Assignment Follow-Ups", _g(10, 5))], 95.0, P.EMAIL_GROUP_LABELS)
check("bar colours: met = green, below = red",
      (_html.count(f"color:#{CE.GREEN};text-align:right"), _html.count(f"color:#{CE.RED};text-align:right")),
      (1, 1))
check("period with no tasks → a clear 'no tasks' line, no bars",
      "No tasks were generated in this period." in CE.goals_section(
          [("Learner Attendance Follow-Ups", _g(0, 0))], 95.0, P.EMAIL_GROUP_LABELS), True)

# ── "Overall Completion %" first ──────────────────────────────────────────
_titles = re.findall(r"<td style='font-size:13px;color:#1a2a48'>([^<]+)</td>", _goals)
check("first metric is 'Overall Completion %', then the task groups in their existing order",
      _titles, ["Overall Completion %"] + [d for d, *_r in _rows])
_ov = _goals.split("Overall Completion %", 1)[1].split("</div></div>", 1)[0]
_tot = [r for r in _dash.iter_rows(values_only=True) if r and r[0] == "All task groups"][0]
check("Overall figures = the report's own total (Dashboard 'All task groups': completed / tasks · %)",
      ("1 / 3 &middot; 33.3%" in _ov, (_tot[2], _tot[1], _tot[4])), (True, (1, 3, 33.3)))
check("Overall below 95% → red, with the 95% benchmark marker",
      (f"color:#{CE.RED};text-align:right" in _ov, f"#{CE.GREEN};width" in _ov, "left:95%" in _ov),
      (True, False, True))
_ovs = lambda n, d: {"tasks": n, "completed": d, "completion_pct": (100.0 * d / n) if n else None}
check("Overall threshold: >= 95% green, < 95% red (judged on the shown 0.1% value)",
      [CE.overall_goal_row(_ovs(*x), 95.0)[4:] for x in ((20, 19), (100000, 94960), (100000, 94940), (4, 4))],
      [(95.0, True), (95.0, True), (94.9, False), (100.0, True)])
_h2 = CE.goals_section([("Learner Attendance Follow-Ups", _g(20, 19)), ("Learner Assignment Follow-Ups", _g(10, 5))],
                       95.0, P.EMAIL_GROUP_LABELS, overall=_ovs(30, 24))
check("Overall reconciles with the groups below it (24 / 30 = 19+5 / 20+10) and is red at 80%",
      re.findall(r"text-align:right'>([^<]+)</td>", _h2)[:1] == ["24 / 30 &middot; 80.0%"]
      and f"color:#{CE.RED};text-align:right'>24 / 30" in _h2, True)
check("period with no tasks → no Overall bar either, just the 'no tasks' line",
      ("Overall Completion %" in _wk, "No tasks were generated in this period." in _wk), (False, True))

SENT.clear()
P.SEND_EMAIL = False
UPLOADED.clear()
res2, errs2 = P.generate()
check("SEND_EMAIL=False: report still generated and uploaded",
      [r["name"] + ".xlsx" in UPLOADED for r in res2], [True, True])
check("SEND_EMAIL=False: same report figures", res2[0]["summary"], res[0]["summary"])
check("SEND_EMAIL=False: no e-mail", len(SENT), 0)
check("SEND_EMAIL=False: still no local file", (tree(os.path.join(ROOT, "output")), [r.get("local") for r in res2]),
      (_out_before, [None, None]))
check("SEND_EMAIL=False: run still succeeds (Monday = Daily + Weekly)", [bool(r.get("failed")) for r in res2],
      [False, False])

P.SEND_EMAIL = True
P.EMAIL_SENDER = "someone.else@example.com"
SENT.clear()
check("sender other than the configured Gmail account is refused", P.email_results(res), False)
check("… and nothing is sent", len(SENT), 0)
P.EMAIL_SENDER = "info@intellibiinnovationstechnologies.in"
P.SEND_EMAIL = False

# =============================================================================
print("\n== 3. Dashboard layout ==")
wbk = uploaded_wb(res2[0])
dash = wbk["Dashboard"]
cells = [str(c.value) for row in dash.iter_rows() for c in row if c.value is not None]
check("tabs: Pending & Overdue removed, the rest unchanged", wbk.sheetnames,
      ["Dashboard", "Progress Trend", "Task Register", "Data Coverage & Rules"])
check("NEEDS MANAGEMENT ATTENTION section and its rows removed",
      [c for c in cells if "MANAGEMENT ATTENTION" in c.upper() or c in ("Act now", "Watch", "Info", "Data note")
       or "Pending & Overdue" in c], [])
check("'Missed (Overdue)' tile removed", "Missed (Overdue)" in cells, False)
check("ASSESSMENT section removed", "  ASSESSMENT" in cells or "ASSESSMENT" in cells, False)
hdr_row = [r for r in dash.iter_rows(values_only=True) if r and r[0] == "Task Group"][0]
check("scorecard columns", [h for h in hdr_row if h],
      ["Task Group", "Tasks", "Completed", "Pending", "Completion %", "Pending %", "Status"])
check("dashboard is 7 columns wide (no empty columns left behind)", dash.max_column, 7)
seen, overlap = set(), False
for mr in dash.merged_cells.ranges:
    for r in range(mr.min_row, mr.max_row + 1):
        for c in range(mr.min_col, mr.max_col + 1):
            overlap = overlap or (r, c) in seen
            seen.add((r, c))
check("no overlapping merged cells", overlap, False)
check("no merge wider than the grid", max(mr.max_col for mr in dash.merged_cells.ranges), 7)
_names = ["OVERALL COORDINATOR PERFORMANCE", "TASK GROUP SCORECARD", "NEEDS MANAGEMENT ATTENTION",
          "COMPLETED VS PENDING BY TASK GROUP", "ASSESSMENT"]
sections = [n for r in range(1, dash.max_row + 1)
            for n in _names if str(dash.cell(row=r, column=1).value or "").strip().startswith(n)]
check("section order", sections,
      ["OVERALL COORDINATOR PERFORMANCE", "TASK GROUP SCORECARD", "COMPLETED VS PENDING BY TASK GROUP"])
_tot = [i for i, r in enumerate(dash.iter_rows(values_only=True), 1) if r and r[0] == "All task groups"][0]
_chs = [i for i, r in enumerate(dash.iter_rows(values_only=True), 1)
        if r and str(r[0] or "").strip().startswith("COMPLETED VS PENDING")][0]
check("chart section follows the scorecard total row after one spacer row", _chs - _tot, 2)
charts = dash._charts
refs = [s.val.numRef.f for s in charts[0].series] if charts else []
hdr_r = [i for i, r in enumerate(dash.iter_rows(values_only=True), 1) if r and r[0] == "Task Group"][0]
check("chart reads the scorecard's Completed / Pending columns",
      [f.split("!")[1].split(":")[0][:2] for f in refs], ["$C", "$D"])
check("… starting under the scorecard header", refs[0].split("!")[1].split(":")[0], f"$C${hdr_r + 1}")
check("print area set to the 7-column grid", dash.print_area.split("!")[1].split(":")[0], "$A$1")
blank_runs = 0
run = 0
for r in range(1, dash.max_row + 1):
    empty = all(dash.cell(row=r, column=c).value in (None, "") for c in range(1, 8))
    run = run + 1 if empty else 0
    blank_runs = max(blank_runs, run)
check("no blank gap longer than one row between sections", blank_runs <= 1, True)
# every pending task (and what the old tab showed about it) is in the Task Register
_reg = wbk["Task Register"]
_rh = [c.value for c in _reg[3]]
_rr = [dict(zip(_rh, [c.value for c in r])) for r in _reg.iter_rows(min_row=4) if r[0].value]
_pend = [t for t in P.scope_ledger(P.build_ledger(versions, lambda v: parsed, datetime(2026, 10, 5, 19, 0)),
                                   D, D)["tasks"] if not t["completed"]]
check("Task Register columns unchanged", _rh,
      ["Report Day", "Task Group", "Item", "Context", "What Was Flagged", "Generated At", "Status",
       "Completed At", "Time to Complete", "Action Taken", "Follow-Up Comment", "Days on List",
       "Seen in Versions", "Recorded in Versions"])
check("every pending task is in the Task Register with its status (filterable)",
      sorted((d["Item"], d["Status"]) for d in _rr if d["Status"] in (P.ST_OPEN, P.ST_MISSED)),
      sorted((t["label"], t["status"]) for t in _pend))
check("… 'attempted' is visible there as 'Recorded in Versions' (same rule)",
      sorted((d["Item"], d["Recorded in Versions"] not in (None, "—")) for d in _rr
             if d["Status"] in (P.ST_OPEN, P.ST_MISSED)),
      sorted((t["label"], t["attempted"]) for t in _pend))
check("Task Register guide explains how to review pending tasks",
      "filter Status = Missed / Open" in str(_reg["A2"].value), True)
check("e-mail closing no longer mentions the removed parts",
      [x for x in ("management attention", "every pending task.") if x in _body], [])

# =============================================================================
print("\n== 4. Batch Coordinator report ==")
# Run the REAL generate() with the data load and the Google calls replaced.
import utils as _utils                                                    # noqa: E402
import pandas as pd                                                        # noqa: E402
_utils.get_sheets_service = lambda *a, **k: None
BC.AR.load_all_data = lambda *a, **k: tuple(pd.DataFrame() for _ in range(5))
BC.AR._ist_today = lambda: date(2026, 10, 5)
GEN = []


_bcwb = openpyxl.Workbook()
_bcwb.remove(_bcwb.active)
BC.build_admission_formalities(_bcwb.create_sheet("Learner Admission Formalities"),
                               [adm("Asha Rao"), adm("Bala K")], D)
BC.build_assignment_followups(_bcwb.create_sheet("Learner Assignment Follow-Ups"), {}, D)
check("task counts read from the tabs' own guide counts", BC._task_counts(_bcwb),
      {"Learner Admission Formalities": 2, "Learner Assignment Follow-Ups": 0})


def _fake_daily(service, report_date, *a):
    GEN.append(report_date)
    return {"link": "https://docs.google.com/spreadsheets/d/x", "report_date": report_date.isoformat(),
            "daily_rows": 21, "agg_groups": 5, "task_counts": BC._task_counts(_bcwb)}


def _fake_period(service, report_type, start, end, plabel, *a):
    GEN.append((report_type, start))
    return {"link": "https://docs.google.com/spreadsheets/d/w", "report_type": report_type,
            "period": plabel, "start": start.isoformat(), "end": end.isoformat()}


BC._generate_daily = _fake_daily
BC._generate_period = _fake_period
BC.GENERATE_AUTO = True
BC.DAILY_DATE = None               # Daily = the (patched) run date
BC.SEND_EMAIL = True
BC.STAR_EMAIL_IN_GMAIL = False     # starring has its own test: verify_gmail_star.py
SENT.clear()
r1 = BC.generate()
check("Batch Coordinator run writes nothing under output/", tree(os.path.join(ROOT, "output")), _out_before)
# the report builders and the upload work on an in-memory buffer only
import re as _re                                                           # noqa: E402
_src = open(BC.__file__, encoding="utf-8").read()
check("Batch Coordinator: every workbook save goes to an in-memory buffer",
      sorted(set(_re.findall(r"\bwb\.save\(([^)]*)\)", _src))), ["buf"])
check("Batch Coordinator: no code that writes report files to disk",
      _re.findall(r"open\([^)]*['\"][wa]b?['\"]|to_excel\(|MediaFileUpload|NamedTemporaryFile|mkstemp", _src), [])
check("SEND_EMAIL=True: Monday run generated Daily + Weekly", (len(GEN), [bool(x.get("failed")) for x in r1]),
      (2, [False, False]))
check("SEND_EMAIL=True: one e-mail per report, from info@, to both recipients",
      (len(SENT), {x["from"] for x in SENT}, [x["to"] for x in SENT]),
      (2, {"info@intellibiinnovationstechnologies.in"}, [RECIPIENTS, RECIPIENTS]))
check("… subjects", [mail_text(x).split("\n")[0] for x in SENT],
      ["Daily Batch Coordinator Report - 05-Oct-2026", "Weekly Batch Coordinator Report - 28 Sep – 04 Oct 2026"])
check("… each carries its Google Sheet link",
      ["spreadsheets/d/x" in mail_text(SENT[0]), "spreadsheets/d/w" in mail_text(SENT[1])], [True, True])
check("… daily e-mail shows the task count per tab",
      all(x in mail_text(SENT[0]) for x in ("Total Tasks", "Learner Admission Formalities")), True)
BC.SEND_EMAIL = False
SENT.clear()
r2 = BC.generate()
check("SEND_EMAIL=False: reports still generated", (len(GEN), [bool(x.get("failed")) for x in r2]),
      (4, [False, False]))
check("SEND_EMAIL=False: no e-mail", len(SENT), 0)
BC.SEND_EMAIL = True
SENT.clear()
check("failed report → nothing to e-mail", BC.email_results([{"failed": True}]), None)
check("… and nothing sent", len(SENT), 0)
BC.EMAIL_RECIPIENTS = ["info@intellibiinnovationstechnologies.in" "intellibihropsb2ch@gmail.com"]
check("fused recipients (missing comma) are refused, not mailed",
      (BC.email_results([{"link": "L1", "report_date": "2026-09-28"}]), len(SENT)), (False, 0))
BC.EMAIL_RECIPIENTS = list(RECIPIENTS)

print("\nALL CHECKS PASSED" if not FAIL else f"\n{len(FAIL)} CHECK(S) FAILED: {FAIL}")
sys.exit(1 if FAIL else 0)
