"""
End-to-end verification of the Coordinator reporting periods, Drive layout and
scheduling (run from the project root:
  python ops_validation\\verify_coordinator_reporting_periods.py).

No Google access: an in-memory Drive stands in for the real one, and the REAL
code paths are exercised —
  coordinator_periods.plan_jobs / folder_path          (period selection, folders)
  pyCoordinatorTaskListReport.upload_report (versioned upload)
  pyCoordinatorTaskPerformanceReport discover → load → ledger → scope → upload
  scripts/run_reports_action.JOBS / script_path, scripts/run_scheduled.py
Covers: Daily / Weekly / Monthly / Manual, AUTO and flag mode, Monday weekly,
month-end (30/31/28/29-day months), Manual validation, re-runs + versions, no
earlier-period leakage, both reports in the same period folder, scheduler entry,
labels and once-per-day gate.
"""
import io
import os
import re
import sys
import types
import shutil
import tempfile
from datetime import date, datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.modules.setdefault("_bootstrap", types.ModuleType("_bootstrap"))
if "paths" not in sys.modules:
    sys.path.insert(0, os.path.join(ROOT, "common"))
ec = types.ModuleType("email_config"); ec.GMAIL_SENDER = "x@x"; ec.GMAIL_APP_PASS = "x"
sys.modules.setdefault("email_config", ec)
for p in ("common", "ops_reports_action", "co-ordinator reports", "scripts"):
    sys.path.insert(0, os.path.join(ROOT, p))

import openpyxl                                          # noqa: E402
import coordinator_periods as CP                         # noqa: E402
import pyCoordinatorTaskListReport as BC     # noqa: E402
import pyCoordinatorTaskPerformanceReport as P           # noqa: E402

FAIL = []


def check(label, got, want):
    ok = got == want
    print(f"[{'pass' if ok else 'FAIL'}] {label}: {got!r}" + ("" if ok else f"  (expected {want!r})"))
    if not ok:
        FAIL.append(label)


def J(jobs):
    return [(j["kind"], j["start"].isoformat(), j["end"].isoformat()) for j in jobs]


# =============================================================================
print("\n== 1. Period selection (shared by both reports) ==")
auto = dict(auto=True, daily=True, weekly=True, monthly=True, manual=True,
            manual_start="2026-08-21", manual_end="2026-09-22")
check("AUTO, ordinary Thursday -> Daily only (flags ignored)",
      J(CP.plan_jobs(date(2026, 10, 1), **auto)[0]), [("Daily", "2026-10-01", "2026-10-01")])
check("AUTO, Monday 28-Sep-2026 -> Daily + previous Mon-Sun week",
      J(CP.plan_jobs(date(2026, 9, 28), **auto)[0]),
      [("Daily", "2026-09-28", "2026-09-28"), ("Weekly", "2026-09-21", "2026-09-27")])
check("AUTO, 30-Sep-2026 (30-day month end) -> Daily + Monthly 1-30 Sep",
      J(CP.plan_jobs(date(2026, 9, 30), **auto)[0]),
      [("Daily", "2026-09-30", "2026-09-30"), ("Monthly", "2026-09-01", "2026-09-30")])
check("AUTO, 31-Oct-2026 (31-day month end, Saturday)",
      J(CP.plan_jobs(date(2026, 10, 31), **auto)[0])[1], ("Monthly", "2026-10-01", "2026-10-31"))
check("AUTO, 30-Nov-2026 = Monday AND month end -> Daily + Weekly + Monthly",
      J(CP.plan_jobs(date(2026, 11, 30), **auto)[0]),
      [("Daily", "2026-11-30", "2026-11-30"), ("Weekly", "2026-11-23", "2026-11-29"),
       ("Monthly", "2026-11-01", "2026-11-30")])
check("AUTO, 28-Feb-2027 (non-leap) month end", J(CP.plan_jobs(date(2027, 2, 28), **auto)[0])[-1],
      ("Monthly", "2027-02-01", "2027-02-28"))
check("AUTO, 28-Feb-2028 (leap year) is NOT month end",
      [k for k, _a, _b in J(CP.plan_jobs(date(2028, 2, 28), **auto)[0])], ["Daily", "Weekly"])
check("AUTO, 29-Feb-2028 (leap) month end", J(CP.plan_jobs(date(2028, 2, 29), **auto)[0])[-1],
      ("Monthly", "2028-02-01", "2028-02-29"))
check("AUTO, Monday 04-Jan-2027 -> week spanning the year end",
      J(CP.plan_jobs(date(2027, 1, 4), **auto)[0])[1], ("Weekly", "2026-12-28", "2027-01-03"))

flags = dict(auto=False, daily=True, weekly=True, monthly=True, manual=True,
             daily_date="2026-09-15", weekly_reference_date="2026-09-24",
             monthly_month=2, monthly_year=2028, manual_start="2026-08-21", manual_end="2026-09-22")
jobs, errs = CP.plan_jobs(date(2026, 10, 1), **flags)
check("flag mode: all four, pinned dates", J(jobs),
      [("Daily", "2026-09-15", "2026-09-15"), ("Weekly", "2026-09-21", "2026-09-27"),
       ("Monthly", "2028-02-01", "2028-02-29"), ("Manual", "2026-08-21", "2026-09-22")])
check("flag mode: no errors", errs, [])
jobs, errs = CP.plan_jobs(date(2026, 10, 1), auto=False, daily=True)
check("flag mode: DAILY_DATE None -> today", J(jobs), [("Daily", "2026-10-01", "2026-10-01")])
jobs, errs = CP.plan_jobs(date(2026, 10, 1), auto=False, weekly=True, monthly=True)
check("flag mode: None week/month -> current week / month", J(jobs),
      [("Weekly", "2026-09-28", "2026-10-04"), ("Monthly", "2026-10-01", "2026-10-31")])
jobs, errs = CP.plan_jobs(date(2026, 10, 1), auto=False, manual=True, manual_start="2026-09-22",
                          manual_end="2026-08-21")
check("Manual Start > End -> skipped with an error", (J(jobs), len(errs)), ([], 1))
jobs, errs = CP.plan_jobs(date(2026, 10, 1), auto=False, manual=True, manual_start="2026-09-22")
check("Manual missing End -> skipped with an error", (J(jobs), len(errs)), ([], 1))
jobs, errs = CP.plan_jobs(date(2026, 10, 1), auto=False, manual=True, manual_start="2026-09-31",
                          manual_end="2026-10-02")
check("Manual invalid date -> skipped with an error", (J(jobs), len(errs)), ([], 1))
jobs, errs = CP.plan_jobs(date(2026, 10, 1), auto=False, manual=True, manual_start="2026-09-22",
                          manual_end="2026-09-22")
check("Manual single day (Start == End) allowed", J(jobs), [("Manual", "2026-09-22", "2026-09-22")])

# the performance report's own configuration block drives the same planner
P.GENERATE_AUTO = False
P.GENERATE_DAILY, P.GENERATE_WEEKLY, P.GENERATE_MONTHLY, P.GENERATE_MANUAL = False, False, False, True
check("performance report MANUAL_START/END config -> Manual 21-Aug..22-Sep",
      J(P._plan_jobs(date(2026, 10, 1))[0]), [("Manual", "2026-08-21", "2026-09-22")])
P.GENERATE_AUTO = True

print("\n== 2. Folder names ==")
check("Daily", CP.folder_path("Daily", date(2026, 10, 1), date(2026, 10, 1)),
      ("Daily Coordinator Reports", "Daily 01-Oct-2026"))
check("Weekly", CP.folder_path("Weekly", date(2026, 9, 21), date(2026, 9, 27)),
      ("Weekly Coordinator Reports", "Weekly 21-Sep-2026 to 27-Sep-2026"))
check("Monthly", CP.folder_path("Monthly", date(2026, 9, 1), date(2026, 9, 30)),
      ("Monthly Coordinator Reports", "Monthly Sep-2026"))
check("Manual", CP.folder_path("Manual", date(2026, 8, 21), date(2026, 9, 22)),
      ("Manual Coordinator Reports", "Manual 21-Aug-2026 to 22-Sep-2026"))
check("Daily folder name <-> date round trip",
      CP.daily_folder_date(CP.period_folder_name("Daily", date(2026, 10, 1), date(2026, 10, 1))),
      date(2026, 10, 1))


# =============================================================================
#  In-memory Drive (only the calls the Coordinator code makes)
# =============================================================================
class _Req:
    def __init__(self, fn):
        self.fn = fn

    def execute(self):
        return self.fn()


class FakeDrive:
    FOLDER = "application/vnd.google-apps.folder"

    def __init__(self, root):
        self.items = {root: {"id": root, "name": "ROOT", "mimeType": self.FOLDER, "parents": []}}
        self.n = 0
        self.clock = datetime(2026, 9, 1, 4, 0)       # UTC

    def files(self):
        return self

    def _new_id(self):
        self.n += 1
        return f"f{self.n}"

    def _match(self, it, q):
        for clause in [c.strip() for c in q.split(" and ")]:
            m = re.match(r"^'(.+)' in parents$", clause)
            if m:
                if m.group(1) not in it["parents"]:
                    return False
                continue
            m = re.match(r"^name\s*=\s*'(.*)'$", clause)
            if m:
                if it["name"] != m.group(1).replace("\\'", "'"):
                    return False
                continue
            m = re.match(r"^name contains '(.*)'$", clause)
            if m:
                if m.group(1).replace("\\'", "'") not in it["name"]:
                    return False
                continue
            m = re.match(r"^mimeType\s*=\s*'(.*)'$", clause)
            if m:
                if it["mimeType"] != m.group(1):
                    return False
                continue
            if clause.replace(" ", "") == "trashed=false":
                continue
            raise AssertionError(f"unsupported query clause: {clause}")
        return True

    def list(self, q=None, **kw):
        hits = [dict(it) for it in self.items.values() if self._match(it, q)]
        return _Req(lambda: {"files": hits})

    def create(self, body=None, media_body=None, fields=None, **kw):
        def go():
            fid = self._new_id()
            self.clock += timedelta(minutes=1)
            data = None
            if media_body is not None:
                data = media_body.getbytes(0, media_body.size())
            self.items[fid] = {"id": fid, "name": body["name"], "mimeType": body["mimeType"],
                               "parents": list(body.get("parents", [])), "data": data,
                               "createdTime": self.clock.strftime("%Y-%m-%dT%H:%M:%S.000Z"),
                               "modifiedTime": self.clock.strftime("%Y-%m-%dT%H:%M:%S.000Z")}
            return {"id": fid, "webViewLink": f"https://drive/{fid}"}
        return _Req(go)

    def export(self, fileId=None, mimeType=None):
        return _Req(lambda: self.items[fileId]["data"])

    def update(self, fileId=None, addParents=None, removeParents=None, **kw):
        def go():
            it = self.items[fileId]
            it["parents"] = [p for p in it["parents"] if p != removeParents] + [addParents]
            return {"id": fileId}
        return _Req(go)

    # helpers for the checks
    def path_of(self, fid):
        names, cur = [], self.items[fid]
        while cur["parents"]:
            names.append(cur["name"])
            cur = self.items[cur["parents"][0]]
        return "/".join(reversed(names))

    def files_in(self, *names):
        return sorted(self.path_of(i) for i, it in self.items.items()
                      if it["mimeType"] != self.FOLDER and self.path_of(i).startswith("/".join(names) + "/"))


drive = FakeDrive(BC.PARENT_FOLDER_ID)
BC._drive_client = lambda: drive
BC._enable_followup_timestamps = lambda sid: None
BC.PROTECT_SHEETS = False            # protection is checked in verify_coordinator_sheet_protection.py
P._drive_service = lambda: drive
P._call = lambda fn, what: fn()
_cache = tempfile.mkdtemp(prefix="coord_cache_")
P.CACHE_DIR = _cache


# =============================================================================
#  Coordinator task reports built with the REAL Coordinator builders
# =============================================================================
def adm(name, email):
    return {"Student Name": name, "Email ID": email, "Phone Number": "+919000000001",
            "Batch Name": "DAAI1026", "Joined On": "2026-10-01", "Request Form Name": "",
            "Recipient Status": "Form Not Sent", "Request Status": "", "Sent Date": "",
            "Signed Date": "", "Expiry Date": ""}


def task_report(day, names, done=None):
    """A daily Coordinator workbook (Admission tab) with follow-ups typed in:
    done = {name: (Yes/No, stamp)}."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    BC.build_admission_formalities(wb.create_sheet("Learner Admission Formalities"),
                                   [adm(n, n.split()[0].lower() + "@x.com") for n in names], day)
    ws = wb.worksheets[0]
    hdr = None
    for r in range(1, ws.max_row + 1):
        txt = [str(ws.cell(row=r, column=c).value or "") for c in range(1, ws.max_column + 1)]
        if "Why Flagged" in txt:
            hdr = {t.replace("✎", "").strip(): i + 1 for i, t in enumerate(txt) if t}
            continue
        for c in range(1, ws.max_column + 1):            # Google keeps values, not the formula
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.startswith("=IF("):
                ws.cell(row=r, column=c).value = None
        if hdr and txt[0] in (done or {}):
            d, stamp = done[txt[0]]
            ws.cell(row=r, column=hdr["Follow-Up Done?"]).value = d
            ws.cell(row=r, column=hdr["Follow-Up DateTime"]).value = stamp
    buf = io.BytesIO()
    wb.save(buf)
    return buf


def bc_upload_daily(day, names, done=None, clock_utc=None):
    if clock_utc:
        drive.clock = clock_utc - timedelta(minutes=1)
    fname = f"{BC.REPORT_BASENAME}_{(day - timedelta(days=1)).strftime('%d-%b-%Y')}_12.00_PM_-_{day.strftime('%d-%b-%Y')}_09.30_AM.xlsx"
    return BC.upload_datewise(day, fname, task_report(day, names, done))


print("\n== 3. Batch Coordinator report → Daily / period folders, versioning ==")
D0, D1, D2 = date(2026, 9, 20), date(2026, 9, 21), date(2026, 9, 22)    # Sun, Mon, Tue
# D0 (previous week): Asha + Bala pending all day — must never leak into the week of D1/D2
bc_upload_daily(D0, ["Asha Rao", "Bala K"], clock_utc=datetime(2026, 9, 20, 4, 0))
# D1: two versions; Asha done on time in V1, Chitra done in V2
bc_upload_daily(D1, ["Asha Rao", "Chitra M"], {"Asha Rao": ("Yes", datetime(2026, 9, 21, 11, 0))},
                clock_utc=datetime(2026, 9, 21, 4, 0))
bc_upload_daily(D1, ["Asha Rao", "Chitra M"], {"Chitra M": ("Yes", datetime(2026, 9, 21, 16, 0))},
                clock_utc=datetime(2026, 9, 21, 8, 0))
# D2: Chitra still pending, Dev new, Asha done
bc_upload_daily(D2, ["Asha Rao", "Chitra M", "Dev N"], {"Asha Rao": ("Yes", datetime(2026, 9, 22, 12, 0))},
                clock_utc=datetime(2026, 9, 22, 4, 0))
check("daily reports in Daily/<Daily DD-Mon-YYYY>, re-run = Version 2",
      drive.files_in("Daily Coordinator Reports", "Daily 21-Sep-2026"),
      ["Daily Coordinator Reports/Daily 21-Sep-2026/" + n for n in sorted([
          f"{BC.REPORT_BASENAME}_20-Sep-2026_12.00_PM_-_21-Sep-2026_09.30_AM",
          f"{BC.REPORT_BASENAME}_20-Sep-2026_12.00_PM_-_21-Sep-2026_09.30_AM - Version 2"])])
check("no daily report written to a legacy YYYY-MM-DD folder",
      [it["name"] for it in drive.items.values() if re.match(r"^\d{4}-\d{2}-\d{2}$", it["name"])], [])

# BC weekly roll-up for the same week goes to the Weekly period folder (as _generate_period does)
_wk = CP.folder_path("Weekly", date(2026, 9, 21), date(2026, 9, 27))
BC.upload_report(_wk, "IntelliBI_Batch_Coordinator_Weekly_Follow_Ups_21_Sep_to_27_Sep_2026.xlsx",
                 io.BytesIO(b"x"), "IntelliBI_Batch_Coordinator_Weekly_Follow_Ups")
# a LEGACY day folder from before the new layout (still read by the performance report)
_leg = BC._find_or_create_folder(drive, BC.PARENT_FOLDER_ID, D2.strftime("%Y-%m-%d"))
drive.clock = datetime(2026, 9, 22, 2, 0)
drive.create(body={"name": f"{BC.REPORT_BASENAME}_LEGACY", "parents": [_leg],
                   "mimeType": "application/vnd.google-apps.spreadsheet"},
             media_body=type("M", (), {"getbytes": lambda s, a, b: task_report(D2, ["Esha L"], {"Esha L": ("Yes", datetime(2026, 9, 22, 10, 0))}).getvalue(),
                                       "size": lambda s: 1})()).execute()

print("\n== 4. Performance report: discovery = task origin, period only ==")
vers = P.discover_report_versions(drive, [(D1, D2)])
check("discovers only daily task reports of the period (new + legacy layout, both versions of D1)",
      sorted((v["day"].isoformat(), v["version"], v["folder"]) for v in vers),
      [("2026-09-21", 1, "Daily 21-Sep-2026"), ("2026-09-21", 2, "Daily 21-Sep-2026"),
       ("2026-09-22", 1, "2026-09-22"), ("2026-09-22", 1, "Daily 22-Sep-2026")])
check("previous-week day (20-Sep) not even read", any(v["day"] == D0 for v in vers), False)

now = datetime(2026, 9, 28, 19, 0)       # Monday 7 PM IST
loader = lambda v: P.load_version(drive, v)
all_vers = P.discover_report_versions(drive, [(D0, D2)])
ledger = P.build_ledger(all_vers, loader, now)
wk = P.scope_ledger(ledger, date(2026, 9, 21), date(2026, 9, 27))
check("Weekly 21-27 Sep: only tasks that originated in the week (D0 excluded)",
      sorted({t["day"].isoformat() for t in wk["tasks"]}), ["2026-09-21", "2026-09-22"])
check("Weekly task count: D1 (Asha, Chitra) + D2 (Asha, Chitra, Dev, Esha legacy)", len(wk["tasks"]), 6)
s = P.summarise(wk["tasks"])
check("Weekly completed / missed (Asha D1, Chitra D1, Asha D2, Esha D2 done; Chitra D2, Dev D2 missed)",
      (s["completed"], s["missed"], s["open"]), (4, 2, 0))
chitra2 = [t for t in wk["tasks"] if t["day"] == D2 and t["label"].startswith("Chitra")][0]
check("Days on List counted within the period only (Chitra D2 = 1: D1 Chitra was done)", chitra2["days_on_list"], 1)
asha_d0 = [t for t in ledger["tasks"] if t["day"] == D0 and t["label"].startswith("Asha")][0]
check("D0 Asha pending (exists in history) …", asha_d0["completed"], False)
dl = P.scope_ledger(ledger, D2, D2)
check("Daily 22-Sep: only that day's 4 tasks", (len(dl["tasks"]), {t["day"] for t in dl["tasks"]}), (4, {D2}))
check("Daily 22-Sep: Days on List never looks back into 21-Sep",
      max(t["days_on_list"] for t in dl["tasks"]), 1)
mo = P.scope_ledger(ledger, date(2026, 9, 1), date(2026, 9, 30))
check("Monthly Sep includes D0, D1, D2", sorted({t["day"].isoformat() for t in mo["tasks"]}),
      ["2026-09-20", "2026-09-21", "2026-09-22"])
oc = P.scope_ledger(ledger, date(2026, 10, 1), date(2026, 10, 31))
check("Monthly Oct: no September task leaks in", len(oc["tasks"]), 0)
check("coverage rows limited to the period", [c["day"] for c in wk["coverage"]], [D1, D2])

print("\n== 5. Performance report delivered next to the Batch Coordinator report ==")
jobs = [CP.make_job("Daily", D2, D2), CP.make_job("Weekly", date(2026, 9, 21), date(2026, 9, 27))]
up = lambda folders, name, buf: BC.upload_report(folders, name + ".xlsx", buf, name)
outdir = tempfile.mkdtemp(prefix="coord_out_")
_cwd = os.getcwd()
os.chdir(outdir)                       # any stray relative write would land here
res = P.run_jobs(jobs, P.discover_report_versions(drive, [(j["start"], j["end"]) for j in jobs]),
                 loader, now, upload=up)
os.chdir(_cwd)
check("both jobs delivered", [r.get("failed", False) for r in res], [False, False])
check("Daily 22-Sep folder: task report + performance report together",
      [p.split("/")[-1].split("_2")[0] for p in drive.files_in("Daily Coordinator Reports", "Daily 22-Sep-2026")],
      ["IntelliBI_Batch_Coordinator_Daily_Attendance_Report", "IntelliBI_Coordinator_Task_Performance_Report_Daily"])
check("Weekly folder: Batch Coordinator weekly roll-up + weekly performance report together",
      drive.files_in(*_wk),
      sorted(["/".join(_wk) + "/IntelliBI_Batch_Coordinator_Weekly_Follow_Ups_21_Sep_to_27_Sep_2026",
              "/".join(_wk) + "/IntelliBI_Coordinator_Task_Performance_Report_Weekly_21_Sep_2026_to_27_Sep_2026"]))
check("Daily performance counts only 22-Sep tasks", res[0]["summary"]["tasks"], 4)
check("Weekly performance counts only the week's tasks", res[1]["summary"]["tasks"], 6)
check("Drive only: no local copy, nothing written to disk",
      ([r.get("local") for r in res], [f for _d, _s, fs in os.walk(outdir) for f in fs]),
      ([None, None], []))
res2 = P.run_jobs(jobs[:1], P.discover_report_versions(drive, [(D2, D2)]), loader, now, upload=up)
check("re-run -> Version 2, Version 1 kept",
      [p.split("/")[-1] for p in drive.files_in("Daily Coordinator Reports", "Daily 22-Sep-2026")
       if "Performance" in p],
      ["IntelliBI_Coordinator_Task_Performance_Report_Daily_22_Sep_2026",
       "IntelliBI_Coordinator_Task_Performance_Report_Daily_22_Sep_2026 - Version 2"])
check("performance reports are never read back as task lists",
      any("Performance" in v["name"] for v in P.discover_report_versions(drive, [(D0, D2)])), False)
check("Batch Coordinator versioning unaffected by the performance files in the same folder",
      [p.split(" - ")[-1] if " - Version" in p else "v1"
       for p in drive.files_in("Daily Coordinator Reports", "Daily 21-Sep-2026")], ["v1", "Version 2"])

# a failing job is reported, the next job still runs
def _boom(folders, name, buf):
    if "Daily" in name:
        raise RuntimeError("simulated Drive outage")
    return up(folders, name, buf)
res3 = P.run_jobs(jobs, P.discover_report_versions(drive, [(D1, D2)]), loader, now, upload=_boom)
check("failed job recorded, other job still delivered", [bool(r.get("failed")) for r in res3], [True, False])

print("\n== 6. Legacy migration classifier ==")
check("legacy daily file -> Daily folder",
      CP.classify_legacy_file(f"{BC.REPORT_BASENAME}_x - Version 3", date(2026, 9, 30)),
      ("Daily", date(2026, 9, 30), date(2026, 9, 30)))
check("legacy monthly roll-up -> Monthly folder",
      CP.classify_legacy_file("IntelliBI_Batch_Coordinator_Monthly_Follow_Ups_September_2026 - Version 8", None),
      ("Monthly", date(2026, 9, 1), date(2026, 9, 30)))
check("unrelated files stay put",
      CP.classify_legacy_file("Coordinator Attendance Tasks — 2026-09-30", date(2026, 9, 30)), None)

print("\n== 7. Scheduler integration ==")
import run_reports_action as RRA                                    # noqa: E402
bc_job = [j for j in RRA.JOBS if j[0] == "pyCoordinatorTaskListReport"]
check("Batch Coordinator report is in the Morning batch (Layer 2)", len(bc_job), 1)
check("… gated on all four Layer-1 refreshes", sorted(bc_job[0][2]),
      sorted(["pyAssignmentSubmissions", "pySessionAttendanceStudentTeacherFeedbacks",
              "pyStudentPaymentClassesStudentEnrolled", "pyZohoSignatureStatusRefresh"]))
check("… resolved from 'co-ordinator reports/'",
      os.path.basename(os.path.dirname(str(RRA.script_path(bc_job[0][0])))), "co-ordinator reports")
check("existing Layer-2 scripts still resolved from ops_reports_action/",
      os.path.basename(os.path.dirname(str(RRA.script_path("pyAttendaceFeedbackReport")))), "ops_reports_action")
import run_evening_reports as EVE                                   # noqa: E402
check("Evening batch runs the Coordinator Task Performance report",
      [s for s, _l in EVE.JOBS], ["pyCoordinatorTaskPerformanceReport"])
check("Evening job script exists", os.path.exists(os.path.join(
    ROOT, "co-ordinator reports", "pyCoordinatorTaskPerformanceReport.py")), True)

import run_scheduled as RS                                          # noqa: E402
import paths as PATHS                                               # noqa: E402
tmp_scripts = tempfile.mkdtemp(prefix="sched_")
tmp_cache = tempfile.mkdtemp(prefix="schedc_")
with open(os.path.join(tmp_scripts, "run_evening_reports.py"), "w") as fh:
    fh.write("import sys, os\nopen(os.environ['CALLED'], 'a').write('evening\\n')\nsys.exit(0)\n")
with open(os.path.join(tmp_scripts, "run_all.py"), "w") as fh:
    fh.write("import sys, os\nopen(os.environ['CALLED'], 'a').write('morning\\n')\nsys.exit(1)\n")
called = os.path.join(tmp_cache, "called.txt")
os.environ["CALLED"] = called
from pathlib import Path                                            # noqa: E402
_orig = (PATHS.SCRIPTS_DIR, PATHS.CACHE_DIR, PATHS.LOGS_DIR)
PATHS.SCRIPTS_DIR, PATHS.CACHE_DIR = Path(tmp_scripts), Path(tmp_cache)
PATHS.LOGS_DIR = Path(tmp_cache) / "logs"       # keep test triggers out of the real logs/


def _sched(*argv):
    sys.argv = ["run_scheduled.py", *argv]
    return RS.main()


rc1 = _sched("--label", "ops_evening", "--once-per-day", "--entry", "run_evening_reports.py")
rc2 = _sched("--label", "ops_evening", "--once-per-day", "--entry", "run_evening_reports.py")
rc3 = _sched("--label", "ops", "--once-per-day")
rc4 = _sched("--label", "ops", "--once-per-day")
calls = open(called).read().split()
check("evening entry launched once; 20:00 retry skipped after success", (rc1, rc2, calls.count("evening")), (0, 0, 1))
check("morning batch (default entry) independent label; failure → retry window runs again",
      (rc3, rc4, calls.count("morning")), (1, 1, 2))
check("unknown entry rejected", _sched("--label", "x", "--entry", "nope.py"), 2)
PATHS.SCRIPTS_DIR, PATHS.CACHE_DIR, PATHS.LOGS_DIR = _orig

with open(os.path.join(ROOT, "scripts", "setup_schedule.ps1"), encoding="utf-8") as fh:
    ps = fh.read()
check("setup_schedule.ps1: evening task at 19:00 (20:00 retry) with the evening entry",
      ('@("19:00","20:00")' in ps, "--label ops_evening --once-per-day --entry run_evening_reports.py" in ps,
       '@("10:30","11:30")' in ps), (True, True, True))

shutil.rmtree(_cache, ignore_errors=True)
shutil.rmtree(outdir, ignore_errors=True)
shutil.rmtree(tmp_scripts, ignore_errors=True)
shutil.rmtree(tmp_cache, ignore_errors=True)
print("\nALL CHECKS PASSED" if not FAIL else f"\n{len(FAIL)} CHECK(S) FAILED: {FAIL}")
sys.exit(1 if FAIL else 0)
