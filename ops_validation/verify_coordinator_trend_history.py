"""
Verification of WITHIN-DAY PROGRESS (Progress Trend tab) for HISTORICAL report
days in pyCoordinatorTaskPerformanceReport.py
(run from the project root:  python ops_validation\\verify_coordinator_trend_history.py).

Builds REAL Coordinator task lists with pyCoordinatorTaskListReport's builders,
fills follow-ups the way the Coordinator does, and builds the performance
workbook for a past Daily date while "now" is a later day. Checks:
  * a day whose task list existed during the day → hour-by-hour figures from
    THAT day's activity only (not today's hours), late completions explained
  * a day whose task list was produced only AFTER the day ended (re-run for a
    past date) → no blank section: a clear explanation with the task count and
    when the list was first generated; tasks still on Dashboard / Task Register
  * a mix of both → table for the on-day tasks + a note for the later ones
  * no task list at all → the existing "No trackable tasks" message
  * Weekly / multi-day pooling still works and explains after-day tasks
No Google access needed.
"""
import os
import sys
import types
from datetime import date, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.modules.setdefault("_bootstrap", types.ModuleType("_bootstrap"))
paths = types.ModuleType("paths")
paths.CREDENTIALS_DIR = paths.CONFIG_DIR = paths.LOGS_DIR = paths.CACHE_DIR = HERE
sys.modules.setdefault("paths", paths)
ec = types.ModuleType("email_config"); ec.GMAIL_SENDER = "x@x"; ec.GMAIL_APP_PASS = "x"
sys.modules["email_config"] = ec
for p in ("common", "ops_reports_action", "co-ordinator reports"):
    sys.path.insert(0, os.path.join(ROOT, p))

import openpyxl                                          # noqa: E402
import pyCoordinatorTaskListReport as BC                 # noqa: E402
import pyCoordinatorTaskPerformanceReport as P           # noqa: E402

FAIL = []


def check(label, got, want):
    ok = got == want
    print(f"[{'pass' if ok else 'FAIL'}] {label}: {got!r}" + ("" if ok else f"  (expected {want!r})"))
    if not ok:
        FAIL.append(label)


def adm(name):
    return {"Student Name": name, "Email ID": name.split()[0].lower() + "@x.com",
            "Phone Number": "+919000000001", "Batch Name": "DAAI1026", "Joined On": "2026-10-01",
            "Request Form Name": "", "Recipient Status": "Form Not Sent", "Request Status": "",
            "Sent Date": "", "Signed Date": "", "Expiry Date": ""}


ADM = "Learner Admission Formalities"


def task_list(day, names):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    BC.build_admission_formalities(wb.create_sheet(ADM), [adm(n) for n in names], day)
    for ws in wb.worksheets:                       # Google exports the stamp VALUE, not the formula
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=IF("):
                    c.value = None
    return wb


def fill(wb, who, done, stamp):
    ws = wb[ADM]
    hdr = None
    for r in range(1, ws.max_row + 1):
        txt = [str(ws.cell(row=r, column=c).value or "") for c in range(1, ws.max_column + 1)]
        if any(t.strip() == "Why Flagged" for t in txt):
            hdr = {t.replace("✎", "").strip(): i + 1 for i, t in enumerate(txt) if t}
            continue
        if hdr and who in txt:
            ws.cell(row=r, column=hdr["Follow-Up Done?"]).value = done
            ws.cell(row=r, column=hdr["Follow-Up DateTime"]).value = stamp
            ws.cell(row=r, column=hdr["Action Taken"]).value = "Call"
            return
    raise AssertionError(who)


def trend_text(ledger, start, end, is_daily, now):
    wb, tasks, s = P.build_workbook(ledger, start, end, "Test", is_daily, now)
    ws = wb["Progress Trend"]
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    return cells, tasks, s, wb


def ledger_of(books, versions, now):
    return P.build_ledger(versions, lambda v: P.parse_report_workbook(books[v["id"]]), now)


D, NEXT = date(2026, 10, 2), date(2026, 10, 3)
NOW = datetime(2026, 10, 3, 16, 21)                # the run: the day after the report day

# =============================================================================
print("\n== 1. Task list produced only AFTER the report day (re-run for a past date) ==")
late = task_list(D, ["Asha Rao", "Bala K", "Chitra M"])
fill(late, "Asha Rao", "Yes", datetime(2026, 10, 3, 16, 0))      # actioned next day
books = {"late": late}
vers = [{"id": "late", "name": "late", "day": D, "version": 1,
         "created": datetime(2026, 10, 3, 15, 26), "modified_raw": ""}]
L = ledger_of(books, vers, NOW)
cells, tasks, s, wb = trend_text(L, D, D, True, NOW)
txt = " ".join(cells)
check("ledger has the day's 3 tasks", len(tasks), 3)
check("hourly figures empty (nothing happened on the report day)", P.hourly_progress(tasks, [D], NOW), [])
check("NOT the misleading 'No trackable tasks' message", "No trackable tasks" in txt, False)
check("explains why, with the first-generated time",
      all(x in txt for x in ("No within-day progress can be shown for 02-Oct-2026",
                             "generated only after the day had ended",
                             "03-Oct-2026 03:26 PM", "none of its 3 task(s)")), True)
check("points to where the tasks are counted", "Dashboard and the Task Register" in txt, True)
check("tasks still counted (Dashboard summary): 1 completed late, 2 missed",
      (s["tasks"], s["completed"], s["missed"]), (3, 1, 2))
check("Task Register lists all 3", sum(1 for row in wb["Task Register"].iter_rows(values_only=True)
                                       if any(str(v) in ("Asha Rao", "Bala K", "Chitra M")
                                              or str(v).startswith(("Asha Rao", "Bala K", "Chitra M"))
                                              for v in row if v)), 3)

# =============================================================================
print("\n== 2. Task list existed during the day → that day's own hours ==")
v1 = task_list(D, ["Asha Rao", "Bala K", "Chitra M"])
fill(v1, "Asha Rao", "Yes", datetime(2026, 10, 2, 11, 15))
fill(v1, "Bala K", "Yes", datetime(2026, 10, 2, 14, 40))
fill(v1, "Chitra M", "Yes", datetime(2026, 10, 3, 9, 5))          # next morning → late
books = {"v1": v1}
vers = [{"id": "v1", "name": "v1", "day": D, "version": 1,
         "created": datetime(2026, 10, 2, 10, 32), "modified_raw": ""}]
L = ledger_of(books, vers, NOW)
hp = P.hourly_progress(L["tasks"], [D], NOW)
check("hours run from generation (10 AM) to the last on-day completion (2 PM) — not to now (4 PM)",
      [r["label"] for r in hp], ["10 AM", "11 AM", "12 PM", "1 PM", "2 PM"])
check("cumulative generated / completed / open at 2 PM",
      (hp[-1]["generated"], hp[-1]["completed"], hp[-1]["open"]), (3, 2, 1))
cells, tasks, s, wb = trend_text(L, D, D, True, NOW)
txt = " ".join(cells)
check("table present", "Hour (IST)" in txt, True)
check("late completion explained under the table",
      "1 completion(s) were recorded after the report day" in txt, True)
check("no 'generated after' note when the list existed during the day",
      "generated only after" in txt, False)

# =============================================================================
print("\n== 3. Mixed: on-day version + a new task first listed the next day ==")
m1 = task_list(D, ["Asha Rao", "Bala K"])
fill(m1, "Asha Rao", "Yes", datetime(2026, 10, 2, 12, 0))
m2 = task_list(D, ["Asha Rao", "Bala K", "Dev N"])                   # re-run on 03-Oct adds Dev
books = {"m1": m1, "m2": m2}
vers = [{"id": "m1", "name": "m1", "day": D, "version": 1, "created": datetime(2026, 10, 2, 10, 30),
         "modified_raw": ""},
        {"id": "m2", "name": "m2", "day": D, "version": 2, "created": datetime(2026, 10, 3, 15, 26),
         "modified_raw": ""}]
L = ledger_of(books, vers, NOW)
hp = P.hourly_progress(L["tasks"], [D], NOW)
check("hours only from the on-day version's 2 tasks", (hp[-1]["generated"], hp[-1]["completed"]), (2, 1))
cells, tasks, s, wb = trend_text(L, D, D, True, NOW)
txt = " ".join(cells)
check("table + note for the task listed after the day",
      ("Hour (IST)" in txt, "1 task(s) were generated only after their report day had ended" in txt),
      (True, True))
check("all 3 tasks still in the report", s["tasks"], 3)

# =============================================================================
print("\n== 4. No task list for the day ==")
L = ledger_of({}, [], NOW)
cells, tasks, s, wb = trend_text(L, D, D, True, NOW)
check("existing message unchanged", "No trackable tasks on this report day." in " ".join(cells), True)

# =============================================================================
print("\n== 5. Period (Weekly) pooling unchanged, after-day tasks explained ==")
w1 = task_list(date(2026, 10, 1), ["Asha Rao"])
fill(w1, "Asha Rao", "Yes", datetime(2026, 10, 1, 11, 0))
books = {"w1": w1, "late": late}
vers = [{"id": "w1", "name": "w1", "day": date(2026, 10, 1), "version": 1,
         "created": datetime(2026, 10, 1, 10, 30), "modified_raw": ""},
        {"id": "late", "name": "late", "day": D, "version": 1,
         "created": datetime(2026, 10, 3, 15, 26), "modified_raw": ""}]
L = ledger_of(books, vers, NOW)
hp = P.hourly_progress(L["tasks"], [date(2026, 10, 1), D], NOW)
check("pooled hours from 01-Oct only", [(r["label"], r["generated"], r["completed"]) for r in hp],
      [("10 AM", 1, 0), ("11 AM", 1, 1)])
cells, tasks, s, wb = trend_text(L, date(2026, 9, 28), date(2026, 10, 4), False, NOW)
txt = " ".join(cells)
check("period: table + note for 02-Oct's after-day tasks",
      ("Hour (IST)" in txt, "3 task(s) were generated only after their report day had ended" in txt),
      (True, True))
check("day-wise section still present", "DAY-WISE PROGRESS" in txt, True)

print("\nALL CHECKS PASSED" if not FAIL else f"\n{len(FAIL)} CHECK(S) FAILED: {FAIL}")
sys.exit(1 if FAIL else 0)
